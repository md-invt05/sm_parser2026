"""Sui Move-package discovery and conservative asset accounting via Blockberry.

Sui is deliberately kept outside the EVM address/coverage model.  A package is
not an EVM account and a package's assets can live in shared or object-owned
state objects.  The totals produced here are therefore documented lower bounds.
"""

from __future__ import annotations

import asyncio
import json
import math
import logging
import os
import random
import sqlite3
import threading
import time
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable

import httpx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


log = logging.getLogger("scanner.sui")
UTC = timezone.utc
BLOCKBERRY_BASE = "https://api.blockberry.one/sui/v1"
SYSTEM_PACKAGES = {f"0x{value:064x}" for value in (0, 1, 2, 3)}


def utc_now() -> str:
    return datetime.now(UTC).isoformat()


def normalize_sui_id(value: str) -> str:
    value = str(value or "").strip().lower()
    if not value.startswith("0x"):
        value = "0x" + value
    body = value[2:]
    if not body or len(body) > 64 or any(ch not in "0123456789abcdef" for ch in body):
        raise ValueError(f"invalid Sui id: {value[:18]}")
    return "0x" + body.zfill(64)


def normalize_coin_type(value: str) -> str:
    parts = str(value or "").strip().split("::")
    if len(parts) < 3:
        raise ValueError("invalid Sui coin type")
    parts[0] = normalize_sui_id(parts[0])
    return "::".join(parts)


def _owner_kind(owner: Any) -> tuple[str | None, str | None]:
    if not isinstance(owner, dict):
        return None, None
    if "Shared" in owner:
        return "shared", None
    nested = owner.get("ObjectOwner")
    if nested:
        try:
            return "object", normalize_sui_id(str(nested))
        except ValueError:
            return "object", None
    # Blockberry/Sui clients sometimes use snake/camel-case variants.
    nested = owner.get("objectOwner") or owner.get("object_owner")
    if nested:
        try:
            return "object", normalize_sui_id(str(nested))
        except ValueError:
            return "object", None
    return None, None


def parse_raw_transaction(payload: dict[str, Any]) -> dict[str, Any]:
    """Extract directly called/published packages and attributable state objects."""
    result = payload.get("result", payload)
    transaction = (((result.get("transaction") or {}).get("data") or {}).get("transaction") or {})
    commands = transaction.get("transactions") or []
    calls: set[str] = set()
    saw_publish = False
    for command in commands:
        if not isinstance(command, dict):
            continue
        move = command.get("MoveCall") or command.get("moveCall")
        if isinstance(move, dict) and move.get("package"):
            try:
                package = normalize_sui_id(move["package"])
                if package not in SYSTEM_PACKAGES:
                    calls.add(package)
            except ValueError:
                pass
        if "Publish" in command or "publish" in command:
            saw_publish = True

    published: set[str] = set()
    objects: list[dict[str, Any]] = []
    for change in result.get("objectChanges") or []:
        if not isinstance(change, dict):
            continue
        kind = str(change.get("type") or "").lower()
        if kind == "published":
            candidate = change.get("packageId") or change.get("package") or change.get("objectId")
            if candidate:
                try:
                    package = normalize_sui_id(candidate)
                    if package not in SYSTEM_PACKAGES:
                        published.add(package)
                except ValueError:
                    pass
            continue
        object_type = str(change.get("objectType") or "")
        if "::" not in object_type or not change.get("objectId"):
            continue
        try:
            package = normalize_sui_id(object_type.split("::", 1)[0])
            object_id = normalize_sui_id(change["objectId"])
        except ValueError:
            continue
        owner_kind, parent = _owner_kind(change.get("owner"))
        if package in SYSTEM_PACKAGES or owner_kind not in {"shared", "object"}:
            continue
        objects.append({
            "package_id": package,
            "object_id": object_id,
            "object_type": object_type,
            "owner_kind": owner_kind,
            "parent_object": parent,
        })
    return {
        "calls": calls,
        "published": published,
        "objects": objects,
        "saw_publish": saw_publish,
    }


@dataclass
class SuiConfig:
    enabled: bool = True
    discovery_days: int = 30
    poll_sec: int = 30
    page_size: int = 100
    max_pages: int = 10
    raw_enrich_per_pass: int = 100
    balance_recheck_sec: int = 86400
    balance_concurrency: int = 2
    timeout_sec: float = 30.0
    max_retries: int = 6
    requests_per_second: float = 2.0
    defi_sync_sec: int = 900
    pool_verify_sec: int = 21600
    pool_limit: int = 5000
    pool_min_liquidity_usd: float = 25000.0
    tvl_anomaly_usd: float = 10_000_000_000.0

    @classmethod
    def from_mapping(cls, value: dict[str, Any] | None) -> "SuiConfig":
        raw = value or {}
        return cls(
            enabled=bool(raw.get("enabled", True)),
            discovery_days=max(1, int(raw.get("discovery_days", 30))),
            poll_sec=max(10, int(raw.get("poll_sec", 30))),
            page_size=max(1, min(100, int(raw.get("page_size", 100)))),
            max_pages=max(1, int(raw.get("max_pages", 500))),
            raw_enrich_per_pass=max(0, int(raw.get("raw_enrich_per_pass", 100))),
            balance_recheck_sec=max(60, int(raw.get("balance_recheck_sec", 86400))),
            balance_concurrency=1,
            timeout_sec=max(1.0, float(raw.get("timeout_sec", 30))),
            max_retries=max(1, min(3, int(raw.get("max_retries", 3)))),
            requests_per_second=max(0.1, min(2.0, float(raw.get("requests_per_second", 2)))),
            defi_sync_sec=max(60, int(raw.get("defi_sync_sec", 900))),
            pool_verify_sec=max(300, int(raw.get("pool_verify_sec", 21600))),
            pool_limit=max(1, int(raw.get("pool_limit", 5000))),
            pool_min_liquidity_usd=max(
                0.0, float(raw.get("pool_min_liquidity_usd", 25000))
            ),
            tvl_anomaly_usd=max(1.0, float(raw.get("tvl_anomaly_usd", 10_000_000_000))),
        )


class BlockberryError(RuntimeError):
    def __init__(
        self, kind: str, message: str, *, permanent: bool = False,
        status_code: int | None = None, endpoint: str | None = None,
    ):
        super().__init__(message)
        self.kind = kind
        self.permanent = permanent
        self.status_code = status_code
        self.endpoint = endpoint


class BlockberryClient:
    """Small, rate-limit-aware client. URLs/keys are never included in logs."""

    def __init__(self, api_key: str, cfg: SuiConfig, global_sem: asyncio.Semaphore | None = None,
                 fallback_urls: list[str] | None = None):
        if not api_key.strip():
            raise ValueError("BLOCKBERRY_API_KEY is required for Sui")
        self.api_key = api_key.strip()
        self.cfg = cfg
        self.global_sem = global_sem or asyncio.Semaphore(cfg.balance_concurrency)
        self.network_sem = asyncio.Semaphore(1)
        self.client: httpx.AsyncClient | None = None
        self.fallback_urls = fallback_urls if fallback_urls is not None else [
            value.strip() for value in os.getenv("SUI_RPC", "").split(",") if value.strip()
        ]
        self.fallback_client: httpx.AsyncClient | None = None
        self._metrics: list[tuple[bool, float, str]] = []
        self._rate_lock = asyncio.Lock()
        self._next_request_at = 0.0

    async def __aenter__(self) -> "BlockberryClient":
        self.client = httpx.AsyncClient(
            base_url=BLOCKBERRY_BASE,
            headers={"x-api-key": self.api_key, "accept": "application/json"},
            timeout=self.cfg.timeout_sec,
        )
        if self.fallback_urls:
            # Separate client is intentional: never forward the Blockberry key to RPC providers.
            self.fallback_client = httpx.AsyncClient(timeout=self.cfg.timeout_sec)
        return self

    async def __aexit__(self, *_args: Any) -> None:
        if self.client:
            await self.client.aclose()
        if self.fallback_client:
            await self.fallback_client.aclose()

    async def request(self, method: str, path: str, **kwargs: Any) -> Any:
        if self.client is None:
            raise RuntimeError("Blockberry client is not open")
        last: Exception | None = None
        for attempt in range(self.cfg.max_retries):
            started = time.perf_counter()
            kind = "ok"
            try:
                async with self._rate_lock:
                    delay = max(0.0, self._next_request_at - time.monotonic())
                    if delay:
                        await asyncio.sleep(delay)
                    self._next_request_at = time.monotonic() + 1.0 / self.cfg.requests_per_second
                async with self.global_sem, self.network_sem:
                    response = await self.client.request(method, path, **kwargs)
                latency = (time.perf_counter() - started) * 1000
                if response.status_code in (401, 403):
                    kind = "auth"
                    self._metrics.append((False, latency, kind))
                    raise BlockberryError(
                        kind, f"Blockberry {path} HTTP {response.status_code}",
                        permanent=True, status_code=response.status_code, endpoint=path,
                    )
                if response.status_code == 429:
                    kind = "rate_limit"
                    self._metrics.append((False, latency, kind))
                    wait = response.headers.get("retry-after")
                    delay = float(wait) if wait and wait.isdigit() else min(30.0, 1.5 * (2 ** attempt))
                    await asyncio.sleep(delay + random.random())
                    continue
                if response.status_code >= 500:
                    kind = "server"
                    self._metrics.append((False, latency, kind))
                    await asyncio.sleep(min(20.0, 2 ** attempt) + random.random())
                    continue
                if response.status_code >= 400:
                    kind = {
                        400: "bad_request", 404: "not_found", 405: "method_not_allowed",
                    }.get(response.status_code, "client_error")
                    self._metrics.append((False, latency, kind))
                    raise BlockberryError(
                        kind, f"Blockberry {path} HTTP {response.status_code}",
                        permanent=response.status_code in {400, 404, 405},
                        status_code=response.status_code, endpoint=path,
                    )
                try:
                    data = response.json()
                except ValueError as exc:
                    kind = "malformed_json"
                    self._metrics.append((False, latency, kind))
                    raise BlockberryError(kind, "Blockberry returned malformed JSON") from exc
                if isinstance(data, dict) and data.get("error"):
                    kind = "api_error"
                    self._metrics.append((False, latency, kind))
                    raise BlockberryError(kind, "Blockberry returned an API error")
                self._metrics.append((True, latency, "ok"))
                return data
            except BlockberryError:
                raise
            except (httpx.TimeoutException, httpx.TransportError) as exc:
                last = exc
                kind = "timeout" if isinstance(exc, httpx.TimeoutException) else "network"
                self._metrics.append((False, (time.perf_counter() - started) * 1000, kind))
                if attempt + 1 < self.cfg.max_retries:
                    await asyncio.sleep(min(20.0, 2 ** attempt) + random.random())
        raise BlockberryError("unavailable", f"Blockberry unavailable ({type(last).__name__})")

    async def transactions(self, page: int, size: int) -> dict[str, Any]:
        data = await self.request(
            "POST", "/transactions",
            params={"page": page, "size": size, "orderBy": "DESC", "sortBy": "AGE"},
            json={},
        )
        if not isinstance(data, dict) or not isinstance(data.get("content"), list):
            raise BlockberryError("partial", "transaction page has no content list")
        return data

    async def packages(self, page: int, size: int) -> dict[str, Any]:
        data = await self.request(
            "GET", "/packages",
            params={"page": page, "size": size, "orderBy": "DESC", "sortBy": "AGE"},
        )
        if not isinstance(data, dict) or not isinstance(data.get("content"), list):
            raise BlockberryError("partial", "package page has no content list")
        return data

    async def raw_transaction(self, digest: str) -> dict[str, Any]:
        data = await self.request("GET", f"/raw-transactions/{digest}")
        if not isinstance(data, dict) or not isinstance(data.get("result"), dict):
            raise BlockberryError("partial", "raw transaction has no result")
        return data

    async def account_balance(self, owner: str) -> list[dict[str, Any]]:
        normalized = normalize_sui_id(owner)
        try:
            data = await self.request("GET", f"/accounts/{normalized}/balance")
        except BlockberryError:
            if not self.fallback_urls:
                raise
            return await self._fallback_all_balances(normalized)
        rows = data if isinstance(data, list) else data.get("value") if isinstance(data, dict) else None
        if not isinstance(rows, list):
            raise BlockberryError("partial", "account balance has no value list")
        return [row for row in rows if isinstance(row, dict)]

    @staticmethod
    def _rows(data: Any, label: str) -> list[dict[str, Any]]:
        if isinstance(data, list):
            rows = data
        elif isinstance(data, dict):
            rows = None
            for key in ("content", "value", "data"):
                if key in data and data[key] is not None:
                    rows = data[key]
                    break
            if isinstance(rows, dict):
                nested = rows
                rows = None
                for key in ("content", "value", "items"):
                    if key in nested and nested[key] is not None:
                        rows = nested[key]
                        break
        else:
            rows = None
        if not isinstance(rows, list):
            raise BlockberryError("partial", f"{label} has no result list")
        return [row for row in rows if isinstance(row, dict)]

    async def _dex_page(self, page: int, size: int) -> dict[str, Any]:
        data = await self.request(
            "POST", "/dex",
            params={
                "page": page, "size": size, "orderBy": "DESC",
                "period": "DAY", "sortBy": "CURRENT_TVL",
            },
            json={"withTvlOnly": False},
        )
        if not isinstance(data, dict) or not isinstance(data.get("content"), list):
            raise BlockberryError("partial", "DEX project page has no content list")
        return data

    async def defi_projects(self) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        for page in range(self.cfg.max_pages):
            data = await self._dex_page(page, self.cfg.page_size)
            page_rows = [row for row in data["content"] if isinstance(row, dict)]
            rows.extend(page_rows)
            total_pages = data.get("totalPages")
            complete = bool(data.get("last")) or not page_rows
            if isinstance(total_pages, int) and page + 1 >= total_pages:
                complete = True
            if complete:
                if not rows:
                    raise BlockberryError("partial", "DEX project snapshot is empty")
                return rows
        raise BlockberryError(
            "partial", f"DEX project snapshot exceeded {self.cfg.max_pages} pages"
        )

    async def dex_pools(self, page: int, size: int) -> list[dict[str, Any]]:
        data = await self.request(
            "POST", "/dex/pools",
            params={
                "page": page, "size": size, "orderBy": "DESC",
                "period": "DAY", "sortBy": "LIQUIDITY_IN_USD",
            },
            json={"poolFactoryId": [], "poolFactoryEmpty": True},
        )
        return self._rows(data, "DEX pools")

    async def object_lookup(self, object_id: str) -> dict[str, Any] | None:
        if self.fallback_client is None:
            self.fallback_client = httpx.AsyncClient(timeout=self.cfg.timeout_sec)
        normalized = normalize_sui_id(object_id)
        for index, url in enumerate(self.fallback_urls):
            try:
                async with self.global_sem, self.network_sem:
                    response = await self.fallback_client.post(url, json={
                        "jsonrpc": "2.0", "id": index + 501,
                        "method": "sui_getObject",
                        "params": [normalized, {"showType": True, "showOwner": True}],
                    })
                response.raise_for_status()
                payload = response.json()
                result = payload.get("result") if isinstance(payload, dict) else None
                if isinstance(result, dict) and not result.get("error"):
                    return result
            except (httpx.HTTPError, ValueError):
                continue
        return None

    async def _fallback_all_balances(self, owner: str) -> list[dict[str, Any]]:
        """Alchemy/standard Sui RPC fallback; amounts stay unpriced by design."""
        if self.fallback_client is None:
            self.fallback_client = httpx.AsyncClient(timeout=self.cfg.timeout_sec)
        last: Exception | None = None
        for index, url in enumerate(self.fallback_urls):
            started = time.perf_counter()
            try:
                async with self.global_sem, self.network_sem:
                    response = await self.fallback_client.post(url, json={
                        "jsonrpc": "2.0", "id": index + 1,
                        "method": "suix_getAllBalances", "params": [owner],
                    })
                response.raise_for_status()
                payload = response.json()
                if payload.get("error") or not isinstance(payload.get("result"), list):
                    raise BlockberryError("rpc_error", "Sui fallback returned an RPC error")
                self._metrics.append((True, (time.perf_counter() - started) * 1000, "fallback"))
                return [{
                    "coinType": row.get("coinType"), "coinSymbol": None,
                    "decimals": None, "balance": row.get("totalBalance", "0"),
                    "coinPrice": None, "balanceUsd": None,
                } for row in payload["result"] if isinstance(row, dict)]
            except (httpx.HTTPError, ValueError, BlockberryError) as exc:
                last = exc
                self._metrics.append((False, (time.perf_counter() - started) * 1000, "fallback"))
        raise BlockberryError("unavailable", f"all Sui fallbacks failed ({type(last).__name__})")

    async def preflight(self) -> dict[str, Any]:
        page = await self.transactions(0, 1)
        # Blockberry's indexed feed can occasionally include a non-object
        # placeholder in ``content``.  Discovery itself already skips such
        # records, so preflight must do the same instead of turning one bad
        # item into a TypeError that disables the whole Sui worker.
        rows = [row for row in page["content"] if isinstance(row, dict)]
        head = 0
        for row in rows:
            try:
                head = max(head, int(row.get("checkpoint") or 0))
            except (TypeError, ValueError):
                continue
        defi_ok = False
        defi_error: str | None = None
        try:
            await self._dex_page(0, 1)
            defi_ok = True
        except BlockberryError as exc:
            defi_error = exc.kind
        return {
            "ok": bool(rows), "head": head, "provider": "Blockberry indexed API",
            "discovery_ok": bool(rows), "defi_ok": defi_ok, "defi_error": defi_error,
        }

    def take_metrics(self) -> dict[str, Any]:
        rows, self._metrics = self._metrics, []
        latencies = sorted(row[1] for row in rows)
        def pct(q: float) -> float | None:
            if not latencies:
                return None
            return latencies[min(len(latencies) - 1, round((len(latencies) - 1) * q))]
        errors: dict[str, int] = {}
        for ok, _latency, kind in rows:
            if not ok:
                errors[kind] = errors.get(kind, 0) + 1
        return {
            "requests": len(rows), "successes": sum(1 for row in rows if row[0]),
            "errors": sum(1 for row in rows if not row[0]),
            "latency_p50_ms": pct(0.50), "latency_p95_ms": pct(0.95),
            "errors_by_type": errors,
        }


SUI_SCHEMA = """
PRAGMA journal_mode=WAL;
PRAGMA synchronous=NORMAL;
CREATE TABLE IF NOT EXISTS sui_schema_meta(version INTEGER NOT NULL);
INSERT INTO sui_schema_meta(version)
SELECT 1 WHERE NOT EXISTS(SELECT 1 FROM sui_schema_meta);
UPDATE sui_schema_meta SET version=2 WHERE version<2;
CREATE TABLE IF NOT EXISTS sui_state(
  id INTEGER PRIMARY KEY CHECK(id=1), last_checkpoint INTEGER NOT NULL DEFAULT 0,
  window_start_ms INTEGER, provider_window_limited INTEGER NOT NULL DEFAULT 1,
  updated_at TEXT, note TEXT
);
INSERT OR IGNORE INTO sui_state(id,last_checkpoint,provider_window_limited) VALUES(1,0,1);
CREATE TABLE IF NOT EXISTS sui_packages(
  package_id TEXT PRIMARY KEY, lineage_id TEXT NOT NULL, name TEXT, publisher TEXT,
  version INTEGER, project_name TEXT, first_checkpoint INTEGER, last_checkpoint INTEGER,
  first_tx TEXT, last_tx TEXT, first_seen_at TEXT NOT NULL, last_seen_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS sui_discoveries(
  package_id TEXT NOT NULL, source TEXT NOT NULL, observed_checkpoint INTEGER,
  observed_tx TEXT, actor TEXT, first_seen_at TEXT NOT NULL,
  PRIMARY KEY(package_id,source)
);
CREATE TABLE IF NOT EXISTS sui_seen_transactions(
  digest TEXT PRIMARY KEY, checkpoint INTEGER NOT NULL, timestamp_ms INTEGER,
  enriched INTEGER NOT NULL DEFAULT 0, seen_at TEXT NOT NULL,
  failure_streak INTEGER NOT NULL DEFAULT 0, next_retry_at REAL NOT NULL DEFAULT 0,
  last_error TEXT
);
CREATE TABLE IF NOT EXISTS sui_package_objects(
  package_id TEXT NOT NULL, object_id TEXT NOT NULL, object_type TEXT,
  owner_kind TEXT NOT NULL, parent_object TEXT, depth INTEGER NOT NULL DEFAULT 0,
  first_checkpoint INTEGER, last_checkpoint INTEGER, active INTEGER NOT NULL DEFAULT 1,
  PRIMARY KEY(package_id,object_id)
);
CREATE TABLE IF NOT EXISTS sui_package_scans(
  id INTEGER PRIMARY KEY AUTOINCREMENT, package_id TEXT NOT NULL, scanned_at TEXT NOT NULL,
  status TEXT NOT NULL, total_usd REAL NOT NULL, object_count INTEGER NOT NULL,
  complete INTEGER NOT NULL, provider TEXT NOT NULL, note TEXT
);
CREATE TABLE IF NOT EXISTS sui_package_token_scans(
  scan_id INTEGER NOT NULL REFERENCES sui_package_scans(id) ON DELETE CASCADE,
  object_id TEXT NOT NULL, coin_type TEXT NOT NULL, symbol TEXT, decimals INTEGER,
  amount TEXT, price_usd REAL, usd_value REAL, priced INTEGER NOT NULL,
  PRIMARY KEY(scan_id,object_id,coin_type)
);
CREATE TABLE IF NOT EXISTS sui_balance_cache(
  owner TEXT PRIMARY KEY, checked_at TEXT NOT NULL, status TEXT NOT NULL,
  total_usd REAL NOT NULL, payload_json TEXT NOT NULL, note TEXT
);
CREATE INDEX IF NOT EXISTS idx_sui_seen_enriched ON sui_seen_transactions(enriched,checkpoint);
CREATE INDEX IF NOT EXISTS idx_sui_scans_latest ON sui_package_scans(package_id,scanned_at);
CREATE TABLE IF NOT EXISTS sui_defi_projects(
  project_key TEXT PRIMARY KEY, project_name TEXT NOT NULL, indexed_tvl REAL,
  status TEXT NOT NULL, packages_json TEXT NOT NULL, synced_at TEXT NOT NULL,
  provider_complete INTEGER NOT NULL DEFAULT 1, note TEXT,
  raw_tvl TEXT, valuation_status TEXT NOT NULL DEFAULT 'included'
);
CREATE TABLE IF NOT EXISTS sui_tvl_policies(
  project_key TEXT PRIMARY KEY, policy TEXT NOT NULL, reason TEXT NOT NULL,
  updated_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS sui_project_packages(
  project_key TEXT NOT NULL, package_id TEXT NOT NULL, first_seen_at TEXT NOT NULL,
  PRIMARY KEY(project_key,package_id)
);
CREATE TABLE IF NOT EXISTS sui_dex_pools(
  pool_id TEXT PRIMARY KEY, project_key TEXT, project_name TEXT,
  liquidity_usd REAL, payload_json TEXT NOT NULL, indexed_at TEXT NOT NULL,
  verified_at TEXT, object_exists INTEGER, object_type TEXT, owner_kind TEXT,
  verification_note TEXT
);
CREATE INDEX IF NOT EXISTS idx_sui_pools_project ON sui_dex_pools(project_key,liquidity_usd);
CREATE TABLE IF NOT EXISTS sui_sync_state(
  key TEXT PRIMARY KEY, value TEXT NOT NULL, updated_at TEXT NOT NULL
);
"""


class SuiStore:
    def __init__(self, path: Path, *, read_only: bool = False):
        self.path = path
        self._lock = threading.Lock()
        if read_only:
            self.conn = sqlite3.connect(
                f"file:{path.as_posix()}?mode=ro", uri=True,
                check_same_thread=False, timeout=60,
            )
        else:
            self.conn = sqlite3.connect(path, check_same_thread=False, timeout=60)
        self.conn.row_factory = sqlite3.Row
        if read_only:
            self.conn.execute("PRAGMA query_only=ON")
        else:
            with self.conn:
                self.conn.executescript(SUI_SCHEMA)
                columns = {row["name"] for row in self.conn.execute(
                    "PRAGMA table_info(sui_defi_projects)"
                )}
                if "raw_tvl" not in columns:
                    self.conn.execute("ALTER TABLE sui_defi_projects ADD COLUMN raw_tvl TEXT")
                if "valuation_status" not in columns:
                    self.conn.execute(
                        "ALTER TABLE sui_defi_projects ADD COLUMN valuation_status TEXT "
                        "NOT NULL DEFAULT 'included'"
                    )
                self.conn.execute(
                    "UPDATE sui_defi_projects SET raw_tvl=CAST(indexed_tvl AS TEXT) "
                    "WHERE raw_tvl IS NULL AND indexed_tvl IS NOT NULL"
                )
                seen_columns = {row["name"] for row in self.conn.execute(
                    "PRAGMA table_info(sui_seen_transactions)"
                )}
                if "failure_streak" not in seen_columns:
                    self.conn.execute(
                        "ALTER TABLE sui_seen_transactions ADD COLUMN failure_streak INTEGER NOT NULL DEFAULT 0"
                    )
                if "next_retry_at" not in seen_columns:
                    self.conn.execute(
                        "ALTER TABLE sui_seen_transactions ADD COLUMN next_retry_at REAL NOT NULL DEFAULT 0"
                    )
                if "last_error" not in seen_columns:
                    self.conn.execute("ALTER TABLE sui_seen_transactions ADD COLUMN last_error TEXT")
                self.conn.execute("UPDATE sui_schema_meta SET version=3 WHERE version<3")

    def close(self) -> None:
        self.conn.close()

    def _bump_revision_locked(self) -> None:
        try:
            row = self.conn.execute(
                "SELECT CAST(value AS INTEGER) FROM runtime_meta WHERE key='data_revision'"
            ).fetchone()
            revision = int(row[0] if row else 0) + 1
            self.conn.execute(
                """INSERT INTO runtime_meta(key,value,updated_at) VALUES('data_revision',?,?)
                   ON CONFLICT(key) DO UPDATE SET value=excluded.value,updated_at=excluded.updated_at""",
                (str(revision), utc_now()),
            )
        except sqlite3.OperationalError:
            pass

    def mark_dirty(self) -> None:
        with self._lock, self.conn:
            self._bump_revision_locked()

    def seed_tvl_policies(self, policies: list[dict[str, Any]]) -> None:
        now = utc_now()
        with self._lock, self.conn:
            for item in policies:
                policy = str(item.get("policy") or "")
                reason = str(item.get("reason") or "").strip()
                key = str(item.get("project_key") or "").strip().lower()
                if not key or not reason or policy not in {"include_verified", "quarantine"}:
                    raise ValueError("Sui TVL policy requires project_key, policy and reason")
                self.conn.execute(
                    """INSERT INTO sui_tvl_policies(project_key,policy,reason,updated_at)
                       VALUES(?,?,?,?) ON CONFLICT(project_key) DO UPDATE SET
                         policy=excluded.policy,reason=excluded.reason,
                         updated_at=excluded.updated_at""",
                    (key, policy, reason, now),
                )

    def seen(self, digest: str) -> bool:
        with self._lock:
            return self.conn.execute(
                "SELECT 1 FROM sui_seen_transactions WHERE digest=?", (digest,)
            ).fetchone() is not None

    def add_seen(self, digest: str, checkpoint: int, timestamp_ms: int | None) -> None:
        with self._lock, self.conn:
            self.conn.execute(
                "INSERT OR IGNORE INTO sui_seen_transactions(digest,checkpoint,timestamp_ms,seen_at) VALUES(?,?,?,?)",
                (digest, checkpoint, timestamp_ms, utc_now()),
            )

    def upsert_package(
        self, package_id: str, source: str, checkpoint: int | None, digest: str | None,
        actor: str | None = None, name: str | None = None, publisher: str | None = None,
        version: int | None = None, project_name: str | None = None,
        lineage_id: str | None = None,
    ) -> None:
        package_id = normalize_sui_id(package_id)
        if package_id in SYSTEM_PACKAGES:
            return
        now = utc_now()
        lineage_id = normalize_sui_id(lineage_id or package_id)
        with self._lock, self.conn:
            self.conn.execute(
                """INSERT INTO sui_packages(package_id,lineage_id,name,publisher,version,project_name,
                       first_checkpoint,last_checkpoint,first_tx,last_tx,first_seen_at,last_seen_at)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
                   ON CONFLICT(package_id) DO UPDATE SET
                     lineage_id=COALESCE(excluded.lineage_id,sui_packages.lineage_id),
                     name=COALESCE(excluded.name,sui_packages.name),
                     publisher=COALESCE(excluded.publisher,sui_packages.publisher),
                     version=COALESCE(excluded.version,sui_packages.version),
                     project_name=COALESCE(excluded.project_name,sui_packages.project_name),
                     first_checkpoint=COALESCE(sui_packages.first_checkpoint,excluded.first_checkpoint),
                     first_tx=COALESCE(sui_packages.first_tx,excluded.first_tx),
                     last_checkpoint=MAX(COALESCE(sui_packages.last_checkpoint,0),COALESCE(excluded.last_checkpoint,0)),
                     last_tx=COALESCE(excluded.last_tx,sui_packages.last_tx), last_seen_at=excluded.last_seen_at""",
                (package_id, lineage_id, name, publisher, version, project_name,
                 checkpoint, checkpoint, digest, digest, now, now),
            )
            self.conn.execute(
                """INSERT INTO sui_discoveries(package_id,source,observed_checkpoint,
                       observed_tx,actor,first_seen_at) VALUES(?,?,?,?,?,?)
                   ON CONFLICT(package_id,source) DO UPDATE SET
                     observed_checkpoint=COALESCE(sui_discoveries.observed_checkpoint,excluded.observed_checkpoint),
                     observed_tx=COALESCE(sui_discoveries.observed_tx,excluded.observed_tx),
                     actor=COALESCE(sui_discoveries.actor,excluded.actor)""",
                (package_id, source, checkpoint, digest, actor, now),
            )

    def add_objects(self, rows: Iterable[dict[str, Any]], checkpoint: int) -> None:
        values = [(
            row["package_id"], row["object_id"], row.get("object_type"), row["owner_kind"],
            row.get("parent_object"), 0, checkpoint, checkpoint,
        ) for row in rows]
        if not values:
            return
        with self._lock, self.conn:
            self.conn.executemany(
                """INSERT INTO sui_package_objects(package_id,object_id,object_type,owner_kind,
                       parent_object,depth,first_checkpoint,last_checkpoint)
                   VALUES(?,?,?,?,?,?,?,?)
                   ON CONFLICT(package_id,object_id) DO UPDATE SET
                     object_type=excluded.object_type, owner_kind=excluded.owner_kind,
                     parent_object=excluded.parent_object,
                     last_checkpoint=MAX(COALESCE(sui_package_objects.last_checkpoint,0),excluded.last_checkpoint),
                     active=1""", values,
            )

    def pending_enrichment(self, limit: int) -> list[sqlite3.Row]:
        with self._lock:
            return list(self.conn.execute(
                "SELECT * FROM sui_seen_transactions WHERE enriched=0 AND next_retry_at<=? "
                "ORDER BY checkpoint ASC,digest LIMIT ?",
                (time.time(), limit),
            ))

    def defer_enrichment(self, digest: str, error: str) -> None:
        with self._lock, self.conn:
            row = self.conn.execute(
                "SELECT failure_streak FROM sui_seen_transactions WHERE digest=?", (digest,),
            ).fetchone()
            if row is None:
                return
            streak = int(row[0]) + 1
            delay = min(3600, 30 * 2 ** min(streak - 1, 7))
            self.conn.execute(
                "UPDATE sui_seen_transactions SET failure_streak=?,next_retry_at=?,"
                "last_error=? WHERE digest=?",
                (streak, time.time() + delay, error[:120], digest),
            )

    def finish_enrichment(self, digest: str) -> None:
        with self._lock, self.conn:
            self.conn.execute("UPDATE sui_seen_transactions SET enriched=1 WHERE digest=?", (digest,))

    def update_state(self, checkpoint: int, cutoff_ms: int, limited: bool, note: str) -> None:
        with self._lock, self.conn:
            self.conn.execute(
                "UPDATE sui_state SET last_checkpoint=MAX(last_checkpoint,?),window_start_ms=?,provider_window_limited=?,updated_at=?,note=? WHERE id=1",
                (checkpoint, cutoff_ms, int(limited), utc_now(), note),
            )

    def state(self) -> sqlite3.Row:
        with self._lock:
            return self.conn.execute("SELECT * FROM sui_state WHERE id=1").fetchone()

    def pending_packages(self, recheck_sec: int, limit: int = 100) -> list[sqlite3.Row]:
        cutoff = (datetime.now(UTC) - timedelta(seconds=recheck_sec)).isoformat()
        with self._lock:
            return list(self.conn.execute(
                """SELECT p.* FROM sui_packages p LEFT JOIN sui_package_scans s
                     ON s.id=(SELECT s2.id FROM sui_package_scans s2 WHERE s2.package_id=p.package_id
                              ORDER BY s2.scanned_at DESC,s2.id DESC LIMIT 1)
                   WHERE s.scanned_at IS NULL OR s.scanned_at<?
                   ORDER BY COALESCE(s.scanned_at,''),p.last_checkpoint DESC LIMIT ?""",
                (cutoff, limit),
            ))

    def owners_for(self, package_id: str) -> list[str]:
        with self._lock:
            rows = self.conn.execute(
                "SELECT object_id FROM sui_package_objects WHERE package_id=? AND active=1 ORDER BY object_id",
                (package_id,),
            ).fetchall()
        return list(dict.fromkeys([package_id, *(row["object_id"] for row in rows)]))

    def save_scan(self, package_id: str, status: str, total_usd: float,
                  items: list[dict[str, Any]], complete: bool, note: str) -> int:
        with self._lock, self.conn:
            cur = self.conn.execute(
                """INSERT INTO sui_package_scans(package_id,scanned_at,status,total_usd,
                     object_count,complete,provider,note) VALUES(?,?,?,?,?,?,?,?)""",
                (package_id, utc_now(), status, total_usd,
                 len({row["object_id"] for row in items}), int(complete),
                 "Blockberry indexed API", note),
            )
            scan_id = int(cur.lastrowid)
            self.conn.executemany(
                """INSERT INTO sui_package_token_scans(scan_id,object_id,coin_type,symbol,
                     decimals,amount,price_usd,usd_value,priced) VALUES(?,?,?,?,?,?,?,?,?)""",
                [(scan_id, row["object_id"], row["coin_type"], row.get("symbol"),
                  row.get("decimals"), str(row.get("amount") or "0"), row.get("price_usd"),
                  row.get("usd_value"), int(row["priced"])) for row in items],
            )
            return scan_id

    def cache_balance(self, owner: str, rows: list[dict[str, Any]], total: float) -> None:
        with self._lock, self.conn:
            self.conn.execute(
                """INSERT INTO sui_balance_cache(owner,checked_at,status,total_usd,payload_json,note)
                   VALUES(?,?,'complete',?,?,NULL)
                   ON CONFLICT(owner) DO UPDATE SET checked_at=excluded.checked_at,
                     status=excluded.status,total_usd=excluded.total_usd,payload_json=excluded.payload_json,note=NULL""",
                (owner, utc_now(), total, json.dumps(rows, ensure_ascii=False)),
            )

    def cached_balance(self, owner: str, max_age_sec: int) -> list[dict[str, Any]] | None:
        cutoff = (datetime.now(UTC) - timedelta(seconds=max_age_sec)).isoformat()
        with self._lock:
            row = self.conn.execute(
                "SELECT payload_json FROM sui_balance_cache WHERE owner=? AND checked_at>=? AND status='complete'",
                (owner, cutoff),
            ).fetchone()
        return json.loads(row["payload_json"]) if row else None

    def latest_scans(self) -> list[sqlite3.Row]:
        with self._lock:
            return list(self.conn.execute(
                """SELECT s.*,p.lineage_id,p.name,p.first_checkpoint,p.last_checkpoint,p.first_seen_at,p.last_seen_at
                   FROM sui_package_scans s JOIN sui_packages p ON p.package_id=s.package_id
                   WHERE s.id=(SELECT s2.id FROM sui_package_scans s2 WHERE s2.package_id=s.package_id
                               ORDER BY s2.scanned_at DESC,s2.id DESC LIMIT 1)
                   ORDER BY s.total_usd DESC,s.package_id"""
            ))

    def scan_tokens(self, scan_id: int) -> list[sqlite3.Row]:
        with self._lock:
            return list(self.conn.execute(
                "SELECT * FROM sui_package_token_scans WHERE scan_id=? AND (usd_value>0 OR CAST(amount AS REAL)>0) ORDER BY COALESCE(usd_value,-1) DESC",
                (scan_id,),
            ))

    def sources(self, package_id: str) -> list[sqlite3.Row]:
        with self._lock:
            return list(self.conn.execute(
                "SELECT * FROM sui_discoveries WHERE package_id=? ORDER BY first_seen_at,source",
                (package_id,),
            ))

    def sync_defi_projects(
        self, rows: list[dict[str, Any]], min_usd: float,
        tvl_anomaly_usd: float = 10_000_000_000.0,
    ) -> int:
        now = utc_now()
        seen: set[str] = set()
        with self._lock, self.conn:
            for raw in rows:
                name = str(
                    raw.get("projectName") or raw.get("name") or raw.get("dexName") or ""
                ).strip()
                if not name:
                    continue
                key = str(raw.get("id") or raw.get("projectId") or name).strip().lower()
                tvl_raw = raw.get("currTvl", raw.get("currentTvl", raw.get("tvl")))
                try:
                    tvl = float(tvl_raw) if tvl_raw is not None else None
                except (TypeError, ValueError):
                    tvl = None
                invalid_tvl = tvl is not None and not math.isfinite(tvl)
                if invalid_tvl:
                    tvl = None
                packages: list[str] = []
                for item in raw.get("packages") or raw.get("projectPackages") or []:
                    candidate = (
                        item.get("packageAddress") or item.get("packageId") or item.get("address")
                        if isinstance(item, dict) else item
                    )
                    try:
                        packages.append(normalize_sui_id(candidate))
                    except (TypeError, ValueError):
                        continue
                packages = sorted(set(packages))
                policy = self.conn.execute(
                    "SELECT policy,reason FROM sui_tvl_policies WHERE project_key=?", (key,),
                ).fetchone()
                quarantine = (
                    policy is not None and policy["policy"] == "quarantine"
                ) or (
                    (invalid_tvl or (tvl is not None and tvl >= tvl_anomaly_usd))
                    and not (policy is not None and policy["policy"] == "include_verified" and not invalid_tvl)
                )
                status = (
                    "incomplete" if tvl is None or not packages or quarantine
                    else "qualifying" if tvl >= min_usd else "below"
                )
                note = (
                    f"indexed TVL quarantined (invalid, >= ${tvl_anomaly_usd:,.0f}, or policy)"
                    if quarantine else None if packages else "Blockberry project has no package mapping"
                )
                valuation_status = "anomalous_balance" if quarantine else "included"
                self.conn.execute(
                    """INSERT INTO sui_defi_projects(
                         project_key,project_name,indexed_tvl,status,packages_json,synced_at,
                         provider_complete,note,raw_tvl,valuation_status
                       ) VALUES(?,?,?,?,?,?,?,?,?,?)
                       ON CONFLICT(project_key) DO UPDATE SET
                         project_name=excluded.project_name,indexed_tvl=excluded.indexed_tvl,
                         status=excluded.status,packages_json=excluded.packages_json,
                         synced_at=excluded.synced_at,provider_complete=excluded.provider_complete,
                         note=excluded.note,raw_tvl=excluded.raw_tvl,
                         valuation_status=excluded.valuation_status""",
                    (key, name, tvl, status, json.dumps(packages), now, 1, note,
                     str(tvl_raw) if tvl_raw is not None else None, valuation_status),
                )
                self.conn.execute(
                    "DELETE FROM sui_project_packages WHERE project_key=?", (key,)
                )
                self.conn.executemany(
                    "INSERT INTO sui_project_packages(project_key,package_id,first_seen_at) VALUES(?,?,?)",
                    [(key, package, now) for package in packages],
                )
                seen.add(key)
            if seen:
                placeholders = ",".join("?" for _ in seen)
                self.conn.execute(
                    f"UPDATE sui_defi_projects SET provider_complete=0,note='missing from latest provider snapshot' WHERE project_key NOT IN ({placeholders})",
                    tuple(seen),
                )
            if seen:
                self._bump_revision_locked()
        return len(seen)

    def save_pools(self, rows: list[dict[str, Any]]) -> int:
        now = utc_now()
        with self._lock:
            project_by_name = {
                str(row["project_name"]).strip().lower(): str(row["project_key"])
                for row in self.conn.execute(
                    "SELECT project_key,project_name FROM sui_defi_projects"
                )
            }
            known_keys = set(project_by_name.values())
        values = []
        for raw in rows:
            candidate = raw.get("poolId") or raw.get("poolAddress") or raw.get("objectId")
            try:
                pool_id = normalize_sui_id(candidate)
            except (TypeError, ValueError):
                continue
            project_name = str(raw.get("projectName") or raw.get("dexName") or "").strip()
            candidate_key = str(raw.get("projectId") or "").strip().lower()
            project_key = (
                candidate_key if candidate_key in known_keys
                else project_by_name.get(project_name.lower())
                or candidate_key or project_name.lower() or None
            )
            try:
                liquidity = float(raw.get("liquidityInUsd") or raw.get("liquidityUsd") or 0)
            except (TypeError, ValueError):
                liquidity = 0.0
            values.append((
                pool_id, project_key, project_name or None, liquidity,
                json.dumps(raw, ensure_ascii=False), now,
            ))
        with self._lock, self.conn:
            self.conn.executemany(
                """INSERT INTO sui_dex_pools(
                     pool_id,project_key,project_name,liquidity_usd,payload_json,indexed_at
                   ) VALUES(?,?,?,?,?,?) ON CONFLICT(pool_id) DO UPDATE SET
                     project_key=excluded.project_key,project_name=excluded.project_name,
                     liquidity_usd=excluded.liquidity_usd,payload_json=excluded.payload_json,
                     indexed_at=excluded.indexed_at""",
                values,
            )
        return len(values)

    def pools_due_verification(self, max_age_sec: int, limit: int) -> list[sqlite3.Row]:
        cutoff = (datetime.now(UTC) - timedelta(seconds=max_age_sec)).isoformat()
        with self._lock:
            return list(self.conn.execute(
                """SELECT * FROM sui_dex_pools
                   WHERE verified_at IS NULL OR verified_at<?
                   ORDER BY liquidity_usd DESC LIMIT ?""",
                (cutoff, limit),
            ))

    def save_pool_verification(self, pool_id: str, result: dict[str, Any] | None) -> None:
        exists = result is not None
        data = result.get("data") if isinstance(result, dict) else None
        owner = data.get("owner") if isinstance(data, dict) else None
        owner_kind, _parent = _owner_kind(owner)
        with self._lock, self.conn:
            self.conn.execute(
                """UPDATE sui_dex_pools SET verified_at=?,object_exists=?,object_type=?,
                     owner_kind=?,verification_note=? WHERE pool_id=?""",
                (utc_now(), int(exists), data.get("type") if isinstance(data, dict) else None,
                 owner_kind, None if exists else "Sui RPC object lookup failed", pool_id),
            )

    def latest_projects(self) -> list[sqlite3.Row]:
        with self._lock:
            return list(self.conn.execute(
                """SELECT p.*,
                     COALESCE((SELECT SUM(liquidity_usd) FROM sui_dex_pools d
                               WHERE d.project_key=p.project_key AND d.object_exists=1),0) AS verified_pool_tvl,
                     (SELECT COUNT(*) FROM sui_dex_pools d WHERE d.project_key=p.project_key) AS pool_count,
                     (SELECT COUNT(*) FROM sui_dex_pools d WHERE d.project_key=p.project_key
                                                       AND d.object_exists=1) AS verified_pool_count
                   FROM sui_defi_projects p ORDER BY COALESCE(p.indexed_tvl,-1) DESC,p.project_name"""
            ))

    def sync_value(self, key: str, default: str = "") -> str:
        with self._lock:
            row = self.conn.execute(
                "SELECT value FROM sui_sync_state WHERE key=?", (key,)
            ).fetchone()
        return str(row["value"]) if row else default

    def set_sync_value(self, key: str, value: Any) -> None:
        with self._lock, self.conn:
            self.conn.execute(
                """INSERT INTO sui_sync_state(key,value,updated_at) VALUES(?,?,?)
                   ON CONFLICT(key) DO UPDATE SET value=excluded.value,updated_at=excluded.updated_at""",
                (key, str(value), utc_now()),
            )

    def mark_defi_sync_failed(self, note: str) -> None:
        """Record the outage without mutating the last valid TVL snapshot."""
        safe_note = str(note or "Blockberry DeFi sync failed")[:500]
        with self._lock, self.conn:
            self.conn.execute(
                """INSERT INTO sui_sync_state(key,value,updated_at) VALUES('last_defi_error',?,?)
                   ON CONFLICT(key) DO UPDATE SET value=excluded.value,updated_at=excluded.updated_at""",
                (safe_note, utc_now()),
            )

    def revalue_suspicious_projects(self, threshold_usd: float) -> int:
        """Idempotently quarantine existing huge TVL, preserving indexed_tvl."""
        with self._lock, self.conn:
            rows = self.conn.execute(
                """SELECT p.project_key FROM sui_defi_projects p
                   LEFT JOIN sui_tvl_policies v ON v.project_key=p.project_key
                   WHERE p.valuation_status!='anomalous_balance'
                     AND (v.policy='quarantine' OR
                          (p.indexed_tvl>=? AND COALESCE(v.policy,'')!='include_verified'))""",
                (threshold_usd,),
            ).fetchall()
            for row in rows:
                self.conn.execute(
                    "UPDATE sui_defi_projects SET status='incomplete',"
                    "valuation_status='anomalous_balance',note=? WHERE project_key=?",
                    (f"indexed TVL quarantined (>= ${threshold_usd:,.0f} or policy)",
                     row["project_key"]),
                )
            if rows:
                self._bump_revision_locked()
        return len(rows)

    def pool_summary(self) -> dict[str, int]:
        with self._lock:
            row = self.conn.execute(
                """SELECT COUNT(*) AS total,
                          SUM(CASE WHEN object_exists=1 THEN 1 ELSE 0 END) AS verified
                   FROM sui_dex_pools"""
            ).fetchone()
        return {"total": int(row["total"] or 0), "verified": int(row["verified"] or 0)}

    def technical_packages(self) -> list[sqlite3.Row]:
        with self._lock:
            return list(self.conn.execute(
                "SELECT * FROM sui_packages ORDER BY last_checkpoint DESC,package_id"
            ))

    def summary(self) -> dict[str, int]:
        with self._lock:
            packages = self.conn.execute("SELECT COUNT(*) FROM sui_packages").fetchone()[0]
            objects = self.conn.execute("SELECT COUNT(*) FROM sui_package_objects WHERE active=1").fetchone()[0]
            pending = self.conn.execute("SELECT COUNT(*) FROM sui_seen_transactions WHERE enriched=0").fetchone()[0]
            statuses = {row["status"]: row["n"] for row in self.conn.execute(
                "SELECT status,COUNT(*) n FROM sui_defi_projects GROUP BY status")}
        return {"packages": int(packages), "objects": int(objects), "enrichment_pending": int(pending),
                "qualifying": int(statuses.get("qualifying", 0)),
                "below": int(statuses.get("below", 0)),
                "incomplete": int(statuses.get("incomplete", 0))}


def _package_from_metadata(row: dict[str, Any]) -> tuple[str, dict[str, Any]] | None:
    candidate = row.get("id") or row.get("packageId")
    if not candidate:
        return None
    try:
        package_id = normalize_sui_id(candidate)
    except ValueError:
        return None
    if package_id in SYSTEM_PACKAGES:
        return None
    return package_id, row


async def discover_sui_once(
    store: SuiStore, client: BlockberryClient, cfg: SuiConfig,
    from_checkpoint: int | None = None, to_checkpoint: int | None = None,
    *, enrich: bool = True,
) -> dict[str, Any]:
    packages_before = int(store.conn.execute("SELECT COUNT(*) FROM sui_packages").fetchone()[0])
    cutoff_ms = int((datetime.now(UTC) - timedelta(days=cfg.discovery_days)).timestamp() * 1000)
    maximum_checkpoint = 0
    oldest_ms: int | None = None
    stopped_at_cutoff = False
    reached_saved_cursor = False
    new_txs = 0
    for page_number in range(cfg.max_pages):
        page = await client.transactions(page_number, cfg.page_size)
        rows = page["content"]
        if not rows:
            stopped_at_cutoff = True
            break
        for tx in rows:
            if not isinstance(tx, dict):
                continue
            if str(tx.get("txStatus") or "").upper() != "SUCCESS":
                continue
            checkpoint = int(tx.get("checkpoint") or 0)
            timestamp_ms = int(tx.get("timestamp") or 0) or None
            digest = str(tx.get("txHash") or "")
            maximum_checkpoint = max(maximum_checkpoint, checkpoint)
            if timestamp_ms:
                oldest_ms = timestamp_ms if oldest_ms is None else min(oldest_ms, timestamp_ms)
                if timestamp_ms < cutoff_ms:
                    stopped_at_cutoff = True
                    continue
            if from_checkpoint is not None and checkpoint < from_checkpoint:
                continue
            if to_checkpoint is not None and checkpoint > to_checkpoint:
                continue
            if digest and store.seen(digest):
                reached_saved_cursor = True
                continue
            if not digest:
                continue
            actor = tx.get("senderAddress")
            for meta in tx.get("packagesMetadata") or []:
                parsed = _package_from_metadata(meta) if isinstance(meta, dict) else None
                if not parsed:
                    continue
                package_id, details = parsed
                store.upsert_package(
                    package_id, "active_call", checkpoint, digest, actor,
                    name=details.get("name"), project_name=details.get("projectName"),
                )
            store.add_seen(digest, checkpoint, timestamp_ms)
            new_txs += 1
        if stopped_at_cutoff or reached_saved_cursor:
            break

    # The packages catalog is useful for Publish discovery and lineage metadata.
    catalog_complete = False
    for page_number in range(cfg.max_pages):
        page = await client.packages(page_number, cfg.page_size)
        rows = page["content"]
        if not rows:
            catalog_complete = True
            break
        page_older_than_cutoff = False
        for row in rows:
            parsed = _package_from_metadata(row)
            if not parsed:
                continue
            package_id, details = parsed
            created = int(details.get("createTimestamp") or 0) or None
            if created and created < cutoff_ms:
                page_older_than_cutoff = True
                continue
            project_packages = details.get("projectPackages") or []
            lineage = package_id
            if isinstance(project_packages, list) and project_packages:
                ids = []
                for candidate in project_packages:
                    try:
                        ids.append(normalize_sui_id(
                            candidate.get("packageId") if isinstance(candidate, dict) else candidate
                        ))
                    except (ValueError, TypeError):
                        pass
                if ids:
                    lineage = min(ids)
            store.upsert_package(
                package_id, "publish", None, None,
                actor=details.get("publisherAddress"), name=details.get("packageName"),
                publisher=details.get("publisherAddress"), version=details.get("version"),
                project_name=details.get("projectName"), lineage_id=lineage,
            )
        if page_older_than_cutoff:
            catalog_complete = True
            break

    enriched = await enrich_sui_once(store, client, cfg) if enrich else 0

    limited = not (stopped_at_cutoff or reached_saved_cursor)
    note = (
        f"coverage gap: saved checkpoint/cutoff not reached in {cfg.max_pages} pages"
        if limited else "saved checkpoint or 30-day cutoff reached"
    )
    if not catalog_complete:
        note += "; package catalog coverage incomplete"
    store.update_state(maximum_checkpoint, cutoff_ms, limited, note)
    packages_after = int(store.conn.execute("SELECT COUNT(*) FROM sui_packages").fetchone()[0])
    if new_txs or enriched or packages_after != packages_before:
        store.mark_dirty()
    return {"head": maximum_checkpoint, "new_transactions": new_txs,
            "enriched": enriched, "window_limited": limited, "oldest_ms": oldest_ms}


async def enrich_sui_once(store: SuiStore, client: BlockberryClient, cfg: SuiConfig) -> int:
    enriched = 0
    for row in store.pending_enrichment(cfg.raw_enrich_per_pass):
        try:
            raw = await client.raw_transaction(row["digest"])
            parsed = parse_raw_transaction(raw)
            result = raw["result"]
            actor = (((result.get("transaction") or {}).get("data") or {}).get("sender"))
            for package_id in parsed["calls"]:
                store.upsert_package(package_id, "active_call", row["checkpoint"], row["digest"], actor)
            for package_id in parsed["published"]:
                store.upsert_package(package_id, "publish", row["checkpoint"], row["digest"], actor)
            store.add_objects(parsed["objects"], int(row["checkpoint"]))
            store.finish_enrichment(row["digest"])
            enriched += 1
        except (BlockberryError, ValueError, TypeError, KeyError) as exc:
            store.defer_enrichment(row["digest"], type(exc).__name__)
            log.warning("[sui/enrichment] deferred %s: %s", row["digest"][:12], type(exc).__name__)
    if enriched:
        store.mark_dirty()
    return enriched


async def sui_enrichment_loop(
    store: SuiStore, client: BlockberryClient, cfg: SuiConfig, stop: asyncio.Event,
) -> None:
    while not stop.is_set():
        enriched = await enrich_sui_once(store, client, cfg)
        if enriched:
            log.info("[sui/enrichment] enriched=%s pending=%s", enriched,
                     store.summary()["enrichment_pending"])
        try:
            await asyncio.wait_for(stop.wait(), timeout=1 if enriched else cfg.poll_sec)
        except asyncio.TimeoutError:
            pass


def _balance_item(owner: str, row: dict[str, Any]) -> dict[str, Any] | None:
    coin_type = row.get("coinType") or row.get("coin_type")
    if not coin_type:
        return None
    try:
        coin_type = normalize_coin_type(coin_type)
    except ValueError:
        coin_type = str(coin_type)
    amount = row.get("balance")
    try:
        positive = float(amount or 0) > 0
    except (TypeError, ValueError, OverflowError):
        positive = False
    if not positive:
        return None
    usd = row.get("balanceUsd")
    price = row.get("coinPrice")
    try:
        usd_value = float(usd) if usd is not None else None
    except (TypeError, ValueError):
        usd_value = None
    try:
        price_usd = float(price) if price is not None else None
    except (TypeError, ValueError):
        price_usd = None
    return {
        "object_id": owner, "coin_type": coin_type,
        "symbol": row.get("coinSymbol") or row.get("symbol"),
        "decimals": row.get("decimals"), "amount": str(amount),
        "price_usd": price_usd, "usd_value": usd_value,
        "priced": usd_value is not None,
    }


async def scan_sui_package(
    store: SuiStore, client: BlockberryClient, package_id: str, min_usd: float,
) -> dict[str, Any]:
    # A Move package/shared object is not an account. Keep this compatibility
    # entry as technical history without manufacturing a wallet balance.
    note = (
        "legacy package scan; package/shared-object account balances are intentionally disabled; "
        "financial valuation is reported once per Blockberry DeFi project"
    )
    store.save_scan(package_id, "incomplete", 0.0, [], False, note)
    return {"package_id": package_id, "status": "incomplete", "total_usd": 0.0}


async def balance_sui_once(
    store: SuiStore, client: BlockberryClient, cfg: SuiConfig, min_usd: float,
    limit: int = 100, stop: asyncio.Event | None = None,
) -> list[dict[str, Any]]:
    projects = await client.defi_projects()
    store.sync_defi_projects(projects, min_usd, cfg.tvl_anomaly_usd)
    store.set_sync_value("last_defi_sync", time.time())
    store.set_sync_value("last_defi_error", "")

    last_pool_sync = float(store.sync_value("last_pool_sync", "0") or 0)
    if time.time() - last_pool_sync >= cfg.pool_verify_sec:
        collected: list[dict[str, Any]] = []
        stop_paging = False
        for page in range(max(1, (cfg.pool_limit + cfg.page_size - 1) // cfg.page_size)):
            if stop is not None and stop.is_set():
                break
            rows = await client.dex_pools(page, cfg.page_size)
            if not rows:
                break
            for row in rows:
                try:
                    liquidity = float(
                        row.get("liquidityInUsd") or row.get("liquidityUsd") or 0
                    )
                except (TypeError, ValueError):
                    liquidity = 0.0
                if liquidity < cfg.pool_min_liquidity_usd:
                    stop_paging = True
                    break
                collected.append(row)
                if len(collected) >= cfg.pool_limit:
                    stop_paging = True
                    break
            if stop_paging:
                break
        store.save_pools(collected)
        for row in store.pools_due_verification(cfg.pool_verify_sec, cfg.pool_limit):
            if stop is not None and stop.is_set():
                break
            result = await client.object_lookup(row["pool_id"])
            store.save_pool_verification(row["pool_id"], result)
        if stop is None or not stop.is_set():
            store.set_sync_value("last_pool_sync", time.time())
    latest = [dict(row) for row in store.latest_projects()]
    summary = store.summary()
    pools = store.pool_summary()
    log.info(
        "[sui/balances] projects=%s qualifying=%s below=%s incomplete=%s "
        "pools=%s verified=%s snapshot=%s",
        len(latest), summary["qualifying"], summary["below"], summary["incomplete"],
        pools["total"], pools["verified"], utc_now(),
    )
    return latest


async def sui_discovery_loop(
    store: SuiStore, client: BlockberryClient, cfg: SuiConfig, stop: asyncio.Event,
    *, once: bool = False, from_checkpoint: int | None = None,
    to_checkpoint: int | None = None, monitor: Any = None, run_id: int | None = None,
    discovery_allowed: Any = None,
) -> None:
    while not stop.is_set():
        if discovery_allowed is not None and not discovery_allowed():
            try:
                await asyncio.wait_for(stop.wait(), timeout=5)
            except asyncio.TimeoutError:
                pass
            continue
        result = await discover_sui_once(
            store, client, cfg, from_checkpoint, to_checkpoint,
            enrich=discovery_allowed is None,
        )
        summary = store.summary()
        metrics = client.take_metrics()
        if monitor is not None:
            monitor.add_chain_sample(
                run_id=run_id, chain="sui", role="discovery+balance",
                cursor=result["head"], safe_head=result["head"], lag=0,
                blocks_per_hour=None, contracts=summary["packages"],
                direct_deploy=sum(1 for _ in store.conn.execute("SELECT 1 FROM sui_discoveries WHERE source='publish'")),
                active_call=sum(1 for _ in store.conn.execute("SELECT 1 FROM sui_discoveries WHERE source='active_call'")),
                active_rpc="Blockberry indexed API", cooldown_sec=0,
                rpc_requests=metrics["requests"], rpc_successes=metrics["successes"],
                rpc_errors=metrics["errors"], latency_p50_ms=metrics["latency_p50_ms"],
                latency_p95_ms=metrics["latency_p95_ms"], errors_json=metrics["errors_by_type"],
            )
        log.info("[sui] head=%s new_tx=%s enriched=%s packages=%s limited=%s",
                 result["head"], result["new_transactions"], result["enriched"],
                 summary["packages"], result["window_limited"])
        if once or to_checkpoint is not None:
            return
        try:
            await asyncio.wait_for(stop.wait(), timeout=cfg.poll_sec)
        except asyncio.TimeoutError:
            pass


async def sui_balance_loop(
    store: SuiStore, client: BlockberryClient, cfg: SuiConfig, min_usd: float,
    stop: asyncio.Event, *, once: bool = False,
) -> None:
    failures = 0
    while not stop.is_set():
        try:
            await balance_sui_once(store, client, cfg, min_usd, stop=stop)
            failures = 0
            delay = cfg.defi_sync_sec
        except BlockberryError as exc:
            failures += 1
            delay = min(cfg.defi_sync_sec, max(30, 30 * (2 ** min(failures - 1, 5))))
            status = f"HTTP {exc.status_code}" if exc.status_code is not None else exc.kind
            note = f"Blockberry DeFi sync failed: {status} at {exc.endpoint or 'indexed API'}"
            store.mark_defi_sync_failed(note)
            log.warning("[sui/balances] %s; retry in %ss", note, delay)
            if once:
                raise
        except Exception as exc:
            failures += 1
            delay = min(cfg.defi_sync_sec, max(30, 30 * (2 ** min(failures - 1, 5))))
            note = f"Sui balance sync failed: {type(exc).__name__}"
            store.mark_defi_sync_failed(note)
            log.warning("[sui/balances] %s; retry in %ss", note, delay)
            if once:
                raise
        else:
            if once:
                return
        try:
            await asyncio.wait_for(stop.wait(), timeout=delay)
        except asyncio.TimeoutError:
            pass


def export_sui_xlsx(store: SuiStore, export_dir: Path, min_usd: float,
                    snapshot: bool = False,
                    mode: str = "full") -> tuple[Path, Path, Path, Path]:
    export_dir.mkdir(parents=True, exist_ok=True)
    if snapshot:
        folder = export_dir / "archive"
        folder.mkdir(parents=True, exist_ok=True)
        suffix = "_" + datetime.now(UTC).strftime("%Y%m%d")
    else:
        folder, suffix = export_dir, ""
    paths = (
        folder / f"sui_qualifying{suffix}.xlsx",
        folder / f"sui_below_threshold{suffix}.xlsx",
        folder / f"sui_incomplete{suffix}.xlsx",
        folder / f"sui_packages{suffix}.xlsx",
    )
    grouped = {"qualifying": [], "below": [], "incomplete": []}
    now = datetime.now(UTC)
    for row in store.latest_projects():
        try:
            age = (now - datetime.fromisoformat(row["synced_at"])).total_seconds()
        except (TypeError, ValueError):
            age = float("inf")
        tvl = float(row["indexed_tvl"] or 0.0)
        status = (
            "incomplete" if age > 1800 or row["status"] == "incomplete" or not row["provider_complete"]
            or not json.loads(row["packages_json"] or "[]")
            else "qualifying" if tvl >= min_usd else "below"
        )
        grouped[status].append(row)
    headers = [
        "Project", "Indexed TVL USD", "Status", "Packages", "Verified pool TVL USD",
        "Verified pools", "Known pools", "Verification coverage", "Synced (UTC)",
        "Provider / notes",
    ]
    for path, status in zip(paths[:3], ("qualifying", "below", "incomplete")):
        if mode == "qualifying" and status != "qualifying":
            continue
        wb = Workbook()
        ws = wb.active
        ws.title = "Sui DeFi projects"
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for row in grouped.get(status, []):
            packages = json.loads(row["packages_json"] or "[]")
            pool_count = int(row["pool_count"] or 0)
            verified_count = int(row["verified_pool_count"] or 0)
            coverage = f"{verified_count}/{pool_count}" if pool_count else "0/0"
            ws.append([
                row["project_name"], float(row["indexed_tvl"] or 0.0), status,
                "; ".join(packages)[:32000], float(row["verified_pool_tvl"] or 0.0),
                verified_count, pool_count, coverage, row["synced_at"],
                "Blockberry indexed TVL; verified pool TVL is evidence only and is not added "
                "to indexed TVL" + (f"; {row['note']}" if row["note"] else ""),
            ])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for index, width in enumerate((30, 20, 14, 80, 24, 16, 14, 20, 24, 75), 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        for cell in ws["B"][1:]:
            cell.number_format = '#,##0.00"$"'
        for cell in ws["E"][1:]:
            cell.number_format = '#,##0.00"$"'
        temporary = path.with_suffix(".tmp.xlsx")
        wb.save(temporary)
        temporary.replace(path)

    if mode == "qualifying":
        return paths

    technical = Workbook()
    ws = technical.active
    ws.title = "Move packages (technical)"
    ws.append([
        "Package ID", "Lineage", "Name", "Project hint", "Publisher", "Version",
        "First checkpoint", "Last checkpoint", "First tx", "Last tx",
        "First seen (UTC)", "Last seen (UTC)", "Notice",
    ])
    for row in store.technical_packages():
        ws.append([
            row["package_id"], row["lineage_id"], row["name"], row["project_name"],
            row["publisher"], row["version"], row["first_checkpoint"], row["last_checkpoint"],
            row["first_tx"], row["last_tx"], row["first_seen_at"], row["last_seen_at"],
            "Technical package discovery only; no wallet balance or duplicated project TVL",
        ])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
    for index, width in enumerate((69, 69, 24, 24, 69, 10, 18, 18, 69, 69, 23, 23, 75), 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    temporary = paths[3].with_suffix(".tmp.xlsx")
    technical.save(temporary)
    temporary.replace(paths[3])
    return paths
