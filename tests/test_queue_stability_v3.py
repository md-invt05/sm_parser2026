"""Regression tests for balance-only draining and isolated valuation/export work."""

import asyncio
import hashlib
import os
import sqlite3
import time
import tracemalloc
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
import pytest

import scan_defi as scanner
from monitoring import MonitorStore
from sui_support import SuiStore, export_sui_xlsx
from telegram_bot import BotService


ADDRESS = "0x" + "ab" * 20
TOKEN = "0x" + "cd" * 20
ROOT = Path(__file__).resolve().parents[1]


def test_token_only_partial_migration_does_not_schedule_rpc_retry(tmp_path):
    db = scanner.DB(tmp_path / "db.sqlite")
    db.upsert_contracts([("ethereum", ADDRESS, 1, "0x1", None, "2026-01-01")])
    db.save_address_chain_state(ADDRESS, {
        "chain": "ethereum", "status": "partial", "has_code": 1,
        "total_usd": 0, "note": "token_coverage_partial: Transfer logs pending",
        "native_raw": "7",
    }, [], 500_000)
    # Recreate the old persisted status, before v3 separated token coverage.
    db.conn.execute(
        "UPDATE address_chain_state SET status='partial',coverage_state='verified' WHERE address=?",
        (ADDRESS,),
    )
    db.save_address_scan(ADDRESS, "incomplete", 0, 1, 1, [{
        "chain": "ethereum", "status": "partial", "has_code": 1,
        "total_usd": 0,
    }], [])
    before = db.conn.execute("SELECT COUNT(*) FROM address_scans").fetchone()[0]
    assert db.rephase_balance_schedule(500_000) == 1
    row = db.conn.execute("SELECT * FROM address_chain_state").fetchone()
    assert row["status"] == "complete"
    assert row["coverage_state"] == "historical_pending"
    assert row["failure_streak"] == 0
    assert row["native_raw"] == "7"
    assert row["next_retry_at"] > (datetime.now(timezone.utc) + timedelta(days=89)).isoformat()
    assert db.rephase_balance_schedule(500_000) == 0
    assert db.conn.execute("SELECT COUNT(*) FROM address_scans").fetchone()[0] == before
    db.close()


def test_aggregate_tiers_and_zero_chain_are_independent(tmp_path):
    db = scanner.DB(tmp_path / "db.sqlite")
    db.upsert_contracts([("ethereum", ADDRESS, 1, "0x1", None, "2026-01-01")])
    for chain, amount in (("ethereum", 80_000), ("base", 80_000)):
        db.save_address_chain_state(ADDRESS, {
            "chain": chain, "status": "complete", "has_code": 1,
            "total_usd": amount,
        }, [], 500_000)
    db.save_address_chain_state(ADDRESS, {
        "chain": "bsc", "status": "complete", "has_code": 1, "total_usd": 0,
    }, [], 500_000)
    db.apply_address_schedule(ADDRESS, "incomplete", 160_000, 500_000)
    rows = {row["chain"]: row for row in db.conn.execute("SELECT * FROM address_chain_state")}
    now = datetime.now(timezone.utc)
    assert (datetime.fromisoformat(rows["ethereum"]["next_retry_at"]) - now).days <= 3
    assert (datetime.fromisoformat(rows["base"]["next_retry_at"]) - now).days <= 3
    assert (datetime.fromisoformat(rows["bsc"]["next_retry_at"]) - now).days >= 89
    db.close()


def test_price_refresh_uses_stored_amount_and_does_not_rescan(tmp_path):
    db = scanner.DB(tmp_path / "db.sqlite")
    cfg = scanner.load_config(ROOT / "config.yaml")
    db.upsert_contracts([("ethereum", ADDRESS, 1, "0x1", None, "2026-01-01")])
    row = {
        "chain": "ethereum", "status": "price_missing", "has_code": 1,
        "native_raw": "0", "native_amount": 0, "included_native_usd": 0,
        "tokens_usd": 0, "total_usd": 0,
    }
    token = {
        "chain": "ethereum", "token": TOKEN, "raw_amount": "100", "amount": 100,
        "symbol": "TKN", "priced": False, "usd_value": None,
    }
    db.save_address_chain_state(ADDRESS, row, [token], cfg.min_usd)
    db.save_address_scan(ADDRESS, "incomplete", 0, 1, 1, [row], [token])
    scan_count = db.conn.execute("SELECT COUNT(*) FROM address_scans").fetchone()[0]
    assert db.apply_price_refresh(ADDRESS, "ethereum", {TOKEN: 2}, cfg.chains, cfg) == 1
    state = db.conn.execute("SELECT * FROM address_chain_state").fetchone()
    assert state["status"] == "complete"
    assert state["total_usd"] == 200
    assert state["price_retry_at"] is None
    assert db.conn.execute("SELECT raw_amount FROM address_token_state").fetchone()[0] == "100"
    assert db.conn.execute("SELECT COUNT(*) FROM address_scans").fetchone()[0] == scan_count
    assert db.latest_address_scans()[0]["total_usd"] == 200
    db.close()


def test_token_budget_balance_only_and_governor_hysteresis():
    async def check_budget():
        budget = scanner.TokenLogBudget(30, 4)
        budget.balance_only = True
        assert budget.limits() == (10, 2)
        await budget.acquire("ethereum")
        await budget.acquire("ethereum")
        assert "ethereum" not in await budget.available_chains(["ethereum", "base"])
        marker = scanner.TOKEN_LOG_HISTORY.set(True)
        try:
            await budget.acquire("base")
            await budget.acquire("base")
        finally:
            scanner.TOKEN_LOG_HISTORY.reset(marker)
        assert not await budget.history_available()
    asyncio.run(check_budget())
    governor = scanner.LoadGovernor(True, 6, 1, allow_balance_only=True)
    assert governor.evaluate(0, 0, 0, now=0, token_pending=10_000) == (1, 0)
    assert governor.evaluate(0, 0, 0, now=100, token_pending=4_000) == (1, 0)
    assert governor.evaluate(0, 0, 0, now=701, token_pending=4_000) == (6, 1)
    assert governor.evaluate(21_000, 0, 0, now=800) == (0, 0)


@pytest.mark.parametrize("profile", scanner.LOAD_PROFILES)
def test_only_steady_enters_automatic_balance_only(profile):
    limits = scanner.LOAD_PROFILES[profile]
    governor = scanner.LoadGovernor(
        True, limits["discovery_live_slots"], limits["discovery_backfill_slots"],
        allow_balance_only=profile == "steady",
    )
    slots = governor.evaluate(25_000, 7 * 3600, 0, now=0)
    if profile == "steady":
        assert governor.state == "balance_only"
        assert slots == (0, 0)
    else:
        assert governor.state == "drain"
        assert slots == (min(limits["discovery_live_slots"], 3), 0)
        assert governor.evaluate(0, 0, 0, now=1) == slots
        assert governor.evaluate(0, 0, 0, now=602) == (
            limits["discovery_live_slots"], limits["discovery_backfill_slots"]
        )


def test_deferred_low_value_history_is_not_reported_as_runnable(tmp_path):
    db = scanner.DB(tmp_path / "db.sqlite")
    now = datetime.now(timezone.utc).isoformat()
    with db.conn:
        db.conn.executemany(
            """INSERT INTO token_log_tasks(
                   chain,address,first_block,next_block,priority,due_at,updated_at,
                   recent_complete,completed_at)
               VALUES('ethereum',?,1,1,?,?,?,1,NULL)""",
            [(ADDRESS, 10, now, now), (TOKEN, 50, now, now)],
        )
    snapshot = db.token_log_queue_snapshot()["ethereum"]
    assert snapshot["total"] == 2
    assert snapshot["partial"] == 2
    assert snapshot["due"] == 1
    assert db.next_token_log_task(["ethereum"], allow_history=True)["address"] == TOKEN
    db.close()


def test_fast_governor_snapshot_skips_coverage_count(tmp_path):
    db = scanner.DB(tmp_path / "db.sqlite")
    db.upsert_contracts([("ethereum", ADDRESS, 1, "0x1", None, "2026-01-01")])
    db.save_address_chain_state(ADDRESS, {
        "chain": "ethereum", "status": "complete", "has_code": 1,
        "total_usd": 0, "coverage_state": "recent_only",
    }, [], 500_000)
    assert db.balance_queue_snapshot()["balance_token_coverage_waiting"] == 0
    assert db.balance_queue_snapshot(include_coverage=True)["balance_token_coverage_waiting"] == 1
    db.close()


def test_rpc_failure_preserves_independent_token_coverage(tmp_path):
    db = scanner.DB(tmp_path / "db.sqlite")
    db.save_address_chain_state(ADDRESS, {
        "chain": "ethereum", "status": "complete", "has_code": 1,
        "total_usd": 12, "coverage_state": "recent_only",
    }, [], 500_000)
    db.save_address_chain_state(ADDRESS, {
        "chain": "ethereum", "status": "rpc_error", "note": "provider timeout",
    }, [], 500_000)
    state = db.conn.execute("SELECT * FROM address_chain_state").fetchone()
    assert state["status"] == "rpc_error"
    assert state["total_usd"] == 12
    assert state["coverage_state"] == "recent_only"
    db.close()


def test_streaming_export_reads_details_in_bounded_batches(tmp_path):
    db = scanner.DB(tmp_path / "db.sqlite")
    cfg = scanner.load_config(ROOT / "config.yaml")
    cfg.export_dir = tmp_path / "reports"
    now = datetime.now(timezone.utc).isoformat()
    with db.conn:
        db.conn.executemany(
            "INSERT INTO address_scans(address,scanned_at,status,total_usd,coverage,total_networks) "
            "VALUES(?,?,'incomplete',0,1,1)",
            [("0x" + format(i, "040x"), now) for i in range(510)],
        )
    with patch.object(db, "latest_export_parts", wraps=db.latest_export_parts) as parts:
        scanner.export_xlsx(db, cfg, cfg.chains, cfg.min_usd, mode="file:incomplete")
        assert parts.call_count == 3
        assert max(len(call.args[0]) for call in parts.call_args_list) <= 200
    book = load_workbook(cfg.export_dir / "incomplete.xlsx", read_only=True)
    assert sum(1 for _ in book.active.values) == 511
    book.close()
    db.close()


@pytest.mark.skipif(os.getenv("RUN_EXPORT_STRESS") != "1", reason="explicit 50k-row export stress test")
def test_streaming_50k_export_stays_below_memory_budget(tmp_path):
    db = scanner.DB(tmp_path / "db.sqlite")
    cfg = scanner.load_config(ROOT / "config.yaml")
    cfg.export_dir = tmp_path / "reports"
    now = datetime.now(timezone.utc).isoformat()
    with db.conn:
        db.conn.executemany(
            "INSERT INTO address_scans(address,scanned_at,status,total_usd,coverage,total_networks) "
            "VALUES(?,?,'incomplete',0,1,1)",
            (("0x" + format(i, "040x"), now) for i in range(50_000)),
        )
    tracemalloc.start()
    scanner.export_xlsx(db, cfg, cfg.chains, cfg.min_usd, mode="file:incomplete")
    _current, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    assert peak < 128 * 1024 * 1024
    book = load_workbook(cfg.export_dir / "incomplete.xlsx", read_only=True)
    assert sum(1 for _ in book.active.values) == 50_001
    book.close()
    db.close()


def test_sui_stale_snapshot_keeps_last_known_category(tmp_path):
    store = SuiStore(tmp_path / "db.sqlite")
    store.sync_defi_projects([{
        "projectName": "Cetus", "currTvl": 1_000_000,
        "packages": [{"packageAddress": "0x" + "a" * 64}],
    }], 500_000)
    old = (datetime.now(timezone.utc) - timedelta(hours=3)).isoformat()
    store.conn.execute("UPDATE sui_defi_projects SET synced_at=?", (old,))
    store.mark_defi_sync_failed("provider down")
    paths = export_sui_xlsx(store, tmp_path / "reports", 500_000, mode="qualifying")
    book = load_workbook(paths[0], read_only=True)
    rows = list(book.active.values)
    assert rows[1][2] == "qualifying"
    assert "stale" in rows[1][-1]
    book.close()
    store.close()


def test_ram_incident_needs_two_minutes_to_recover(tmp_path):
    async def run():
        store = MonitorStore(tmp_path / "monitoring.db")
        service = BotService.__new__(BotService)
        service.monitor = store
        now = datetime.now(timezone.utc)
        for offset in range(300, -1, -15):
            store.add_resource_sample({
                "ts": (now - timedelta(seconds=offset)).isoformat(),
                "ram_percent": 94, "cpu_percent": 10, "cpu_cores": 4, "load1": 0,
            })
        await service._evaluate_resource_thresholds(now)
        assert store.active_incident("system:ram:critical") is not None
        assert store.active_incident("system:ram:warning") is None
        store.add_resource_sample({
            "ts": (now + timedelta(seconds=15)).isoformat(),
            "ram_percent": 60, "cpu_percent": 10, "cpu_cores": 4, "load1": 0,
        })
        await service._evaluate_resource_thresholds(now + timedelta(seconds=15))
        assert store.active_incident("system:ram:critical") is not None
        for offset in range(30, 151, 15):
            store.add_resource_sample({
                "ts": (now + timedelta(seconds=offset)).isoformat(),
                "ram_percent": 60, "cpu_percent": 10, "cpu_cores": 4, "load1": 0,
            })
        await service._evaluate_resource_thresholds(now + timedelta(seconds=150))
        assert store.active_incident("system:ram:critical") is None
        store.close()
    asyncio.run(run())


@pytest.mark.skipif(os.getenv("RUN_MIGRATION_COPY") != "1", reason="explicit local-DB copy audit")
def test_local_db_copy_migration_preserves_financial_records_and_cursors(tmp_path):
    source_path = ROOT / "data" / "contracts.db"
    if not source_path.exists():
        pytest.skip("no local database")
    source = sqlite3.connect(f"file:{source_path.as_posix()}?mode=ro", uri=True)
    target = sqlite3.connect(tmp_path / "copy.db")
    try:
        source.backup(target)
    finally:
        source.close()
        target.close()
    db = scanner.DB(tmp_path / "copy.db")
    def financial_digest():
        digest = hashlib.sha256()
        for row in db.conn.execute(
            "SELECT id,status,total_usd FROM address_scans ORDER BY id"
        ):
            digest.update(repr(tuple(row)).encode())
        return digest.hexdigest()
    before = financial_digest()
    cursors = [tuple(row) for row in db.conn.execute(
        "SELECT chain,role,next_block,last_committed FROM chain_cursors ORDER BY chain,role"
    )]
    db.rephase_balance_schedule(500_000)
    assert financial_digest() == before
    assert [tuple(row) for row in db.conn.execute(
        "SELECT chain,role,next_block,last_committed FROM chain_cursors ORDER BY chain,role"
    )] == cursors
    assert db.rephase_balance_schedule(500_000) == 0
    db.close()
