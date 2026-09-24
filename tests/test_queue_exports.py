import asyncio
import os
import tempfile
import time
import unittest
from dataclasses import replace
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

import scan_defi as scanner
import telegram_bot
from monitoring import MonitorStore
from sui_support import SuiStore


ADDRESS = "0x" + "ab" * 20
TOKEN = "0x" + "cd" * 20
ROOT = Path(__file__).resolve().parents[1]


class BalanceScheduleTests(unittest.TestCase):
    def test_tiers_and_error_backoff(self):
        delay = scanner.balance_retry_delay
        threshold = 500_000
        self.assertEqual(21600, delay("complete", 1, threshold, threshold))
        self.assertEqual(86400, delay("complete", 1, 50_000, threshold))
        self.assertEqual(259200, delay("complete", 1, 49_999, threshold))
        self.assertEqual(604800, delay("complete", 1, 0, threshold))
        self.assertEqual(2592000, delay("absent", 0, 0, threshold))
        self.assertEqual(21600, delay("price_missing", 1, None, threshold))
        self.assertEqual(86400, delay("anomalous_balance", 1, None, threshold))
        self.assertEqual([600, 1800, 7200, 21600, 86400], [
            delay("rpc_error", None, None, threshold, streak) for streak in range(1, 6)
        ])

    def test_new_token_wakes_only_its_chain_and_repeat_does_not(self):
        db = scanner.DB(Path(":memory:"))
        try:
            for chain in ("ethereum", "base"):
                db.save_address_chain_state(
                    ADDRESS, {"chain": chain, "has_code": 1, "status": "complete", "total_usd": 0},
                    [], 500_000,
                )
            original_base = db.conn.execute(
                "SELECT next_retry_at FROM address_chain_state WHERE address=? AND chain='base'",
                (ADDRESS,),
            ).fetchone()[0]
            db.upsert_contract_tokens([("ethereum", ADDRESS, TOKEN, "transfer", 1)])
            ethereum = db.conn.execute(
                "SELECT next_retry_at FROM address_chain_state WHERE address=? AND chain='ethereum'",
                (ADDRESS,),
            ).fetchone()[0]
            self.assertLessEqual(datetime.fromisoformat(ethereum).timestamp(), time.time() + 1)
            db.upsert_contract_tokens([("ethereum", ADDRESS, TOKEN, "transfer", 2)])
            self.assertEqual(original_base, db.conn.execute(
                "SELECT next_retry_at FROM address_chain_state WHERE address=? AND chain='base'",
                (ADDRESS,),
            ).fetchone()[0])
        finally:
            db.close()

    def test_rephase_is_idempotent_and_preserves_scans(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "contracts.db"
            db = scanner.DB(path)
            db.upsert_contracts([("ethereum", ADDRESS, 1, "0xtx", None, "2026-01-01")])
            db.save_address_chain_state(
                ADDRESS, {"chain": "ethereum", "has_code": 0, "status": "absent", "total_usd": 0},
                [], 500_000,
            )
            db.save_address_scan(
                ADDRESS, "below", 0, 1, 1,
                [{"chain": "ethereum", "has_code": 0, "status": "absent", "total_usd": 0}], [],
            )
            old_checked = (datetime.now(timezone.utc) - timedelta(days=2)).isoformat()
            db.conn.execute(
                "UPDATE address_chain_state SET checked_at=?,next_retry_at=? WHERE address=?",
                (old_checked, old_checked, ADDRESS),
            )
            self.assertEqual(1, db.balance_queue_snapshot()["balance_retry_pending"])
            before = db.conn.execute("SELECT COUNT(*) FROM address_scans").fetchone()[0]
            self.assertEqual(1, db.rephase_balance_schedule(500_000))
            self.assertEqual(0, db.balance_queue_snapshot()["balance_retry_pending"])
            due = db.conn.execute(
                "SELECT next_retry_at FROM address_chain_state WHERE address=?", (ADDRESS,),
            ).fetchone()[0]
            self.assertAlmostEqual(
                datetime.fromisoformat(due).timestamp() - datetime.fromisoformat(old_checked).timestamp(),
                30 * 86400, delta=1,
            )
            self.assertEqual(0, db.rephase_balance_schedule(500_000))
            self.assertEqual(before, db.conn.execute("SELECT COUNT(*) FROM address_scans").fetchone()[0])
            db.close()
            reopened = scanner.DB(path)
            self.assertEqual(0, reopened.rephase_balance_schedule(500_000))
            self.assertEqual(due, reopened.conn.execute(
                "SELECT next_retry_at FROM address_chain_state WHERE address=?", (ADDRESS,),
            ).fetchone()[0])
            reopened.close()

    def test_queue_categories_and_rechecks_are_counted(self):
        db = scanner.DB(Path(":memory:"))
        try:
            other = "0x" + "ef" * 20
            db.upsert_contracts([
                ("ethereum", ADDRESS, 1, "0xa", None, "2026-01-01"),
                ("ethereum", other, 2, "0xb", None, "2026-01-02"),
            ])
            db.save_address_chain_state(
                ADDRESS, {"chain": "ethereum", "has_code": None, "status": "rpc_error", "total_usd": None},
                [], 500_000,
            )
            db.conn.execute(
                "UPDATE address_chain_state SET next_retry_at='2026-01-01T00:00:00+00:00' WHERE address=?",
                (ADDRESS,),
            )
            queue = db.balance_queue_snapshot()
            self.assertEqual((1, 1, 2), (
                queue["balance_new_pending"], queue["balance_retry_pending"], queue["balance_pending"],
            ))
            self.assertEqual(other, db.pending_addresses(86400, limit=2)[0])
            for _ in range(2):
                db.save_address_scan(
                    ADDRESS, "incomplete", 0, 0, 1,
                    [{"chain": "ethereum", "has_code": None, "status": "rpc_error", "total_usd": None}], [],
                )
            snapshot = db.monitoring_snapshot(500_000)
            self.assertEqual(1, snapshot["balance_completed"])
            self.assertEqual(2, snapshot["balance_scans_total"])
        finally:
            db.close()


class ExportTests(unittest.TestCase):
    def test_requested_evm_file_only_and_selected_parts(self):
        with tempfile.TemporaryDirectory() as folder:
            db = scanner.DB(Path(folder) / "contracts.db")
            cfg = replace(scanner.load_config(ROOT / "config.yaml"), export_dir=Path(folder) / "reports")
            db.save_address_scan(
                ADDRESS, "below", 100, 1, 1,
                [{"chain": "ethereum", "has_code": 1, "status": "complete", "total_usd": 100}], [],
            )
            other = "0x" + "ef" * 20
            db.save_address_scan(
                other, "qualifying", 600_000, 1, 1,
                [{"chain": "ethereum", "has_code": 1, "status": "complete", "total_usd": 600_000}], [],
            )
            with patch.object(db, "latest_export_parts", wraps=db.latest_export_parts) as parts:
                scanner.export_xlsx(db, cfg, {"ethereum": cfg.chains["ethereum"]}, 500_000, mode="file:below")
                self.assertEqual([ADDRESS], [row["address"] for row in parts.call_args.args[0]])
            report = cfg.export_dir / "below_threshold.xlsx"
            self.assertTrue(report.exists())
            self.assertFalse((cfg.export_dir / "qualifying.xlsx").exists())
            book = load_workbook(report, read_only=True)
            self.assertEqual(["Address", ADDRESS], [row[0] for row in book.active.iter_rows(values_only=True)])
            book.close()
            db.close()

    def test_export_request_deduplicates_and_migrates_monitoring(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "monitoring.db"
            store = MonitorStore(path)
            first = store.request_export_once("1", "below")
            self.assertEqual(first, store.request_export_once("2", "below"))
            self.assertNotEqual(first, store.request_export_once("2", "incomplete"))
            full = store.request_control("export", "3")
            self.assertEqual(full, store.request_export_once("4", "sui_packages"))
            store.close()
            reopened = MonitorStore(path)
            self.assertEqual(1, reopened.rows(
                "SELECT COUNT(*) n FROM control_requests WHERE id=?", (first,)
            )[0]["n"])
            reopened.close()

    def test_sui_file_export_does_not_rebuild_evm_or_other_sui_files(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            db = scanner.DB(root / "contracts.db")
            db.close()
            sui = SuiStore(root / "contracts.db")
            sui.close()
            cfg = replace(scanner.load_config(ROOT / "config.yaml"), export_dir=root / "reports")
            scanner.export_bundle_process(
                root / "contracts.db", cfg, {"ethereum": cfg.chains["ethereum"]},
                500_000, mode="file", file_key="sui_below",
            )
            self.assertTrue((cfg.export_dir / "sui_below_threshold.xlsx").exists())
            self.assertFalse((cfg.export_dir / "qualifying.xlsx").exists())
            self.assertFalse((cfg.export_dir / "sui_packages.xlsx").exists())


class BotFileTests(unittest.IsolatedAsyncioTestCase):
    def test_repeat_scan_prevents_false_stall(self):
        sample = {
            "pending": 10, "min_done": 100, "max_done": 101,
            "first_ts": "2026-09-24T10:00:00+00:00",
            "last_ts": "2026-09-24T10:10:00+00:00",
        }
        self.assertFalse(telegram_bot.balance_scans_stalled(sample))
        sample["max_done"] = 100
        self.assertTrue(telegram_bot.balance_scans_stalled(sample))

    async def test_recent_file_is_sent_without_export(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "reports").mkdir()
            (root / "reports" / "below_threshold.xlsx").write_bytes(b"PK03")
            service = telegram_bot.BotService.__new__(telegram_bot.BotService)
            service.monitor = MonitorStore(root / "monitoring.db")
            service.guard = AsyncMock(return_value=True)
            service.send_file = AsyncMock()
            update = SimpleNamespace(
                effective_chat=SimpleNamespace(id=1),
                message=SimpleNamespace(reply_text=AsyncMock()),
            )
            with patch.object(telegram_bot, "ROOT", root):
                await service.cmd_file(update, SimpleNamespace(args=["below"], bot=object()))
            self.assertEqual([], service.monitor.pending_controls())
            service.send_file.assert_awaited_once()
            service.monitor.close()

    async def test_stale_file_queues_one_target_export(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "reports").mkdir()
            path = root / "reports" / "below_threshold.xlsx"
            path.write_bytes(b"PK03")
            old = time.time() - telegram_bot.FILE_FRESH_SEC - 1
            os.utime(path, (old, old))
            service = telegram_bot.BotService.__new__(telegram_bot.BotService)
            service.monitor = MonitorStore(root / "monitoring.db")
            service.guard = AsyncMock(return_value=True)
            service.wait_control = AsyncMock(return_value=True)
            service.send_file = AsyncMock()
            update = SimpleNamespace(
                effective_chat=SimpleNamespace(id=1),
                message=SimpleNamespace(reply_text=AsyncMock()),
            )
            with patch.object(telegram_bot, "ROOT", root):
                await asyncio.gather(*(
                    service.cmd_file(update, SimpleNamespace(args=["below"], bot=object()))
                    for _ in range(2)
                ))
            self.assertEqual(1, len(service.monitor.pending_controls()))
            self.assertEqual("export_file", service.monitor.pending_controls()[0].action)
            service.monitor.close()
