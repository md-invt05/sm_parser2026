#!/usr/bin/env python3
"""
DeFi liquidity scanner — автономный режим.

Индексирует прямые создания контрактов, считает мультичейн-балансы в USD,
пишет SQLite и три xlsx-отчёта.
При 429 / CU limit / перегрузе RPC сам уходит в паузу и продолжает.
Логи — в logs/, не в консоль (консоль только старт и редкий heartbeat).

    python scan_defi.py
    python scan_defi.py --rpc-check
"""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import logging
import math
import multiprocessing
import os
import random
import shutil
import signal
import sqlite3
import threading
import time
import traceback
from contextvars import ContextVar
from contextlib import asynccontextmanager
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal
from logging.handlers import RotatingFileHandler
from pathlib import Path
from typing import Any
from urllib.parse import urlparse
from zoneinfo import ZoneInfo

import httpx
import yaml
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from monitoring import LOAD_PROFILES, MonitorStore, percentile
from sui_support import (
    BlockberryClient,
    SuiConfig,
    SuiStore,
    export_sui_xlsx,
    sui_balance_loop,
    sui_discovery_loop,
)

load_dotenv()

ROOT = Path(__file__).resolve().parent
TRANSFER_TOPIC = "0xddf252ad1be2c89b69c2b068fc378daa952ba7f163c4a11628f55a4df523b3ef"
BALANCE_OF_SEL = "0x70a08231"
DECIMALS_SEL = "0x313ce567"
SYMBOL_SEL = "0x95d89b41"
ZERO = "0x0000000000000000000000000000000000000000"
LLAMA_PRICES = "https://coins.llama.fi/prices/current/"
# Task-local policy: balance probes must not inherit the indexer's long retries.
BALANCE_RPC = ContextVar("balance_rpc", default=False)

RATE_HINTS = (
    "rate limit",
    "too many requests",
    "too many request",
    "429",
    "compute units",
    "cu limit",
    "exceeded",
    "over rate",
    "capacity",
    "throttl",
    "slow down",
    "temporarily unavailable",
    "try again",
    "call rate",
    "credits",
    "quota",
    "limit reached",
    "projected cu",
)
RANGE_HINTS = (
    "block range",
    "query returned more",
    "response size",
    "log response size",
    "more than",
    "range limit",
    "too many results",
    "timeout",
    "timed out",
    "-32602",
)
AUTH_HINTS = (
    "unauthorized",
    "forbidden",
    "invalid api key",
    "missing api key",
    "api key required",
    "access token",
    "authentication required",
    "project id",
)
RETRYABLE_HTTP = {408, 425, 429, 500, 502, 503, 504, 520, 522, 524}

log = logging.getLogger("scanner")
console = logging.getLogger("scanner.console")


# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------


@dataclass
class ChainCfg:
    key: str
    enabled: bool
    chain_id: int
    display_name: str
    llama_chain: str
    native_symbol: str
    native_coingecko: str
    native_decimals: int
    start_block: int
    confirmations: int
    logs_max_range: int
    explorer: str
    rpc: list[str]
    lifecycle: str = "active"


@dataclass
class AppCfg:
    min_usd: float
    db_path: Path
    export_dir: Path
    log_dir: Path
    monitoring_db_path: Path
    recheck_interval_sec: int
    global_rpc_concurrency: int
    rpc_concurrency: int
    balance_concurrency: int
    balance_chain_concurrency: int
    balance_chain_timeout_sec: float
    balance_address_timeout_sec: float
    balance_retry_sec: int
    rabby_fallback: bool
    rabby_request_interval_sec: float
    rabby_timeout_sec: float
    rabby_uncertainty: float
    block_batch_size: int
    receipt_batch_size: int
    eth_call_batch_size: int
    price_batch_size: int
    http_timeout_sec: int
    max_retries: int
    cooldown_min_sec: float
    cooldown_max_sec: float
    all_down_sleep_sec: float
    adaptive_batch: bool
    heartbeat_sec: int
    export_every_sec: int
    daily_snapshot: bool
    task_restart_sec: int
    run_indexer: bool
    run_balance_checker: bool
    discover_tokens_from_transfers: bool
    discover_tx_to_contracts: bool
    code_cache_ttl_sec: int
    native_anomaly_usd: float
    rabby_token_discovery_after_sec: int
    valuation_policies: list[dict[str, Any]]
    default_chains: list[str]
    chains: dict[str, ChainCfg]
    sui: dict[str, Any]


def _env_rpcs(chain_key: str) -> list[str]:
    raw = os.getenv(f"{chain_key.upper()}_RPC", "").strip()
    if not raw:
        return []
    return [u.strip() for u in raw.split(",") if u.strip()]


DRPC_FALLBACK_SLUGS = {
    "ethereum": "ethereum", "bsc": "bsc", "polygon": "polygon",
    "arbitrum": "arbitrum", "optimism": "optimism", "base": "base",
    "zk": "polygon-zkevm", "zksync": "zksync", "robinhood": "robinhood",
    "linea": "linea", "scroll": "scroll", "mantle": "mantle", "blast": "blast",
    "celo": "celo", "gnosis": "gnosis", "cronos": "cronos", "kava": "kava",
    "metis": "metis", "harmony": "harmony-0",
}


def _drpc_fallback_rpcs(chain_key: str) -> list[str]:
    """Build secondary keyed dRPC URLs without duplicating every network URL in .env."""
    slug = DRPC_FALLBACK_SLUGS.get(chain_key)
    raw = os.getenv("DRPC_FALLBACK_KEYS", "").strip()
    if not slug or not raw:
        return []
    return [
        f"https://lb.drpc.live/{slug}/{token.strip()}"
        for token in raw.split(",")
        if token.strip()
    ]


def load_config(path: Path) -> AppCfg:
    raw = yaml.safe_load(path.read_text())
    uncertainty = float(raw.get("rabby_uncertainty", 0.10))
    if not 0 < uncertainty < 1:
        raise ValueError("rabby_uncertainty must be between 0 and 1")
    chains: dict[str, ChainCfg] = {}
    for key, c in raw["chains"].items():
        # Keyed/private endpoints have priority, public endpoints stay as fallback.
        rpcs = list(dict.fromkeys(
            _env_rpcs(key) + _drpc_fallback_rpcs(key) + list(c.get("rpc") or [])
        ))
        chains[key] = ChainCfg(
            key=key,
            enabled=bool(c.get("enabled", True)),
            chain_id=int(c["chain_id"]),
            display_name=c.get("display_name", key),
            llama_chain=c["llama_chain"],
            native_symbol=c["native_symbol"],
            native_coingecko=c["native_coingecko"],
            native_decimals=int(c.get("native_decimals", 18)),
            start_block=int(c.get("start_block") or 1),
            confirmations=int(c.get("confirmations", 20)),
            logs_max_range=int(c.get("logs_max_range") or 500),
            explorer=c.get("explorer", ""),
            rpc=rpcs,
            lifecycle=str(c.get("lifecycle", "active")),
        )
    return AppCfg(
        min_usd=float(raw["min_usd"]),
        db_path=(ROOT / raw.get("db_path", "data/contracts.db")).resolve(),
        export_dir=(ROOT / raw.get("export_dir", "reports")).resolve(),
        log_dir=(ROOT / raw.get("log_dir", "logs")).resolve(),
        monitoring_db_path=(ROOT / raw.get("monitoring_db_path", "data/monitoring.db")).resolve(),
        recheck_interval_sec=int(raw.get("recheck_interval_sec", 86400)),
        global_rpc_concurrency=int(raw.get("global_rpc_concurrency", 12)),
        rpc_concurrency=int(raw.get("rpc_concurrency", 2)),
        balance_concurrency=max(1, int(raw.get("balance_concurrency", 8))),
        balance_chain_concurrency=max(1, int(raw.get("balance_chain_concurrency", 12))),
        balance_chain_timeout_sec=max(0.01, float(raw.get("balance_chain_timeout_sec", 20))),
        balance_address_timeout_sec=max(0.01, float(raw.get("balance_address_timeout_sec", 60))),
        balance_retry_sec=max(1, int(raw.get("balance_retry_sec", 300))),
        rabby_fallback=bool(raw.get("rabby_fallback", True)),
        rabby_request_interval_sec=max(0.0, float(raw.get("rabby_request_interval_sec", 3))),
        rabby_timeout_sec=max(0.01, float(raw.get("rabby_timeout_sec", 12))),
        rabby_uncertainty=uncertainty,
        block_batch_size=int(raw.get("block_batch_size", 8)),
        receipt_batch_size=int(raw.get("receipt_batch_size", 20)),
        eth_call_batch_size=int(raw.get("eth_call_batch_size", 10)),
        price_batch_size=int(raw.get("price_batch_size", 40)),
        http_timeout_sec=int(raw.get("http_timeout_sec", 40)),
        max_retries=int(raw.get("max_retries", 8)),
        cooldown_min_sec=float(raw.get("cooldown_min_sec", 5)),
        cooldown_max_sec=float(raw.get("cooldown_max_sec", 900)),
        all_down_sleep_sec=float(raw.get("all_down_sleep_sec", 30)),
        adaptive_batch=bool(raw.get("adaptive_batch", True)),
        heartbeat_sec=max(10, int(raw.get("heartbeat_sec", 60))),
        export_every_sec=int(raw.get("export_every_sec", 300)),
        daily_snapshot=bool(raw.get("daily_snapshot", True)),
        task_restart_sec=int(raw.get("task_restart_sec", 20)),
        run_indexer=bool(raw.get("run_indexer", True)),
        run_balance_checker=bool(raw.get("run_balance_checker", True)),
        discover_tokens_from_transfers=bool(raw.get("discover_tokens_from_transfers", True)),
        discover_tx_to_contracts=bool(raw.get("discover_tx_to_contracts", True)),
        code_cache_ttl_sec=max(1, int(raw.get("code_cache_ttl_sec", 86400))),
        native_anomaly_usd=max(0.0, float(raw.get("native_anomaly_usd", 100_000_000))),
        rabby_token_discovery_after_sec=max(
            60, int(raw.get("rabby_token_discovery_after_sec", 900))
        ),
        valuation_policies=list(raw.get("valuation_policies") or []),
        default_chains=list(raw.get("default_chains") or [
            key for key, chain in chains.items() if chain.enabled
        ]),
        chains=chains,
        sui=dict(raw.get("sui") or {}),
    )


def apply_load_profile(cfg: AppCfg, profile: str) -> None:
    if profile not in LOAD_PROFILES:
        profile = "normal"
    for name, value in LOAD_PROFILES[profile].items():
        setattr(cfg, name, value)


def load_token_seed(path: Path) -> dict[str, list[dict[str, Any]]]:
    if not path.exists():
        return {}
    return yaml.safe_load(path.read_text()) or {}


def setup_logging(log_dir: Path) -> None:
    log_dir.mkdir(parents=True, exist_ok=True)
    fmt = logging.Formatter("%(asctime)s %(levelname)-7s %(message)s", "%Y-%m-%d %H:%M:%S")

    root = logging.getLogger()
    root.handlers.clear()
    root.setLevel(logging.INFO)

    file_h = RotatingFileHandler(
        log_dir / "scanner.log",
        maxBytes=20 * 1024 * 1024,
        backupCount=14,
        encoding="utf-8",
    )
    file_h.setLevel(logging.INFO)
    file_h.setFormatter(fmt)
    root.addHandler(file_h)

    err_h = RotatingFileHandler(
        log_dir / "errors.log",
        maxBytes=10 * 1024 * 1024,
        backupCount=7,
        encoding="utf-8",
    )
    err_h.setLevel(logging.WARNING)
    err_h.setFormatter(fmt)
    root.addHandler(err_h)
    # httpx INFO records include complete URLs and may leak API keys from .env.
    logging.getLogger("httpx").setLevel(logging.WARNING)
    logging.getLogger("httpcore").setLevel(logging.WARNING)

    # Консоль: только явные сообщения scanner.console (старт / heartbeat).
    ch = logging.StreamHandler()
    ch.setLevel(logging.INFO)
    ch.setFormatter(logging.Formatter("%(message)s"))
    console.setLevel(logging.INFO)
    console.propagate = False
    console.addHandler(ch)


# ---------------------------------------------------------------------------
# SQLite
# ---------------------------------------------------------------------------


SCHEMA = """
PRAGMA journal_mode=WAL;
PRAGMA synchronous=NORMAL;
PRAGMA foreign_keys=ON;

CREATE TABLE IF NOT EXISTS chain_state (
    chain            TEXT PRIMARY KEY,
    start_block      INTEGER NOT NULL,
    last_indexed     INTEGER NOT NULL DEFAULT 0,
    discovery_revision INTEGER NOT NULL DEFAULT 2
);

CREATE TABLE IF NOT EXISTS contracts (
    chain            TEXT NOT NULL,
    address          TEXT NOT NULL,
    created_block    INTEGER,
    created_tx       TEXT,
    creator          TEXT,
    first_seen_at    TEXT NOT NULL,
    last_checked_at  TEXT,
    last_total_usd   REAL,
    last_status      TEXT,
    PRIMARY KEY (chain, address)
);

CREATE TABLE IF NOT EXISTS tokens (
    chain            TEXT NOT NULL,
    address          TEXT NOT NULL,
    symbol           TEXT,
    decimals         INTEGER,
    source           TEXT,
    PRIMARY KEY (chain, address)
);

CREATE TABLE IF NOT EXISTS scans (
    id               INTEGER PRIMARY KEY AUTOINCREMENT,
    chain            TEXT NOT NULL,
    address          TEXT NOT NULL,
    scanned_at       TEXT NOT NULL,
    native_amount    REAL,
    native_usd       REAL,
    tokens_usd       REAL,
    total_usd        REAL,
    meets_threshold  INTEGER,
    note             TEXT
);

CREATE TABLE IF NOT EXISTS scan_tokens (
    scan_id          INTEGER NOT NULL REFERENCES scans(id) ON DELETE CASCADE,
    token            TEXT NOT NULL,
    symbol           TEXT,
    amount           REAL,
    price_usd        REAL,
    usd_value        REAL
);

CREATE INDEX IF NOT EXISTS idx_contracts_check ON contracts(chain, last_checked_at);
CREATE INDEX IF NOT EXISTS idx_scans_addr ON scans(chain, address, scanned_at);

CREATE TABLE IF NOT EXISTS schema_meta (
    version          INTEGER NOT NULL
);

CREATE TABLE IF NOT EXISTS contract_tokens (
    chain            TEXT NOT NULL,
    contract         TEXT NOT NULL,
    token            TEXT NOT NULL,
    source           TEXT NOT NULL DEFAULT 'transfer',
    first_seen_block INTEGER,
    last_seen_block  INTEGER,
    PRIMARY KEY (chain, contract, token)
);

CREATE TABLE IF NOT EXISTS address_scans (
    id               INTEGER PRIMARY KEY AUTOINCREMENT,
    address          TEXT NOT NULL,
    scanned_at       TEXT NOT NULL,
    status           TEXT NOT NULL,
    total_usd        REAL NOT NULL,
    coverage         INTEGER NOT NULL,
    total_networks   INTEGER NOT NULL,
    note             TEXT
);

CREATE TABLE IF NOT EXISTS address_chain_scans (
    scan_id          INTEGER NOT NULL REFERENCES address_scans(id) ON DELETE CASCADE,
    chain            TEXT NOT NULL,
    has_code         INTEGER,
    status           TEXT NOT NULL,
    native_amount    REAL,
    native_usd       REAL,
    tokens_usd       REAL,
    total_usd        REAL,
    note             TEXT,
    PRIMARY KEY (scan_id, chain)
);

CREATE TABLE IF NOT EXISTS address_token_scans (
    scan_id          INTEGER NOT NULL REFERENCES address_scans(id) ON DELETE CASCADE,
    chain            TEXT NOT NULL,
    token            TEXT NOT NULL,
    symbol           TEXT,
    raw_amount       TEXT,
    amount           REAL,
    price_usd        REAL,
    usd_value        REAL,
    priced           INTEGER NOT NULL,
    PRIMARY KEY (scan_id, chain, token)
);

CREATE INDEX IF NOT EXISTS idx_contract_tokens_holder
    ON contract_tokens(chain, contract);
CREATE INDEX IF NOT EXISTS idx_address_scans_latest
    ON address_scans(address, scanned_at);

CREATE TABLE IF NOT EXISTS rabby_estimates (
    address          TEXT PRIMARY KEY,
    checked_at       TEXT NOT NULL,
    status           TEXT NOT NULL,
    estimated_usd    REAL,
    lower_usd        REAL,
    upper_usd        REAL,
    coverage         INTEGER NOT NULL,
    total_networks   INTEGER NOT NULL,
    uncertainty      REAL NOT NULL,
    chain_parts      TEXT NOT NULL,
    token_parts      TEXT NOT NULL,
    note             TEXT,
    rpc_scan_id      INTEGER NOT NULL REFERENCES address_scans(id)
);

CREATE TABLE IF NOT EXISTS contract_discoveries (
    chain            TEXT NOT NULL,
    address          TEXT NOT NULL,
    source           TEXT NOT NULL,
    observed_block   INTEGER,
    observed_tx      TEXT,
    actor            TEXT,
    first_seen_at    TEXT NOT NULL,
    PRIMARY KEY (chain, address, source)
);

CREATE TABLE IF NOT EXISTS contract_code_cache (
    chain            TEXT NOT NULL,
    address          TEXT NOT NULL,
    has_code         INTEGER NOT NULL,
    checked_at       TEXT NOT NULL,
    checked_block    INTEGER,
    PRIMARY KEY (chain, address)
);

CREATE INDEX IF NOT EXISTS idx_contract_discoveries_address
    ON contract_discoveries(address, chain);
CREATE INDEX IF NOT EXISTS idx_contract_code_cache_checked
    ON contract_code_cache(chain, checked_at);

CREATE TABLE IF NOT EXISTS chain_cursors (
    chain            TEXT NOT NULL,
    role             TEXT NOT NULL CHECK(role IN ('live','backfill')),
    next_block       INTEGER NOT NULL,
    anchor_block     INTEGER NOT NULL,
    last_committed   INTEGER NOT NULL,
    status           TEXT NOT NULL DEFAULT 'active',
    updated_at       TEXT NOT NULL,
    note             TEXT,
    PRIMARY KEY (chain, role)
);

CREATE TABLE IF NOT EXISTS address_chain_state (
    address          TEXT NOT NULL,
    chain            TEXT NOT NULL,
    checked_at       TEXT NOT NULL,
    last_success_at  TEXT,
    status           TEXT NOT NULL,
    has_code         INTEGER,
    native_raw       TEXT,
    native_amount    REAL,
    observed_native_usd REAL,
    included_native_usd REAL,
    excluded_usd     REAL NOT NULL DEFAULT 0,
    tokens_usd       REAL,
    total_usd        REAL,
    valuation_status TEXT,
    failure_streak   INTEGER NOT NULL DEFAULT 0,
    next_retry_at    TEXT NOT NULL,
    note             TEXT,
    PRIMARY KEY (address, chain)
);
CREATE INDEX IF NOT EXISTS idx_address_chain_retry
    ON address_chain_state(next_retry_at, address);

CREATE TABLE IF NOT EXISTS address_token_state (
    address          TEXT NOT NULL,
    chain            TEXT NOT NULL,
    token            TEXT NOT NULL,
    checked_at       TEXT NOT NULL,
    raw_amount       TEXT NOT NULL,
    amount           REAL,
    symbol           TEXT,
    price_usd        REAL,
    usd_value        REAL,
    priced           INTEGER NOT NULL,
    valuation_status TEXT NOT NULL DEFAULT 'included',
    PRIMARY KEY (address, chain, token)
);

CREATE TABLE IF NOT EXISTS asset_valuation_policies (
    chain            TEXT NOT NULL,
    address          TEXT NOT NULL,
    asset            TEXT NOT NULL,
    policy           TEXT NOT NULL,
    reason           TEXT NOT NULL,
    source           TEXT,
    updated_at       TEXT NOT NULL,
    PRIMARY KEY (chain, address, asset)
);

CREATE TABLE IF NOT EXISTS anomalous_balances (
    chain            TEXT NOT NULL,
    address          TEXT NOT NULL,
    asset            TEXT NOT NULL,
    detected_at      TEXT NOT NULL,
    last_seen_at     TEXT NOT NULL,
    raw_amount       TEXT NOT NULL,
    observed_usd     REAL,
    second_rpc_match INTEGER,
    early_block_raw  TEXT,
    status           TEXT NOT NULL,
    evidence_json    TEXT NOT NULL DEFAULT '{}',
    PRIMARY KEY (chain, address, asset)
);

CREATE TABLE IF NOT EXISTS rpc_method_health (
    chain            TEXT NOT NULL,
    endpoint_fp      TEXT NOT NULL,
    hostname         TEXT NOT NULL,
    method_group     TEXT NOT NULL,
    fail_streak      INTEGER NOT NULL DEFAULT 0,
    success_streak   INTEGER NOT NULL DEFAULT 0,
    cooldown_until   REAL NOT NULL DEFAULT 0,
    circuit_until    REAL NOT NULL DEFAULT 0,
    permanent_error  TEXT,
    latency_ewma_ms  REAL,
    batch_limit      INTEGER,
    success_since_resize INTEGER NOT NULL DEFAULT 0,
    updated_at       TEXT NOT NULL,
    PRIMARY KEY (chain, endpoint_fp, method_group)
);

CREATE TABLE IF NOT EXISTS runtime_meta (
    key              TEXT PRIMARY KEY,
    value            TEXT NOT NULL,
    updated_at       TEXT NOT NULL
);
"""


class DB:
    def __init__(self, path: Path, *, read_only: bool = False):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.path = path
        self.read_only = read_only
        self._lock = threading.RLock()
        if read_only:
            uri = f"file:{path.as_posix()}?mode=ro"
            self.conn = sqlite3.connect(uri, uri=True, check_same_thread=False, timeout=60)
        else:
            self.conn = sqlite3.connect(path, check_same_thread=False, timeout=60)
        self.conn.row_factory = sqlite3.Row
        if read_only:
            self.conn.execute("PRAGMA query_only=ON")
            return
        self.conn.executescript(SCHEMA)
        chain_state_columns = {
            row["name"] for row in self.conn.execute("PRAGMA table_info(chain_state)")
        }
        if "discovery_revision" not in chain_state_columns:
            # Existing cursors are revision 1 (deployment-only). They are rewound
            # lazily by init_chain only when tx.to discovery is actually enabled.
            self.conn.execute(
                "ALTER TABLE chain_state ADD COLUMN discovery_revision INTEGER NOT NULL DEFAULT 1"
            )
        rabby_columns = {
            row["name"] for row in self.conn.execute("PRAGMA table_info(rabby_estimates)")
        }
        if "rpc_scan_id" not in rabby_columns:
            self.conn.execute("ALTER TABLE rabby_estimates ADD COLUMN rpc_scan_id INTEGER")
        chain_scan_columns = {
            row["name"] for row in self.conn.execute("PRAGMA table_info(address_chain_scans)")
        }
        for name, declaration in {
            "native_raw": "TEXT",
            "observed_native_usd": "REAL",
            "included_native_usd": "REAL",
            "excluded_usd": "REAL NOT NULL DEFAULT 0",
            "valuation_status": "TEXT",
        }.items():
            if name not in chain_scan_columns:
                self.conn.execute(
                    f"ALTER TABLE address_chain_scans ADD COLUMN {name} {declaration}"
                )
        token_scan_columns = {
            row["name"] for row in self.conn.execute("PRAGMA table_info(address_token_scans)")
        }
        if "valuation_status" not in token_scan_columns:
            self.conn.execute(
                "ALTER TABLE address_token_scans ADD COLUMN valuation_status TEXT NOT NULL DEFAULT 'included'"
            )
        self.conn.execute(
            """
            INSERT OR IGNORE INTO contract_discoveries(
                chain, address, source, observed_block, observed_tx, actor, first_seen_at
            )
            SELECT chain, lower(address), 'direct_deploy', created_block, created_tx,
                   creator, first_seen_at
            FROM contracts
            WHERE created_block IS NOT NULL OR created_tx IS NOT NULL
            """
        )
        self.conn.execute(
            "INSERT INTO schema_meta(version) SELECT 7 WHERE NOT EXISTS (SELECT 1 FROM schema_meta)"
        )
        self.conn.execute("UPDATE schema_meta SET version=7 WHERE version < 7")
        self.conn.execute(
            "INSERT OR IGNORE INTO runtime_meta(key,value,updated_at) VALUES('data_revision','0',?)",
            (datetime.now(timezone.utc).isoformat(),),
        )
        self.conn.commit()

    def close(self) -> None:
        self.conn.close()

    def _bump_revision_locked(self) -> int:
        row = self.conn.execute(
            "SELECT CAST(value AS INTEGER) FROM runtime_meta WHERE key='data_revision'"
        ).fetchone()
        revision = int(row[0] if row else 0) + 1
        self.conn.execute(
            """INSERT INTO runtime_meta(key,value,updated_at) VALUES('data_revision',?,?)
               ON CONFLICT(key) DO UPDATE SET value=excluded.value,updated_at=excluded.updated_at""",
            (str(revision), datetime.now(timezone.utc).isoformat()),
        )
        return revision

    def data_revision(self) -> int:
        with self._lock:
            row = self.conn.execute(
                "SELECT value FROM runtime_meta WHERE key='data_revision'"
            ).fetchone()
        return int(row[0]) if row else 0

    def seed_valuation_policies(self, rows: list[dict[str, Any]]) -> None:
        if not rows:
            return
        now = datetime.now(timezone.utc).isoformat()
        values = []
        for row in rows:
            address = normalize_evm_address(row.get("address"))
            if address is None:
                raise ValueError(f"invalid valuation policy address: {row.get('address')}")
            policy = str(row.get("policy") or "").strip()
            if policy not in {"include", "exclude_from_total", "quarantine"}:
                raise ValueError(f"invalid valuation policy: {policy}")
            values.append((
                str(row["chain"]), address, str(row.get("asset") or "native").lower(),
                policy, str(row.get("reason") or "manual policy"), row.get("source"), now,
            ))
        with self._lock, self.conn:
            self.conn.executemany(
                """INSERT INTO asset_valuation_policies(
                       chain,address,asset,policy,reason,source,updated_at
                   ) VALUES(?,?,?,?,?,?,?)
                   ON CONFLICT(chain,address,asset) DO UPDATE SET
                     policy=excluded.policy,reason=excluded.reason,
                     source=excluded.source,updated_at=excluded.updated_at""",
                values,
            )

    def valuation_policy(self, chain: str, address: str, asset: str) -> sqlite3.Row | None:
        with self._lock:
            return self.conn.execute(
                "SELECT * FROM asset_valuation_policies WHERE chain=? AND address=? AND asset=?",
                (chain, address.lower(), asset.lower()),
            ).fetchone()

    def revalue_system_balances(self, min_usd: float) -> int:
        """Revalue historical aggregate scans without deleting their raw observations."""
        with self._lock, self.conn:
            policies = self.conn.execute(
                """SELECT chain,address,reason FROM asset_valuation_policies
                   WHERE asset='native' AND policy='exclude_from_total'"""
            ).fetchall()
            touched: set[int] = set()
            for policy in policies:
                rows = self.conn.execute(
                    """SELECT c.scan_id,c.native_usd,c.total_usd
                       FROM address_chain_scans c JOIN address_scans a ON a.id=c.scan_id
                       WHERE c.chain=? AND lower(a.address)=? AND c.has_code=1
                         AND COALESCE(c.valuation_status,'included')!='exclude_from_total'""",
                    (policy["chain"], policy["address"]),
                ).fetchall()
                for row in rows:
                    observed = float(row["native_usd"] or 0.0)
                    self.conn.execute(
                        """UPDATE address_chain_scans SET
                             native_raw=COALESCE(native_raw,(
                               SELECT raw_amount FROM address_token_scans t
                               WHERE t.scan_id=address_chain_scans.scan_id
                                 AND t.chain=address_chain_scans.chain AND t.token='native'
                             )),
                             observed_native_usd=COALESCE(observed_native_usd,native_usd),
                             included_native_usd=0,
                             excluded_usd=MAX(excluded_usd,?),
                             valuation_status='exclude_from_total',
                             total_usd=MAX(0,COALESCE(total_usd,0)-?)
                           WHERE scan_id=? AND chain=?""",
                        (observed, observed, row["scan_id"], policy["chain"]),
                    )
                    self.conn.execute(
                        """UPDATE address_token_scans SET valuation_status='exclude_from_total'
                           WHERE scan_id=? AND chain=? AND token='native'""",
                        (row["scan_id"], policy["chain"]),
                    )
                    touched.add(int(row["scan_id"]))
            for scan_id in touched:
                total = float(self.conn.execute(
                    "SELECT COALESCE(SUM(total_usd),0) FROM address_chain_scans WHERE scan_id=?",
                    (scan_id,),
                ).fetchone()[0])
                parts = self.conn.execute(
                    "SELECT status FROM address_chain_scans WHERE scan_id=?", (scan_id,)
                ).fetchall()
                status = (
                    "qualifying" if total >= min_usd
                    else "below" if parts and all(row["status"] in ("complete", "absent") for row in parts)
                    else "incomplete"
                )
                self.conn.execute(
                    "UPDATE address_scans SET total_usd=?,status=?,note=TRIM(COALESCE(note,'') || ', system_balance_excluded',', ') WHERE id=?",
                    (total, status, scan_id),
                )
            if touched:
                self._bump_revision_locked()
            return len(touched)

    def reclassify_latest_scans(self, min_usd: float) -> int:
        changed = 0
        with self._lock, self.conn:
            latest = self.conn.execute(
                """SELECT a.* FROM address_scans a WHERE a.id=(
                     SELECT a2.id FROM address_scans a2 WHERE a2.address=a.address
                     ORDER BY a2.scanned_at DESC,a2.id DESC LIMIT 1)"""
            ).fetchall()
            for scan in latest:
                parts = self.conn.execute(
                    "SELECT status FROM address_chain_scans WHERE scan_id=?", (scan["id"],)
                ).fetchall()
                total = float(scan["total_usd"] or 0.0)
                status = (
                    "qualifying" if total >= min_usd
                    else "below" if parts and all(
                        row["status"] in ("complete", "absent") for row in parts
                    ) else "incomplete"
                )
                if status != scan["status"]:
                    self.conn.execute(
                        "UPDATE address_scans SET status=? WHERE id=?", (status, scan["id"])
                    )
                    changed += 1
                self.conn.execute(
                    "UPDATE contracts SET last_status=?,last_total_usd=? WHERE lower(address)=?",
                    (status, total, scan["address"].lower()),
                )
            if changed:
                self._bump_revision_locked()
        return changed

    def save_anomaly(
        self, chain: str, address: str, asset: str, raw_amount: int,
        observed_usd: float | None, second_rpc_match: bool | None,
        early_block_raw: int | None, status: str, evidence: dict[str, Any],
    ) -> None:
        now = datetime.now(timezone.utc).isoformat()
        with self._lock, self.conn:
            self.conn.execute(
                """INSERT INTO anomalous_balances(
                       chain,address,asset,detected_at,last_seen_at,raw_amount,
                       observed_usd,second_rpc_match,early_block_raw,status,evidence_json
                   ) VALUES(?,?,?,?,?,?,?,?,?,?,?)
                   ON CONFLICT(chain,address,asset) DO UPDATE SET
                     last_seen_at=excluded.last_seen_at,raw_amount=excluded.raw_amount,
                     observed_usd=excluded.observed_usd,
                     second_rpc_match=excluded.second_rpc_match,
                     early_block_raw=excluded.early_block_raw,status=excluded.status,
                     evidence_json=excluded.evidence_json""",
                (chain, address.lower(), asset.lower(), now, now, str(raw_amount),
                 observed_usd, None if second_rpc_match is None else int(second_rpc_match),
                 None if early_block_raw is None else str(early_block_raw), status,
                 json.dumps(evidence, ensure_ascii=False, sort_keys=True)),
            )

    def init_chain_cursors(
        self, chain: str, start_block: int, safe_head: int, lookback: int = 2,
    ) -> dict[str, sqlite3.Row]:
        now = datetime.now(timezone.utc).isoformat()
        with self._lock, self.conn:
            existing_live_before = self.conn.execute(
                "SELECT last_committed FROM chain_cursors WHERE chain=? AND role='live'",
                (chain,),
            ).fetchone()
            old = self.conn.execute(
                "SELECT last_indexed FROM chain_state WHERE chain=?", (chain,)
            ).fetchone()
            historical = int(old[0]) if old else max(0, start_block - 1)
            backfill_next = max(start_block, historical + 1)
            backfill_status = "complete" if backfill_next > safe_head else "active"
            live_next = max(start_block, safe_head - max(0, lookback) + 1)
            self.conn.execute(
                """INSERT OR IGNORE INTO chain_cursors(
                       chain,role,next_block,anchor_block,last_committed,status,updated_at,note
                   ) VALUES(?,?,?,?,?,?,?,?)""",
                (chain, "backfill", backfill_next, safe_head, backfill_next - 1,
                 backfill_status, now, "migrated from chain_state"),
            )
            self.conn.execute(
                """INSERT OR IGNORE INTO chain_cursors(
                       chain,role,next_block,anchor_block,last_committed,status,updated_at,note
                   ) VALUES(?,?,?,?,?,?,?,?)""",
                (chain, "live", live_next, safe_head, live_next - 1,
                 "active", now, "safe-head anchor with two-block lookback"),
            )
            if existing_live_before is not None:
                replay_from = max(
                    start_block, int(existing_live_before["last_committed"]) - lookback + 1
                )
                self.conn.execute(
                    """UPDATE chain_cursors SET next_block=?,last_committed=?,anchor_block=?,
                         status='active',updated_at=?,note='startup reorg lookback'
                       WHERE chain=? AND role='live'""",
                    (replay_from, replay_from - 1, safe_head, now, chain),
                )
            rows = self.conn.execute(
                "SELECT * FROM chain_cursors WHERE chain=?", (chain,)
            ).fetchall()
        return {str(row["role"]): row for row in rows}

    def cursor(self, chain: str, role: str) -> sqlite3.Row | None:
        with self._lock:
            return self.conn.execute(
                "SELECT * FROM chain_cursors WHERE chain=? AND role=?", (chain, role)
            ).fetchone()

    def advance_cursor_start(self, chain: str, role: str, next_block: int) -> None:
        with self._lock, self.conn:
            row = self.conn.execute(
                "SELECT next_block FROM chain_cursors WHERE chain=? AND role=?", (chain, role)
            ).fetchone()
            if row is not None and int(row["next_block"]) < next_block:
                self.conn.execute(
                    """UPDATE chain_cursors SET next_block=?,last_committed=?,updated_at=?,
                         note='advanced by --from-block' WHERE chain=? AND role=?""",
                    (next_block, next_block - 1, datetime.now(timezone.utc).isoformat(), chain, role),
                )

    def commit_index_range(
        self, chain: str, role: str, range_end: int,
        contract_rows: list[tuple[Any, ...]],
        discovery_rows: list[tuple[Any, ...]],
        cache_rows: list[tuple[Any, ...]],
        token_rows: list[tuple[Any, ...]],
        relation_rows: list[tuple[Any, ...]],
    ) -> tuple[int, int]:
        """Commit all derived data and the matching cursor in one transaction."""
        now = datetime.now(timezone.utc).isoformat()
        with self._lock, self.conn:
            cursor = self.conn.execute(
                "SELECT next_block,anchor_block FROM chain_cursors WHERE chain=? AND role=?",
                (chain, role),
            ).fetchone()
            if cursor is None:
                raise RuntimeError(f"missing {chain}/{role} cursor")
            before = self.conn.total_changes
            self.conn.executemany(
                """INSERT OR IGNORE INTO contracts(
                       chain,address,created_block,created_tx,creator,first_seen_at
                   ) VALUES(?,?,?,?,?,?)""",
                contract_rows,
            )
            inserted = self.conn.total_changes - before
            self.conn.executemany(
                """UPDATE contracts SET created_block=COALESCE(created_block,?),
                     created_tx=COALESCE(created_tx,?),creator=COALESCE(creator,?)
                   WHERE chain=? AND address=? AND ? IS NOT NULL""",
                [(row[2], row[3], row[4], row[0], row[1], row[2]) for row in contract_rows],
            )
            before_sources = self.conn.total_changes
            self.conn.executemany(
                """INSERT OR IGNORE INTO contract_discoveries(
                       chain,address,source,observed_block,observed_tx,actor,first_seen_at
                   ) VALUES(?,?,?,?,?,?,?)""",
                discovery_rows,
            )
            sources_inserted = self.conn.total_changes - before_sources
            self.conn.executemany(
                """INSERT INTO contract_code_cache(chain,address,has_code,checked_at,checked_block)
                   VALUES(?,?,?,?,?) ON CONFLICT(chain,address) DO UPDATE SET
                     has_code=excluded.has_code,checked_at=excluded.checked_at,
                     checked_block=excluded.checked_block""",
                cache_rows,
            )
            self.conn.executemany(
                """INSERT INTO tokens(chain,address,symbol,decimals,source) VALUES(?,?,?,?,?)
                   ON CONFLICT(chain,address) DO UPDATE SET
                     symbol=COALESCE(excluded.symbol,tokens.symbol),
                     decimals=COALESCE(excluded.decimals,tokens.decimals)""",
                token_rows,
            )
            self.conn.executemany(
                """INSERT INTO contract_tokens(
                       chain,contract,token,source,first_seen_block,last_seen_block
                   ) VALUES(?,?,?,?,?,?) ON CONFLICT(chain,contract,token) DO UPDATE SET
                     last_seen_block=MAX(contract_tokens.last_seen_block,excluded.last_seen_block)""",
                [(*row, row[4]) for row in relation_rows],
            )
            anchor = int(cursor["anchor_block"])
            status = "complete" if role == "backfill" and range_end >= anchor else "active"
            self.conn.execute(
                """UPDATE chain_cursors SET next_block=?,last_committed=?,status=?,updated_at=?,note=NULL
                   WHERE chain=? AND role=?""",
                (range_end + 1, range_end, status, now, chain, role),
            )
            if role == "backfill":
                self.conn.execute(
                    "UPDATE chain_state SET last_indexed=MAX(last_indexed,?) WHERE chain=?",
                    (range_end, chain),
                )
            self._bump_revision_locked()
            return int(inserted), int(sources_inserted)

    def save_rpc_health(self, rows: list[dict[str, Any]]) -> None:
        if not rows:
            return
        with self._lock, self.conn:
            self.conn.executemany(
                """INSERT INTO rpc_method_health(
                       chain,endpoint_fp,hostname,method_group,fail_streak,success_streak,
                       cooldown_until,circuit_until,permanent_error,latency_ewma_ms,
                       batch_limit,success_since_resize,updated_at
                   ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)
                   ON CONFLICT(chain,endpoint_fp,method_group) DO UPDATE SET
                     hostname=excluded.hostname,fail_streak=excluded.fail_streak,
                     success_streak=excluded.success_streak,cooldown_until=excluded.cooldown_until,
                     circuit_until=excluded.circuit_until,permanent_error=excluded.permanent_error,
                     latency_ewma_ms=excluded.latency_ewma_ms,batch_limit=excluded.batch_limit,
                     success_since_resize=excluded.success_since_resize,updated_at=excluded.updated_at""",
                [(
                    row["chain"], row["endpoint_fp"], row["hostname"], row["method_group"],
                    row["fail_streak"], row["success_streak"], row["cooldown_until"],
                    row["circuit_until"], row.get("permanent_error"), row.get("latency_ewma_ms"),
                    row.get("batch_limit"), row.get("success_since_resize"), row["updated_at"],
                ) for row in rows],
            )

    def load_rpc_health(self, chain: str) -> list[sqlite3.Row]:
        with self._lock:
            return list(self.conn.execute(
                "SELECT * FROM rpc_method_health WHERE chain=?", (chain,)
            ))

    def init_chain(
        self, chain: str, start_block: int, discover_tx_to_contracts: bool = False
    ) -> int:
        with self._lock:
            row = self.conn.execute(
                """
                SELECT last_indexed, start_block, discovery_revision
                FROM chain_state WHERE chain=?
                """,
                (chain,),
            ).fetchone()
            if row is None:
                revision = 2 if discover_tx_to_contracts else 1
                self.conn.execute(
                    """
                    INSERT INTO chain_state(
                        chain, start_block, last_indexed, discovery_revision
                    ) VALUES(?,?,?,?)
                    """,
                    (chain, start_block, max(start_block - 1, 0), revision),
                )
                self.conn.commit()
                return max(start_block - 1, 0)
            if discover_tx_to_contracts and int(row["discovery_revision"]) < 2:
                rewound = min(int(row["last_indexed"]), max(start_block - 1, 0))
                self.conn.execute(
                    """
                    UPDATE chain_state
                    SET start_block=?, last_indexed=?, discovery_revision=2
                    WHERE chain=?
                    """,
                    (start_block, rewound, chain),
                )
                self.conn.commit()
                log.info(
                    "[%s] tx.to discovery enabled; cursor rewound %s -> %s",
                    chain,
                    row["last_indexed"],
                    rewound,
                )
                return rewound
            return int(row["last_indexed"])

    def last_indexed(self, chain: str) -> int:
        with self._lock:
            row = self.conn.execute(
                "SELECT last_indexed FROM chain_state WHERE chain=?", (chain,)
            ).fetchone()
            return int(row["last_indexed"]) if row else 0

    def discovery_revision(self, chain: str) -> int | None:
        with self._lock:
            row = self.conn.execute(
                "SELECT discovery_revision FROM chain_state WHERE chain=?", (chain,)
            ).fetchone()
            return int(row["discovery_revision"]) if row else None

    def set_last_indexed(self, chain: str, block: int) -> None:
        with self._lock:
            self.conn.execute(
                "UPDATE chain_state SET last_indexed=? WHERE chain=?", (block, chain)
            )
            self.conn.commit()

    def upsert_contracts(self, rows: list[tuple]) -> int:
        if not rows:
            return 0
        with self._lock:
            before = self.conn.total_changes
            self.conn.executemany(
                """
                INSERT OR IGNORE INTO contracts
                    (chain, address, created_block, created_tx, creator, first_seen_at)
                VALUES (?,?,?,?,?,?)
                """,
                rows,
            )
            inserted = self.conn.total_changes - before
            self.conn.executemany(
                """
                UPDATE contracts SET
                    created_block=COALESCE(created_block, ?),
                    created_tx=COALESCE(created_tx, ?),
                    creator=COALESCE(creator, ?)
                WHERE chain=? AND address=? AND ? IS NOT NULL
                """,
                [
                    (row[2], row[3], row[4], row[0], row[1], row[2])
                    for row in rows
                ],
            )
            self.conn.commit()
            return inserted

    def upsert_contract_discoveries(self, rows: list[tuple]) -> int:
        if not rows:
            return 0
        with self._lock:
            before = self.conn.total_changes
            self.conn.executemany(
                """
                INSERT OR IGNORE INTO contract_discoveries(
                    chain, address, source, observed_block, observed_tx, actor, first_seen_at
                ) VALUES (?,?,?,?,?,?,?)
                """,
                rows,
            )
            inserted = self.conn.total_changes - before
            self.conn.commit()
            return inserted

    def discovery_addresses(self, chain: str, source: str) -> set[str]:
        with self._lock:
            return {
                row["address"].lower()
                for row in self.conn.execute(
                    "SELECT address FROM contract_discoveries WHERE chain=? AND source=?",
                    (chain, source),
                )
            }

    def cached_code_statuses(
        self, chain: str, addresses: list[str], max_age_sec: int
    ) -> dict[str, bool]:
        if not addresses:
            return {}
        cutoff = datetime.fromtimestamp(
            time.time() - max_age_sec, timezone.utc
        ).isoformat()
        found: dict[str, bool] = {}
        with self._lock:
            for offset in range(0, len(addresses), 500):
                chunk = addresses[offset : offset + 500]
                placeholders = ",".join("?" for _ in chunk)
                rows = self.conn.execute(
                    f"""
                    SELECT address, has_code FROM contract_code_cache
                    WHERE chain=? AND checked_at>=? AND address IN ({placeholders})
                    """,
                    (chain, cutoff, *chunk),
                ).fetchall()
                found.update(
                    (row["address"].lower(), bool(row["has_code"])) for row in rows
                )
        return found

    def upsert_code_cache(self, rows: list[tuple]) -> None:
        if not rows:
            return
        with self._lock:
            self.conn.executemany(
                """
                INSERT INTO contract_code_cache(
                    chain, address, has_code, checked_at, checked_block
                ) VALUES (?,?,?,?,?)
                ON CONFLICT(chain, address) DO UPDATE SET
                    has_code=excluded.has_code,
                    checked_at=excluded.checked_at,
                    checked_block=excluded.checked_block
                """,
                rows,
            )
            self.conn.commit()

    def upsert_tokens(self, rows: list[tuple]) -> None:
        if not rows:
            return
        with self._lock:
            self.conn.executemany(
                """
                INSERT INTO tokens(chain, address, symbol, decimals, source)
                VALUES (?,?,?,?,?)
                ON CONFLICT(chain, address) DO UPDATE SET
                    symbol=COALESCE(excluded.symbol, tokens.symbol),
                    decimals=COALESCE(excluded.decimals, tokens.decimals)
                """,
                rows,
            )
            self.conn.commit()

    def tokens_for(self, chain: str) -> list[sqlite3.Row]:
        with self._lock:
            return list(self.conn.execute("SELECT * FROM tokens WHERE chain=?", (chain,)))

    def contract_addresses(self, chain: str) -> set[str]:
        with self._lock:
            return {
                r["address"].lower()
                for r in self.conn.execute(
                    "SELECT address FROM contracts WHERE chain=?", (chain,)
                )
            }

    def pending_contracts(self, chain: str, recheck_sec: int, limit: int = 200) -> list[sqlite3.Row]:
        cutoff = datetime.now(timezone.utc).isoformat()
        with self._lock:
            return list(
                self.conn.execute(
                    """
                    SELECT * FROM contracts
                    WHERE chain=?
                      AND (
                        last_checked_at IS NULL
                        OR strftime('%s', last_checked_at) < strftime('%s', ?) - ?
                      )
                    ORDER BY last_checked_at IS NULL DESC, created_block DESC
                    LIMIT ?
                    """,
                    (chain, cutoff, recheck_sec, limit),
                )
            )

    def save_scan(
        self,
        chain: str,
        address: str,
        native_amount: float,
        native_usd: float,
        token_rows: list[dict[str, Any]],
        total_usd: float,
        min_usd: float,
    ) -> None:
        tokens_usd = sum(t["usd_value"] for t in token_rows)
        meets = 1 if total_usd >= min_usd else 0
        if total_usd <= 0:
            status = "zero"
            note = "нулевой баланс"
        elif meets:
            status = "qualifying"
            note = f"{total_usd:,.2f}$ >= {min_usd:,.0f}$"
        else:
            status = "below"
            note = f"{total_usd:,.2f}$ (меньше условия {min_usd:,.0f}$)"
        now = datetime.now(timezone.utc).isoformat()
        with self._lock:
            cur = self.conn.execute(
                """
                INSERT INTO scans(chain, address, scanned_at, native_amount, native_usd,
                                  tokens_usd, total_usd, meets_threshold, note)
                VALUES (?,?,?,?,?,?,?,?,?)
                """,
                (chain, address, now, native_amount, native_usd, tokens_usd, total_usd, meets, note),
            )
            scan_id = cur.lastrowid
            self.conn.executemany(
                """
                INSERT INTO scan_tokens(scan_id, token, symbol, amount, price_usd, usd_value)
                VALUES (?,?,?,?,?,?)
                """,
                [
                    (scan_id, t["token"], t["symbol"], t["amount"], t["price_usd"], t["usd_value"])
                    for t in token_rows
                ],
            )
            self.conn.execute(
                """
                UPDATE contracts
                SET last_checked_at=?, last_total_usd=?, last_status=?
                WHERE chain=? AND address=?
                """,
                (now, total_usd, status, chain, address),
            )
            self.conn.commit()

    def latest_nonzero_scans(self) -> list[sqlite3.Row]:
        with self._lock:
            return list(
                self.conn.execute(
                    """
                    SELECT c.chain, c.address, c.created_block, c.created_tx, c.creator,
                           c.last_status, s.scanned_at, s.native_amount, s.native_usd,
                           s.tokens_usd, s.total_usd, s.note, s.id AS scan_id
                    FROM contracts c
                    JOIN scans s ON s.id = (
                        SELECT id FROM scans
                        WHERE chain=c.chain AND address=c.address
                        ORDER BY scanned_at DESC LIMIT 1
                    )
                    WHERE s.total_usd > 0
                    ORDER BY s.total_usd DESC
                    """
                )
            )

    def tokens_of_scan(self, scan_id: int) -> list[sqlite3.Row]:
        with self._lock:
            return list(
                self.conn.execute(
                    "SELECT * FROM scan_tokens WHERE scan_id=? AND usd_value > 0 ORDER BY usd_value DESC",
                    (scan_id,),
                )
            )

    def chain_counts(self, chain: str) -> dict[str, int]:
        with self._lock:
            total = self.conn.execute(
                "SELECT COUNT(*) FROM contracts WHERE chain=?", (chain,)
            ).fetchone()[0]
            unchecked = self.conn.execute(
                "SELECT COUNT(*) FROM contracts WHERE chain=? AND last_checked_at IS NULL",
                (chain,),
            ).fetchone()[0]
            qual = self.conn.execute(
                "SELECT COUNT(*) FROM contracts WHERE chain=? AND last_status='qualifying'",
                (chain,),
            ).fetchone()[0]
            below = self.conn.execute(
                "SELECT COUNT(*) FROM contracts WHERE chain=? AND last_status='below'",
                (chain,),
            ).fetchone()[0]
        return {
            "contracts": int(total),
            "unchecked": int(unchecked),
            "qualifying": int(qual),
            "below": int(below),
        }

    def upsert_contract_tokens(self, rows: list[tuple[str, str, str, str, int]]) -> None:
        if not rows:
            return
        with self._lock:
            self.conn.executemany(
                """
                INSERT INTO contract_tokens(
                    chain, contract, token, source, first_seen_block, last_seen_block
                ) VALUES (?,?,?,?,?,?)
                ON CONFLICT(chain, contract, token) DO UPDATE SET
                    last_seen_block=CASE
                        WHEN excluded.last_seen_block > contract_tokens.last_seen_block
                        THEN excluded.last_seen_block ELSE contract_tokens.last_seen_block END
                """,
                [(*row, row[4]) for row in rows],
            )
            self.conn.commit()

    def tokens_for_contract(self, chain: str, contract: str) -> list[sqlite3.Row]:
        with self._lock:
            return list(
                self.conn.execute(
                    """
                    SELECT DISTINCT t.*
                    FROM tokens t
                    WHERE t.chain=?
                      AND (
                        t.source='seed'
                        OR EXISTS (
                            SELECT 1 FROM contract_tokens ct
                            WHERE ct.chain=t.chain
                              AND ct.token=t.address
                              AND ct.contract=?
                        )
                      )
                    ORDER BY t.address
                    """,
                    (chain, contract.lower()),
                )
            )

    def pending_addresses(
        self, recheck_sec: int, limit: int = 200, retry_sec: int | None = None,
        as_of: float | None = None,
    ) -> list[str]:
        now_iso = datetime.fromtimestamp(
            time.time() if as_of is None else as_of, timezone.utc
        ).isoformat()
        with self._lock:
            rows = self.conn.execute(
                """
                SELECT lower(c.address) AS address
                FROM contracts c
                GROUP BY lower(c.address)
                HAVING NOT EXISTS (
                    SELECT 1 FROM address_chain_state s
                    WHERE s.address=lower(c.address)
                ) OR EXISTS (
                    SELECT 1 FROM address_chain_state s
                    WHERE s.address=lower(c.address) AND s.next_retry_at<=?
                )
                ORDER BY MIN(c.last_checked_at IS NOT NULL), MIN(COALESCE(c.last_checked_at,'')),
                         lower(c.address)
                LIMIT ?
                """,
                (now_iso, limit),
            ).fetchall()
            return [row["address"] for row in rows]

    def due_address_chains(
        self, address: str, chain_keys: list[str], as_of: float | None = None,
    ) -> list[str]:
        now_iso = datetime.fromtimestamp(
            time.time() if as_of is None else as_of, timezone.utc
        ).isoformat()
        with self._lock:
            rows = {
                str(row["chain"]): row
                for row in self.conn.execute(
                    "SELECT chain,next_retry_at FROM address_chain_state WHERE address=?",
                    (address.lower(),),
                )
            }
        return [
            chain for chain in chain_keys
            if chain not in rows or str(rows[chain]["next_retry_at"]) <= now_iso
        ]

    def save_address_chain_state(
        self, address: str, row: dict[str, Any], token_rows: list[dict[str, Any]],
        min_usd: float,
    ) -> None:
        address = address.lower()
        chain = str(row["chain"])
        now = datetime.now(timezone.utc)
        now_iso = now.isoformat()
        status = str(row.get("status") or "rpc_error")
        with self._lock, self.conn:
            previous = self.conn.execute(
                "SELECT * FROM address_chain_state WHERE address=? AND chain=?",
                (address, chain),
            ).fetchone()
            failed = status in {"rpc_error", "partial", "timeout"}
            failure_streak = (int(previous["failure_streak"]) if previous else 0) + 1 if failed else 0
            if failed:
                delay = min(21_600, 600 * (2 ** min(max(0, failure_streak - 1), 6)))
            elif status in {"price_missing", "anomalous_balance"}:
                delay = 3_600
            elif status == "complete" and float(row.get("total_usd") or 0) >= min_usd:
                delay = 21_600
            else:
                delay = 86_400
            next_retry = datetime.fromtimestamp(now.timestamp() + delay, timezone.utc).isoformat()
            has_observation = row.get("has_code") is not None and row.get("total_usd") is not None
            def observed(name: str, fallback: Any = None) -> Any:
                value = row.get(name)
                if value is None and previous is not None:
                    return previous[name]
                return fallback if value is None else value
            self.conn.execute(
                """INSERT INTO address_chain_state(
                       address,chain,checked_at,last_success_at,status,has_code,native_raw,
                       native_amount,observed_native_usd,included_native_usd,excluded_usd,
                       tokens_usd,total_usd,valuation_status,failure_streak,next_retry_at,note
                   ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                   ON CONFLICT(address,chain) DO UPDATE SET
                     checked_at=excluded.checked_at,last_success_at=excluded.last_success_at,
                     status=excluded.status,has_code=excluded.has_code,native_raw=excluded.native_raw,
                     native_amount=excluded.native_amount,
                     observed_native_usd=excluded.observed_native_usd,
                     included_native_usd=excluded.included_native_usd,
                     excluded_usd=excluded.excluded_usd,tokens_usd=excluded.tokens_usd,
                     total_usd=excluded.total_usd,valuation_status=excluded.valuation_status,
                     failure_streak=excluded.failure_streak,next_retry_at=excluded.next_retry_at,
                     note=excluded.note""",
                (
                    address, chain, now_iso,
                    now_iso if has_observation else (previous["last_success_at"] if previous else None),
                    status, observed("has_code"), observed("native_raw"),
                    # A failed chain has no observation.  Keep its amounts NULL
                    # (or retain a previous successful value), never silently
                    # turn an RPC failure into a zero balance.
                    observed("native_amount"), observed("observed_native_usd"),
                    observed("included_native_usd"), observed("excluded_usd", 0.0),
                    observed("tokens_usd"), observed("total_usd"),
                    observed("valuation_status"), failure_streak, next_retry, row.get("note"),
                ),
            )
            if has_observation:
                if status in {"complete", "price_missing", "anomalous_balance", "absent"}:
                    self.conn.execute(
                        "DELETE FROM address_token_state WHERE address=? AND chain=?",
                        (address, chain),
                    )
                else:
                    checked_tokens = list(row.get("checked_tokens") or [])
                    for offset in range(0, len(checked_tokens), 500):
                        chunk = checked_tokens[offset:offset + 500]
                        placeholders = ",".join("?" for _ in chunk)
                        if chunk:
                            self.conn.execute(
                                f"DELETE FROM address_token_state WHERE address=? AND chain=? AND token IN ({placeholders})",
                                (address, chain, *chunk),
                            )
                self.conn.executemany(
                    """INSERT INTO address_token_state(
                           address,chain,token,checked_at,raw_amount,amount,symbol,
                           price_usd,usd_value,priced,valuation_status
                       ) VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                    [(
                        address, chain, token["token"], now_iso,
                        str(token.get("raw_amount", "0")), token.get("amount"),
                        token.get("symbol"), token.get("price_usd"), token.get("usd_value"),
                        int(bool(token.get("priced"))),
                        token.get("valuation_status", "included"),
                    ) for token in token_rows],
                )
                if status == "partial":
                    token_total = float(self.conn.execute(
                        """SELECT COALESCE(SUM(usd_value),0) FROM address_token_state
                           WHERE address=? AND chain=? AND token!='native'
                             AND valuation_status='included'""",
                        (address, chain),
                    ).fetchone()[0])
                    included_native = float(observed("included_native_usd", 0.0) or 0.0)
                    self.conn.execute(
                        """UPDATE address_chain_state SET tokens_usd=?,total_usd=?
                           WHERE address=? AND chain=?""",
                        (token_total, token_total + included_native, address, chain),
                    )
            self._bump_revision_locked()

    def current_address_parts(
        self, address: str, chain_keys: list[str],
    ) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        address = address.lower()
        with self._lock:
            state = {
                str(row["chain"]): dict(row)
                for row in self.conn.execute(
                    "SELECT * FROM address_chain_state WHERE address=?", (address,)
                )
            }
            tokens = [
                dict(row) for row in self.conn.execute(
                    "SELECT * FROM address_token_state WHERE address=? ORDER BY chain,token",
                    (address,),
                )
            ]
        chain_rows: list[dict[str, Any]] = []
        for chain in chain_keys:
            row = state.get(chain)
            if row is None:
                chain_rows.append({
                    "chain": chain, "has_code": None, "status": "rpc_error",
                    "total_usd": None, "note": "not_checked",
                })
            else:
                row.setdefault("native_usd", row.get("observed_native_usd"))
                chain_rows.append(row)
        return chain_rows, tokens

    def apply_address_schedule(self, address: str, status: str) -> None:
        if status != "qualifying":
            return
        target = datetime.fromtimestamp(time.time() + 21_600, timezone.utc).isoformat()
        with self._lock, self.conn:
            self.conn.execute(
                """UPDATE address_chain_state SET next_retry_at=CASE
                     WHEN next_retry_at>? THEN ? ELSE next_retry_at END
                   WHERE address=? AND status NOT IN ('rpc_error','partial','price_missing','anomalous_balance')""",
                (target, target, address.lower()),
            )

    def ingest_indexed_tokens(
        self, address: str, rows: list[dict[str, Any]], source: str,
    ) -> int:
        token_rows: list[tuple[str, str, str | None, int | None, str]] = []
        relations: list[tuple[str, str, str, str, int]] = []
        for row in rows:
            token = normalize_evm_address(row.get("token"))
            chain = str(row.get("chain") or "")
            if token is None or not chain or token == ZERO:
                continue
            decimals = row.get("decimals")
            try:
                decimals = int(decimals) if decimals is not None else None
            except (TypeError, ValueError):
                decimals = None
            token_rows.append((chain, token, row.get("symbol"), decimals, source))
            relations.append((chain, address.lower(), token, source, 0))
        self.upsert_tokens(token_rows)
        self.upsert_contract_tokens(relations)
        return len(relations)

    def save_rabby_estimate(self, address: str, estimate: dict[str, Any], scan_id: int) -> None:
        with self._lock:
            self.conn.execute(
                """INSERT INTO rabby_estimates(
                     address, checked_at, status, estimated_usd, lower_usd, upper_usd,
                     coverage, total_networks, uncertainty, chain_parts, token_parts, note, rpc_scan_id
                   ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
                   ON CONFLICT(address) DO UPDATE SET
                     checked_at=excluded.checked_at, status=excluded.status,
                     estimated_usd=excluded.estimated_usd, lower_usd=excluded.lower_usd,
                     upper_usd=excluded.upper_usd, coverage=excluded.coverage,
                     total_networks=excluded.total_networks, uncertainty=excluded.uncertainty,
                     chain_parts=excluded.chain_parts, token_parts=excluded.token_parts,
                     note=excluded.note, rpc_scan_id=excluded.rpc_scan_id""",
                (address.lower(), datetime.now(timezone.utc).isoformat(), estimate["status"],
                 estimate["estimated_usd"], estimate["lower_usd"], estimate["upper_usd"],
                 estimate["coverage"], estimate["total_networks"], estimate["uncertainty"],
                 json.dumps(estimate["chains"], ensure_ascii=False),
                 json.dumps(estimate["tokens"], ensure_ascii=False), estimate.get("note"), scan_id),
            )
            self.conn.commit()

    def pending_rabby_scan(
        self, retry_sec: int, as_of: float | None = None,
        minimum_age_sec: int = 900,
    ) -> sqlite3.Row | None:
        now = time.time() if as_of is None else as_of
        cutoff = datetime.fromtimestamp(now - retry_sec, timezone.utc).isoformat()
        old_enough = datetime.fromtimestamp(
            time.time() - minimum_age_sec, timezone.utc
        ).isoformat()
        with self._lock:
            return self.conn.execute(
                """SELECT a.* FROM address_scans a
                   LEFT JOIN rabby_estimates r ON r.address=a.address
                   WHERE a.id=(SELECT MAX(a2.id) FROM address_scans a2 WHERE a2.address=a.address)
                     AND EXISTS (SELECT 1 FROM address_chain_scans c WHERE c.scan_id=a.id
                                 AND c.status NOT IN ('complete','absent'))
                     AND a.scanned_at <= ?
                     AND (r.address IS NULL OR r.checked_at < ?)
                   ORDER BY COALESCE(r.checked_at, '') ASC, a.id LIMIT 1""", (old_enough, cutoff)
            ).fetchone()

    def rabby_estimate(self, address: str) -> sqlite3.Row | None:
        with self._lock:
            return self.conn.execute(
                "SELECT * FROM rabby_estimates WHERE address=?", (address.lower(),)
            ).fetchone()

    def save_address_scan(
        self,
        address: str,
        status: str,
        total_usd: float,
        coverage: int,
        total_networks: int,
        chain_rows: list[dict[str, Any]],
        token_rows: list[dict[str, Any]],
        note: str | None = None,
    ) -> int:
        now = datetime.now(timezone.utc).isoformat()
        address = address.lower()
        with self._lock:
            cur = self.conn.execute(
                """
                INSERT INTO address_scans(
                    address, scanned_at, status, total_usd, coverage, total_networks, note
                ) VALUES (?,?,?,?,?,?,?)
                """,
                (address, now, status, total_usd, coverage, total_networks, note),
            )
            scan_id = int(cur.lastrowid)
            self.conn.executemany(
                """
                INSERT INTO address_chain_scans(
                    scan_id, chain, has_code, status, native_amount, native_usd,
                    tokens_usd, total_usd, note, native_raw, observed_native_usd,
                    included_native_usd, excluded_usd, valuation_status
                ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                [
                    (
                        scan_id,
                        row["chain"],
                        row.get("has_code"),
                        row["status"],
                        row.get("native_amount"),
                        row.get("native_usd"),
                        row.get("tokens_usd"),
                        row.get("total_usd"),
                        row.get("note"),
                        row.get("native_raw"),
                        row.get("observed_native_usd", row.get("native_usd")),
                        row.get("included_native_usd", row.get("native_usd")),
                        row.get("excluded_usd", 0.0),
                        row.get("valuation_status"),
                    )
                    for row in chain_rows
                ],
            )
            self.conn.executemany(
                """
                INSERT INTO address_token_scans(
                    scan_id, chain, token, symbol, raw_amount, amount,
                    price_usd, usd_value, priced, valuation_status
                ) VALUES (?,?,?,?,?,?,?,?,?,?)
                """,
                [
                    (
                        scan_id,
                        row["chain"],
                        row["token"],
                        row.get("symbol"),
                        str(row.get("raw_amount", "0")),
                        row.get("amount"),
                        row.get("price_usd"),
                        row.get("usd_value"),
                        1 if row.get("priced") else 0,
                        row.get("valuation_status", "included"),
                    )
                    for row in token_rows
                ],
            )
            self.conn.execute(
                """
                UPDATE contracts
                SET last_checked_at=?, last_total_usd=?, last_status=?
                WHERE lower(address)=?
                """,
                (now, total_usd, status, address),
            )
            self._bump_revision_locked()
            self.conn.commit()
            return scan_id

    def latest_address_scans(self) -> list[sqlite3.Row]:
        with self._lock:
            return list(
                self.conn.execute(
                    """
                    SELECT a.*
                    FROM address_scans a
                    WHERE a.id=(
                        SELECT a2.id FROM address_scans a2
                        WHERE a2.address=a.address
                        ORDER BY a2.scanned_at DESC, a2.id DESC LIMIT 1
                    )
                    ORDER BY a.total_usd DESC, a.address
                    """
                )
            )

    def address_scan_chains(self, scan_id: int) -> list[sqlite3.Row]:
        with self._lock:
            return list(
                self.conn.execute(
                    "SELECT * FROM address_chain_scans WHERE scan_id=? ORDER BY chain",
                    (scan_id,),
                )
            )

    def address_scan_tokens(self, scan_id: int) -> list[sqlite3.Row]:
        with self._lock:
            return list(
                self.conn.execute(
                    """
                    SELECT * FROM address_token_scans
                    WHERE scan_id=? AND CAST(raw_amount AS INTEGER) > 0
                    ORDER BY chain, COALESCE(usd_value, -1) DESC, token
                    """,
                    (scan_id,),
                )
            )

    def address_sources(self, address: str) -> list[sqlite3.Row]:
        with self._lock:
            return list(
                self.conn.execute(
                    """
                    SELECT chain, source, observed_block, observed_tx, actor, first_seen_at
                    FROM contract_discoveries WHERE lower(address)=?
                    ORDER BY first_seen_at, chain, source
                    """,
                    (address.lower(),),
                )
            )

    def aggregate_counts(self) -> dict[str, int]:
        with self._lock:
            rows = self.conn.execute(
                """
                SELECT status, COUNT(*) AS n FROM address_scans a
                WHERE a.id=(
                    SELECT a2.id FROM address_scans a2
                    WHERE a2.address=a.address
                    ORDER BY a2.scanned_at DESC, a2.id DESC LIMIT 1
                ) GROUP BY status
                """
            ).fetchall()
        return {row["status"]: int(row["n"]) for row in rows}

    def monitoring_snapshot(self, min_usd: float) -> dict[str, Any]:
        """Return one cheap aggregate snapshot for monitoring.db."""
        with self._lock:
            unique_addresses = int(self.conn.execute(
                "SELECT COUNT(DISTINCT lower(address)) FROM contracts"
            ).fetchone()[0])
            contract_instances = int(self.conn.execute(
                "SELECT COUNT(*) FROM contracts"
            ).fetchone()[0])
            sources = {
                str(row["source"]): int(row["n"])
                for row in self.conn.execute(
                    "SELECT source,COUNT(*) AS n FROM contract_discoveries GROUP BY source"
                )
            }
            pending = int(self.conn.execute(
                """SELECT COUNT(DISTINCT lower(c.address)) FROM contracts c
                   WHERE NOT EXISTS(SELECT 1 FROM address_chain_state s WHERE s.address=lower(c.address))
                      OR EXISTS(SELECT 1 FROM address_chain_state s WHERE s.address=lower(c.address)
                                AND strftime('%s',s.next_retry_at)<=strftime('%s','now'))"""
            ).fetchone()[0])
            oldest_age = self.conn.execute(
                """SELECT MAX(strftime('%s','now')-strftime('%s',c.first_seen_at))
                   FROM contracts c WHERE NOT EXISTS(
                     SELECT 1 FROM address_chain_state s WHERE s.address=lower(c.address))"""
            ).fetchone()[0]
            completed = int(self.conn.execute(
                "SELECT COUNT(DISTINCT address) FROM address_scans"
            ).fetchone()[0])
            latest = list(self.conn.execute(
                """SELECT a.* FROM address_scans a
                   WHERE a.id=(SELECT a2.id FROM address_scans a2 WHERE a2.address=a.address
                               ORDER BY a2.scanned_at DESC,a2.id DESC LIMIT 1)"""
            ))
            statuses = {
                "qualifying": sum(float(row["total_usd"] or 0) >= min_usd for row in latest),
                "below": sum(
                    row["status"] == "below" and 0 < float(row["total_usd"] or 0) < min_usd
                    for row in latest
                ),
                "incomplete": sum(
                    row["status"] == "incomplete" and float(row["total_usd"] or 0) < min_usd
                    for row in latest
                ),
            }
            coverage = {
                f"{int(row['coverage'])}/{int(row['total_networks'])}": int(row["n"])
                for row in self.conn.execute(
                    """
                    SELECT coverage,total_networks,COUNT(*) AS n FROM address_scans a
                    WHERE a.id=(SELECT a2.id FROM address_scans a2 WHERE a2.address=a.address
                                ORDER BY a2.scanned_at DESC,a2.id DESC LIMIT 1)
                    GROUP BY coverage,total_networks
                    """
                )
            }
        return {
            "unique_addresses": unique_addresses,
            "contract_instances": contract_instances,
            "direct_deploy": sources.get("direct_deploy", 0),
            "active_call": sources.get("active_call", 0),
            "balance_pending": pending,
            "balance_oldest_age_sec": float(oldest_age) if oldest_age is not None else None,
            "balance_completed": completed,
            "qualifying": statuses.get("qualifying", 0),
            "below_count": statuses.get("below", 0),
            "incomplete": statuses.get("incomplete", 0),
            "coverage_json": coverage,
        }

    def discovery_count(self, chain: str, source: str) -> int:
        with self._lock:
            return int(self.conn.execute(
                "SELECT COUNT(*) FROM contract_discoveries WHERE chain=? AND source=?",
                (chain, source),
            ).fetchone()[0])


# ---------------------------------------------------------------------------
# RPC with cooldowns
# ---------------------------------------------------------------------------


class RpcError(Exception):
    def __init__(
        self,
        kind: str,
        message: str,
        retry_after: float | None = None,
        permanent: bool = False,
    ):
        super().__init__(message)
        self.kind = kind
        self.retry_after = retry_after
        self.permanent = permanent


def _short_url(url: str) -> str:
    p = urlparse(url)
    return p.netloc or url[:40]


def _parse_retry_after(headers: httpx.Headers) -> float | None:
    raw = headers.get("retry-after")
    if not raw:
        return None
    try:
        return max(1.0, float(raw))
    except ValueError:
        return None


def classify_rpc_problem(status: int | None, body: str, headers: httpx.Headers | None = None) -> RpcError | None:
    text = (body or "").lower()
    retry_after = _parse_retry_after(headers) if headers is not None else None
    if status in (401, 403) or any(h in text for h in AUTH_HINTS):
        return RpcError("auth", "authentication rejected", permanent=True)
    if status in (429,) or any(h in text for h in RATE_HINTS):
        return RpcError("rate_limit", "RPC rate limit", retry_after)
    if any(h in text for h in RANGE_HINTS):
        return RpcError("range", "RPC range limit", None)
    if status in RETRYABLE_HTTP:
        return RpcError("http", f"HTTP {status}", retry_after or 10)
    if status and status >= 400:
        return RpcError("http", f"HTTP {status}", retry_after)
    return None


@dataclass
class MethodHealth:
    cooldown_until: float = 0.0
    circuit_until: float = 0.0
    fail_streak: int = 0
    ok_streak: int = 0
    permanent_error: str | None = None
    last_error: str | None = None
    latency_ewma_ms: float | None = None
    success_since_resize: int = 0

    def available(self) -> bool:
        now = time.time()
        return self.permanent_error is None and now >= self.cooldown_until and now >= self.circuit_until


@dataclass
class Endpoint:
    url: str
    cooldown_until: float = 0.0
    fail_streak: int = 0
    ok_streak: int = 0
    permanent_error: str | None = None
    batch_supported: bool = True
    last_error: str | None = None
    chain_verified: bool = False
    priority: int = 0
    private: bool = False
    method_health: dict[str, MethodHealth] = field(default_factory=dict)
    next_request_at: float = 0.0
    rate_lock: asyncio.Lock = field(default_factory=asyncio.Lock)

    @property
    def fingerprint(self) -> str:
        return hashlib.sha256(self.url.encode("utf-8")).hexdigest()[:16]

    def health(self, method_group: str) -> MethodHealth:
        return self.method_health.setdefault(method_group, MethodHealth())

    def available(self, method_group: str = "head") -> bool:
        return (
            self.permanent_error is None
            and time.time() >= self.cooldown_until
            and self.health(method_group).available()
        )

    def cool(self, seconds: float, cap: float, method_group: str = "head") -> float:
        state = self.health(method_group)
        wait = min(cap, max(1.0, seconds))
        state.fail_streak += 1
        state.ok_streak = 0
        self.fail_streak = max(self.fail_streak, state.fail_streak)
        self.ok_streak = 0
        # экспонента от серии фейлов
        wait = min(cap, wait * (2 ** min(state.fail_streak - 1, 6)))
        wait *= random.uniform(0.85, 1.15)
        state.cooldown_until = time.time() + wait
        if state.fail_streak >= 5:
            state.circuit_until = time.time() + 300
        state.last_error = self.last_error
        return wait

    def ok(self, method_group: str = "head", latency_ms: float | None = None) -> None:
        state = self.health(method_group)
        state.fail_streak = 0
        state.ok_streak += 1
        state.success_since_resize += 1
        state.cooldown_until = 0.0
        state.circuit_until = 0.0
        state.last_error = None
        if latency_ms is not None:
            state.latency_ewma_ms = (
                latency_ms if state.latency_ewma_ms is None
                else (state.latency_ewma_ms * 0.8 + latency_ms * 0.2)
            )
        self.fail_streak = 0
        self.ok_streak += 1
        # cooldown_until is the legacy/global gate; successful calls must not
        # clear a manually imposed global pause for unrelated methods.
        self.last_error = None

    def disable(self, reason: str) -> None:
        self.permanent_error = reason
        self.last_error = reason
        for state in self.method_health.values():
            state.permanent_error = reason
            state.last_error = reason


def rpc_method_group(method: str) -> str:
    if method in {"eth_chainId", "eth_blockNumber"}:
        return "head"
    if method == "eth_getBlockByNumber":
        return "blocks"
    if method in {"eth_getTransactionReceipt", "eth_getBlockReceipts"}:
        return "receipts"
    if method == "eth_getLogs":
        return "logs"
    if method == "eth_getCode":
        return "code"
    if method == "eth_getBalance":
        return "balance"
    if method == "eth_call":
        return "eth_call"
    return "head"


class RpcPool:
    def __init__(
        self,
        chain: str,
        urls: list[str],
        cfg: AppCfg,
        expected_chain_id: int | None = None,
        global_sem: asyncio.Semaphore | None = None,
        transport: httpx.AsyncBaseTransport | None = None,
        health_store: DB | None = None,
    ):
        if not urls:
            raise ValueError(f"{chain}: нет RPC URL")
        self.chain = chain
        self.cfg = cfg
        self.expected_chain_id = expected_chain_id
        private_hosts = {"lb.drpc.live", "alchemy.com", "g.alchemy.com"}
        self.endpoints = [
            Endpoint(
                u, priority=index,
                private=(urlparse(u).hostname or "").endswith(tuple(private_hosts)),
            )
            for index, u in enumerate(urls)
        ]
        self.timeout = cfg.http_timeout_sec
        self.retries = cfg.max_retries
        self.sem = asyncio.Semaphore(cfg.rpc_concurrency)
        self.global_sem = global_sem or asyncio.Semaphore(cfg.global_rpc_concurrency)
        self.transport = transport
        self.health_store = health_store
        self.preflight_complete = False
        self._i = 0
        self._client: httpx.AsyncClient | None = None
        self.block_batch = cfg.block_batch_size
        self.receipt_batch = cfg.receipt_batch_size
        self.call_batch = max(1, cfg.eth_call_batch_size)
        if self.endpoints and (urlparse(self.endpoints[0].url).hostname or "") == "lb.drpc.live":
            # dRPC free endpoints reject JSON-RPC batches larger than three.
            self.block_batch = min(self.block_batch, 3)
            self.receipt_batch = min(self.receipt_batch, 3)
            self.call_batch = min(self.call_batch, 3)
        self.paused_until = 0.0
        self.metric_requests = 0
        self.metric_successes = 0
        self.metric_errors: dict[str, int] = {}
        self.metric_latencies_ms: list[float] = []
        self.last_head: int | None = None
        self._batch_successes = {"block": 0, "receipt": 0, "call": 0}
        if health_store is not None:
            persisted = health_store.load_rpc_health(chain)
            by_fp = {endpoint.fingerprint: endpoint for endpoint in self.endpoints}
            now = time.time()
            for row in persisted:
                endpoint = by_fp.get(str(row["endpoint_fp"]))
                if endpoint is None:
                    continue
                state = endpoint.health(str(row["method_group"]))
                state.fail_streak = int(row["fail_streak"])
                state.ok_streak = int(row["success_streak"])
                state.cooldown_until = max(now, float(row["cooldown_until"] or 0)) if float(row["cooldown_until"] or 0) > now else 0.0
                state.circuit_until = max(now, float(row["circuit_until"] or 0)) if float(row["circuit_until"] or 0) > now else 0.0
                # auth/wrong-chain disablement lasts only for the process; cooldowns
                # and learned limits survive restarts.
                state.permanent_error = None
                state.latency_ewma_ms = row["latency_ewma_ms"]
                state.success_since_resize = int(row["success_since_resize"] or 0)
                limit = row["batch_limit"]
                if limit is not None:
                    if row["method_group"] == "blocks":
                        self.block_batch = max(1, min(self.block_batch, int(limit)))
                    elif row["method_group"] == "receipts":
                        self.receipt_batch = max(1, min(self.receipt_batch, int(limit)))
                    elif row["method_group"] in {"code", "eth_call"}:
                        self.call_batch = max(1, min(self.call_batch, int(limit)))

    async def __aenter__(self) -> "RpcPool":
        self._client = httpx.AsyncClient(
            timeout=httpx.Timeout(self.timeout, connect=15.0),
            transport=self.transport,
        )
        return self

    async def __aexit__(self, *a: Any) -> None:
        if self._client:
            await self._client.aclose()

    def shrink_batch(self, kind: str = "block") -> None:
        if not self.cfg.adaptive_batch:
            return
        attr = {"block": "block_batch", "receipt": "receipt_batch", "call": "call_batch"}.get(
            kind, "block_batch"
        )
        old = getattr(self, attr)
        new = max(1, old // 2)
        setattr(self, attr, new)
        self._batch_successes[kind] = 0
        if new != old:
            log.warning("[%s] %s batch %s -> %s", self.chain, kind, old, new)

    def grow_batch(self, kind: str = "block") -> None:
        if not self.cfg.adaptive_batch:
            return
        self._batch_successes[kind] = self._batch_successes.get(kind, 0) + 1
        if self._batch_successes[kind] < 100:
            return
        self._batch_successes[kind] = 0
        attr, limit = {
            "block": ("block_batch", self.cfg.block_batch_size),
            "receipt": ("receipt_batch", self.cfg.receipt_batch_size),
            "call": ("call_batch", self.cfg.eth_call_batch_size),
        }.get(kind, ("block_batch", self.cfg.block_batch_size))
        if getattr(self, attr) < limit:
            setattr(self, attr, min(limit, getattr(self, attr) + 1))

    def _pick(self, method_group: str = "head", exclude: set[str] | None = None) -> Endpoint | None:
        exclude = exclude or set()
        candidates = [
            ep for ep in self.endpoints
            if ep.fingerprint not in exclude
            and ep.available(method_group)
            and (not self.preflight_complete or ep.chain_verified)
        ]
        if not candidates:
            return None
        def score(ep: Endpoint) -> tuple[float, int]:
            state = ep.health(method_group)
            latency = state.latency_ewma_ms if state.latency_ewma_ms is not None else 500.0
            return (ep.priority * 1000 + state.fail_streak * 5000 + latency, ep.priority)
        return min(candidates, key=score)

    def _soonest_wait(self, method_group: str = "head") -> float:
        now = time.time()
        waits = [
            max(0.0, max(ep.health(method_group).cooldown_until,
                         ep.health(method_group).circuit_until) - now)
            for ep in self.endpoints
            if ep.permanent_error is None
            and (not self.preflight_complete or ep.chain_verified)
        ]
        return min(waits) if waits else self.cfg.all_down_sleep_sec

    async def _wait_healthy_endpoint(
        self, method_group: str = "head", exclude: set[str] | None = None,
    ) -> Endpoint:
        while True:
            ep = self._pick(method_group, exclude)
            if ep is not None:
                return ep
            if BALANCE_RPC.get():
                raise RpcError("cooldown", "no ready endpoint for balance probe")
            usable = [
                item for item in self.endpoints
                if item.permanent_error is None
                and (not self.preflight_complete or item.chain_verified)
                and item.fingerprint not in (exclude or set())
            ]
            if not usable:
                reasons = ", ".join(
                    f"{_short_url(item.url)}={item.permanent_error}" for item in self.endpoints
                )
                raise RpcError(
                    "permanent",
                    f"{self.chain}: all RPC endpoints disabled ({reasons})",
                    permanent=True,
                )
            wait = min(
                self.cfg.cooldown_max_sec,
                max(self.cfg.all_down_sleep_sec, self._soonest_wait(method_group)),
            )
            self.paused_until = time.time() + wait
            log.warning("[%s] all RPC endpoints cooling down for %.0fs", self.chain, wait)
            await asyncio.sleep(wait)

    async def _request_endpoint(
        self, ep: Endpoint, payload: Any, method_group: str = "head",
    ) -> Any:
        assert self._client
        started = time.monotonic()
        self.metric_requests += 1
        try:
            try:
                interval = 0.2 if ep.private else 1.0
                async with ep.rate_lock:
                    delay = max(0.0, ep.next_request_at - time.monotonic())
                    if delay:
                        await asyncio.sleep(delay)
                    ep.next_request_at = time.monotonic() + interval
                # Queued work for one chain must not reserve global slots.
                async with self.sem:
                    async with self.global_sem:
                        if BALANCE_RPC.get() and not ep.available(method_group):
                            raise RpcError("cooldown", "endpoint became unavailable while queued")
                        response = await self._client.post(ep.url, json=payload)
            except (httpx.TimeoutException, httpx.NetworkError, httpx.RemoteProtocolError) as exc:
                raise RpcError("network", type(exc).__name__) from exc

            body = response.text[:1200]
            if response.status_code >= 400:
                raise classify_rpc_problem(response.status_code, body, response.headers) or RpcError(
                    "http", f"HTTP {response.status_code}"
                )
            try:
                data = response.json()
            except (ValueError, json.JSONDecodeError) as exc:
                raise RpcError("malformed", "invalid JSON response") from exc
            if isinstance(data, dict) and data.get("error") is not None:
                error_text = json.dumps(data["error"], ensure_ascii=False)
                raise classify_rpc_problem(None, error_text) or RpcError("rpc", "JSON-RPC error")
            if not isinstance(data, (dict, list)):
                raise RpcError("malformed", f"unexpected JSON type: {type(data).__name__}")
            latency_ms = (time.monotonic() - started) * 1000
            ep.ok(method_group, latency_ms)
            self.metric_successes += 1
            return data
        except Exception as exc:
            kind = exc.kind if isinstance(exc, RpcError) else type(exc).__name__
            self.metric_errors[kind] = self.metric_errors.get(kind, 0) + 1
            raise
        finally:
            self.metric_latencies_ms.append((time.monotonic() - started) * 1000)
            if len(self.metric_latencies_ms) > 5000:
                del self.metric_latencies_ms[:-2500]

    async def _post(
        self, payload: Any, method_group: str | None = None,
        exclude: set[str] | None = None,
    ) -> tuple[Endpoint, Any]:
        last: Exception | None = None
        if method_group is None:
            method = payload[0].get("method", "") if isinstance(payload, list) and payload else payload.get("method", "")
            method_group = rpc_method_group(str(method))
        attempts = min(2, self.retries) if BALANCE_RPC.get() else self.retries
        for _attempt in range(max(1, attempts)):
            ep = await self._wait_healthy_endpoint(method_group, exclude)
            try:
                return ep, await self._request_endpoint(ep, payload, method_group)
            except RpcError as exc:
                last = exc
                if exc.kind == "cooldown":
                    continue
                ep.last_error = exc.kind
                if exc.permanent or exc.kind == "auth":
                    ep.disable(exc.kind)
                    log.warning(
                        "[%s] %s disabled for this run (%s)",
                        self.chain,
                        _short_url(ep.url),
                        exc.kind,
                    )
                    continue
                base = exc.retry_after or (
                    60.0 if exc.kind == "rate_limit" else self.cfg.cooldown_min_sec
                )
                ep.health(method_group).last_error = exc.kind
                slept = ep.cool(base, self.cfg.cooldown_max_sec, method_group)
                log.warning(
                    "[%s] %s %s, cooldown %.0fs",
                    self.chain,
                    _short_url(ep.url),
                    exc.kind,
                    slept,
                )
                if not BALANCE_RPC.get():
                    await asyncio.sleep(min(2.0, slept))
            except Exception as exc:
                last = exc
                ep.last_error = type(exc).__name__
                ep.health(method_group).last_error = type(exc).__name__
                slept = ep.cool(
                    self.cfg.cooldown_min_sec, self.cfg.cooldown_max_sec, method_group
                )
                log.warning(
                    "[%s] %s unexpected %s, cooldown %.0fs",
                    self.chain,
                    _short_url(ep.url),
                    type(exc).__name__,
                    slept,
                )
                if not BALANCE_RPC.get():
                    await asyncio.sleep(min(1.0, slept))
        if isinstance(last, RpcError):
            raise last
        raise RpcError("network", "RPC retries exhausted")

    async def call(self, method: str, params: list[Any]) -> Any:
        payload = {"jsonrpc": "2.0", "id": 1, "method": method, "params": params}
        _ep, data = await self._post(payload, rpc_method_group(method))
        if not isinstance(data, dict) or "result" not in data:
            raise RpcError("malformed", f"{method}: missing result")
        result = data["result"]
        if method == "eth_blockNumber" and result is not None:
            self.last_head = hex_int(result)
        return result

    async def confirmed_call(
        self, method: str, params: list[Any], expected: Any,
    ) -> tuple[bool | None, Any | None]:
        """Repeat a sensitive read on a distinct endpoint without trusting it for valuation."""
        group = rpc_method_group(method)
        payload = {"jsonrpc": "2.0", "id": 91, "method": method, "params": params}
        first = self._pick(group)
        if first is None:
            return None, None
        try:
            endpoint, data = await self._post(payload, group, {first.fingerprint})
        except RpcError:
            return None, None
        if not isinstance(data, dict) or "result" not in data:
            return False, None
        return data["result"] == expected, data["result"]

    async def batch_partial(self, calls: list[tuple[str, list[Any]]]) -> list[Any | Exception]:
        if not calls:
            return []
        if self.preflight_complete and not any(
            endpoint.chain_verified
            and endpoint.permanent_error is None
            and endpoint.batch_supported
            for endpoint in self.endpoints
        ):
            return list(
                await asyncio.gather(
                    *(self.call(method, params) for method, params in calls),
                    return_exceptions=True,
                )
            )
        payload = [
            {"jsonrpc": "2.0", "id": i, "method": method, "params": params}
            for i, (method, params) in enumerate(calls)
        ]
        endpoint: Endpoint | None = None
        try:
            group = rpc_method_group(calls[0][0]) if calls else "head"
            endpoint, data = await self._post(payload, group)
        except RpcError:
            data = None
        if not isinstance(data, list):
            if endpoint is not None:
                endpoint.batch_supported = False
            return list(
                await asyncio.gather(
                    *(self.call(method, params) for method, params in calls),
                    return_exceptions=True,
                )
            )

        by_id = {item.get("id"): item for item in data if isinstance(item, dict)}
        out: list[Any | Exception] = [RpcError("rpc", "missing batch item") for _ in calls]
        retry_indexes: list[int] = []
        for i in range(len(calls)):
            item = by_id.get(i)
            if (
                not item
                or item.get("error") is not None
                or "result" not in item
                or item.get("result") is None
            ):
                retry_indexes.append(i)
            else:
                out[i] = item["result"]
        if retry_indexes:
            async def retry_one(index: int) -> Any:
                last_result: Any = None
                for _ in range(max(1, min(len(self.endpoints), self.retries))):
                    last_result = await self.call(*calls[index])
                    if last_result is not None:
                        return last_result
                return last_result

            retried = await asyncio.gather(
                *(retry_one(i) for i in retry_indexes),
                return_exceptions=True,
            )
            for i, result in zip(retry_indexes, retried):
                out[i] = result
        return out

    async def batch(self, calls: list[tuple[str, list[Any]]]) -> list[Any]:
        out = await self.batch_partial(calls)
        for item in out:
            if isinstance(item, Exception):
                raise item
        return out

    async def preflight(self) -> list[dict[str, Any]]:
        """Validate endpoint chain IDs and basic calls without exposing full URLs."""
        assert self._client
        # Startup must not wait behind a dead public endpoint.  This is a
        # diagnostic probe, not an indexing request: a short deadline is
        # enough to choose a healthy fallback and let workers start.
        async def request(ep: Endpoint, payload: Any) -> Any:
            try:
                return await asyncio.wait_for(
                    self._request_endpoint(ep, payload, "head"), timeout=30.0,
                )
            except asyncio.TimeoutError as exc:
                raise RpcError("timeout", "preflight request timed out") from exc

        async def probe(ep: Endpoint) -> dict[str, Any]:
            row: dict[str, Any] = {"endpoint": _short_url(ep.url), "ok": False}
            try:
                chain_data = await request(
                    ep,
                    {"jsonrpc": "2.0", "id": 1, "method": "eth_chainId", "params": []},
                )
                if not isinstance(chain_data, dict) or "result" not in chain_data:
                    raise RpcError("malformed", "eth_chainId missing result")
                actual = hex_int(chain_data["result"])
                row["chain_id"] = actual
                if self.expected_chain_id is not None and actual != self.expected_chain_id:
                    reason = f"wrong chain id {actual}, expected {self.expected_chain_id}"
                    ep.disable(reason)
                    raise RpcError("wrong_chain", reason, permanent=True)
                ep.chain_verified = True

                head_data = await request(
                    ep,
                    {"jsonrpc": "2.0", "id": 2, "method": "eth_blockNumber", "params": []},
                )
                if not isinstance(head_data, dict) or "result" not in head_data:
                    raise RpcError("malformed", "eth_blockNumber missing result")
                row["head"] = hex_int(head_data["result"])
                self.last_head = max(self.last_head or 0, int(row["head"]))
                try:
                    batch_data = await request(
                        ep,
                        [
                            {"jsonrpc": "2.0", "id": 11, "method": "eth_chainId", "params": []},
                            {"jsonrpc": "2.0", "id": 12, "method": "eth_blockNumber", "params": []},
                        ],
                    )
                    ep.batch_supported = (
                        isinstance(batch_data, list)
                        and len(batch_data) == 2
                        and all(
                            isinstance(item, dict)
                            and item.get("error") is None
                            and "result" in item
                            for item in batch_data
                        )
                    )
                except RpcError:
                    ep.batch_supported = False
                row["batch"] = ep.batch_supported
                row["ok"] = True
            except RpcError as exc:
                row["error"] = exc.kind
                if exc.permanent or exc.kind in ("auth", "wrong_chain"):
                    ep.disable(exc.kind)
                else:
                    ep.last_error = exc.kind
            return row

        # Environment variables may contain a long historical rotation of
        # keys.  A startup gate only needs the preferred two plus one
        # fallback; probing every dormant fallback serially behind the
        # per-chain semaphore delayed the scanner by minutes.
        report = list(await asyncio.gather(*(probe(ep) for ep in self.endpoints[:3])))
        self.preflight_complete = True
        return report

    def active_endpoint(self) -> str:
        ep = next(
            (
                candidate for candidate in self.endpoints
                if candidate.available("head")
                and (not self.preflight_complete or candidate.chain_verified)
            ),
            None,
        )
        return _short_url(ep.url) if ep else "none"

    def usable(self) -> bool:
        return any(
            endpoint.permanent_error is None
            and (not self.preflight_complete or endpoint.chain_verified)
            for endpoint in self.endpoints
        )

    def take_metrics(self) -> dict[str, Any]:
        latencies = self.metric_latencies_ms
        result = {
            "requests": self.metric_requests,
            "successes": self.metric_successes,
            "errors": sum(self.metric_errors.values()),
            "errors_by_type": dict(self.metric_errors),
            "latency_p50_ms": percentile(latencies, 0.50),
            "latency_p95_ms": percentile(latencies, 0.95),
        }
        self.metric_requests = 0
        self.metric_successes = 0
        self.metric_errors = {}
        self.metric_latencies_ms = []
        return result

    def health_rows(self) -> list[dict[str, Any]]:
        now = datetime.now(timezone.utc).isoformat()
        rows: list[dict[str, Any]] = []
        for endpoint in self.endpoints:
            hostname = _short_url(endpoint.url)
            for group, state in endpoint.method_health.items():
                limit = {
                    "blocks": self.block_batch,
                    "receipts": self.receipt_batch,
                    "code": self.call_batch,
                    "eth_call": self.call_batch,
                }.get(group)
                rows.append({
                    "chain": self.chain, "endpoint_fp": endpoint.fingerprint,
                    "hostname": hostname, "method_group": group,
                    "fail_streak": state.fail_streak, "success_streak": state.ok_streak,
                    "cooldown_until": state.cooldown_until,
                    "circuit_until": state.circuit_until,
                    "permanent_error": endpoint.permanent_error or state.permanent_error,
                    "latency_ewma_ms": state.latency_ewma_ms, "batch_limit": limit,
                    "success_since_resize": state.success_since_resize,
                    "updated_at": now,
                })
        return rows

    def method_cooldowns(self) -> dict[str, float]:
        now = time.time()
        result: dict[str, float] = {}
        for endpoint in self.endpoints:
            for group, state in endpoint.method_health.items():
                result[f"{_short_url(endpoint.url)}:{group}"] = max(
                    0.0, max(state.cooldown_until, state.circuit_until) - now
                )
        return result


def hex_int(v: Any) -> int:
    if v is None:
        return 0
    if isinstance(v, int):
        return v
    return int(v, 16)


def pad_addr(addr: str) -> str:
    return addr.lower().replace("0x", "").rjust(64, "0")


def topic_to_addr(topic: str) -> str:
    return "0x" + topic[-40:].lower()


def decode_uint(data: str | None) -> int:
    if not data or data == "0x":
        return 0
    return int(data, 16)


def decode_string(data: str | None) -> str | None:
    if not data or data == "0x" or len(data) < 2 + 64:
        return None
    raw = bytes.fromhex(data[2:])
    if len(raw) >= 64 and int.from_bytes(raw[:32], "big") == 32:
        n = int.from_bytes(raw[32:64], "big")
        return raw[64 : 64 + n].decode("utf-8", "replace").strip("\x00")
    try:
        return raw[:32].split(b"\x00", 1)[0].decode("utf-8", "replace") or None
    except Exception:
        return None


# ---------------------------------------------------------------------------
# Prices
# ---------------------------------------------------------------------------


class PriceBook:
    def __init__(self, timeout: int, batch_size: int = 40):
        self.timeout = timeout
        self.batch_size = max(1, batch_size)
        self.cache: dict[str, tuple[float, float]] = {}  # key -> (price, ts)
        self._lock = asyncio.Lock()
        self._missing_until: dict[str, float] = {}
        self._cooldown_until = 0.0

    async def fetch(self, keys: list[str], ttl: float = 120.0) -> dict[str, float]:
        keys = list(dict.fromkeys(keys))
        # Coalesce concurrent misses: addresses usually need the same seed prices.
        async with self._lock:
            now = time.time()
            missing = [k for k in keys if
                       (k not in self.cache or now - self.cache[k][1] > ttl)
                       and self._missing_until.get(k, 0) <= now]
            if missing and now >= self._cooldown_until:
                async with httpx.AsyncClient(timeout=min(self.timeout, 8)) as client:
                    for i in range(0, len(missing), self.batch_size):
                        chunk = missing[i:i + self.batch_size]
                        try:
                            r = await client.get(LLAMA_PRICES + ",".join(chunk))
                            if r.status_code == 429:
                                self._cooldown_until = time.time() + (_parse_retry_after(r.headers) or 30)
                                log.warning("DefiLlama rate limit; prices temporarily unavailable")
                                break
                            r.raise_for_status()
                            coins = r.json().get("coins") or {}
                            ts = time.time()
                            for k in chunk:
                                px = coins.get(k, {}).get("price")
                                if px is not None and math.isfinite(float(px)) and float(px) > 0:
                                    self.cache[k] = (float(px), ts)
                                else:
                                    self._missing_until[k] = ts + 30
                        except Exception as exc:
                            self._cooldown_until = time.time() + 10
                            log.warning("DefiLlama prices error: %s", type(exc).__name__)
                            break
        now = time.time()
        return {k: self.cache[k][0] for k in keys
                if k in self.cache and now - self.cache[k][1] <= ttl}


def llama_key(chain: ChainCfg, token: str | None = None) -> str:
    if token is None:
        return f"coingecko:{chain.native_coingecko}"
    return f"{chain.llama_chain}:{token.lower()}"


# ---------------------------------------------------------------------------
# Indexer
# ---------------------------------------------------------------------------


def normalize_evm_address(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    address = value.lower()
    if len(address) != 42 or not address.startswith("0x"):
        return None
    try:
        bytes.fromhex(address[2:])
    except ValueError:
        return None
    return address


def rpc_code_has_bytecode(value: Any) -> bool:
    if not isinstance(value, str) or not value.startswith("0x") or len(value) % 2:
        raise RpcError("malformed", "eth_getCode returned invalid bytecode")
    try:
        bytes.fromhex(value[2:])
    except ValueError as exc:
        raise RpcError("malformed", "eth_getCode returned invalid bytecode") from exc
    return bool(value[2:].lstrip("0"))


class DiscoverySlots:
    """A fair global 3:1 live/backfill scheduler with independent capacities."""

    def __init__(self, live_slots: int = 2, backfill_slots: int = 1):
        self.live_slots = max(1, live_slots)
        self.backfill_slots = max(1, backfill_slots)
        self.active = {"live": 0, "backfill": 0}
        self.waiting = {"live": 0, "backfill": 0}
        self.live_grants = 0
        self.condition = asyncio.Condition()

    def _allowed(self, role: str) -> bool:
        capacity = self.live_slots if role == "live" else self.backfill_slots
        if self.active[role] >= capacity:
            return False
        if role == "live":
            return self.live_grants < 3 or self.waiting["backfill"] == 0
        return self.live_grants >= 3 or self.waiting["live"] == 0

    @asynccontextmanager
    async def slot(self, role: str):
        async with self.condition:
            self.waiting[role] += 1
            try:
                await self.condition.wait_for(lambda: self._allowed(role))
                self.active[role] += 1
                if role == "live":
                    self.live_grants += 1
                else:
                    self.live_grants = 0
            finally:
                self.waiting[role] -= 1
        try:
            yield
        finally:
            async with self.condition:
                self.active[role] -= 1
                self.condition.notify_all()


async def collect_tx_to_contracts(
    db: DB, chain: ChainCfg, rpc: RpcPool, cfg: AppCfg,
    candidates: dict[str, tuple[dict[str, Any], int]], known_contracts: set[str],
    known_active_calls: set[str], checked_block: int, seen_at: str,
) -> tuple[list[tuple[Any, ...]], list[tuple[Any, ...]], list[tuple[Any, ...]], set[str], int]:
    cached = db.cached_code_statuses(
        chain.key,
        [address for address in candidates if address not in known_contracts],
        cfg.code_cache_ttl_sec,
    )
    contract_addresses = {
        address for address in candidates
        if address in known_contracts or cached.get(address) is True
    }
    unknown = [
        address for address in candidates
        if address not in known_contracts and address not in cached
    ]
    cache_rows: list[tuple[Any, ...]] = []
    offset = 0
    while offset < len(unknown):
        size = max(1, rpc.call_batch)
        chunk = unknown[offset:offset + size]
        results = await rpc.batch([("eth_getCode", [address, "latest"]) for address in chunk])
        if len(results) != len(chunk):
            rpc.shrink_batch("call")
            raise RpcError("partial", "one or more eth_getCode results are missing")
        parsed = [rpc_code_has_bytecode(value) for value in results]
        for address, has_code in zip(chunk, parsed):
            cache_rows.append((chain.key, address, int(has_code), seen_at, checked_block))
            if has_code:
                contract_addresses.add(address)
        rpc.grow_batch("call")
        offset += len(chunk)
    contracts = [
        (chain.key, address, None, None, None, seen_at)
        for address in sorted(contract_addresses) if address not in known_contracts
    ]
    discoveries = []
    for address in sorted(contract_addresses):
        if address in known_active_calls:
            continue
        tx, block_number = candidates[address]
        discoveries.append((
            chain.key, address, "active_call", block_number, tx.get("hash"),
            normalize_evm_address(tx.get("from")), seen_at,
        ))
    return contracts, discoveries, cache_rows, contract_addresses, len(unknown)


async def collect_transfers(
    chain: ChainCfg, rpc: RpcPool, start: int, end: int, known: set[str],
) -> tuple[list[tuple[Any, ...]], list[tuple[Any, ...]]]:
    if not known:
        return [], []
    step = max(1, min(chain.logs_max_range, end - start + 1))
    token_rows: dict[str, tuple[Any, ...]] = {}
    relation_rows: dict[tuple[str, str], tuple[Any, ...]] = {}
    known_list = sorted(known)
    cursor = start
    while cursor <= end:
        range_end = min(end, cursor + step - 1)
        try:
            logs: list[dict[str, Any]] = []
            for offset in range(0, len(known_list), 50):
                holders = known_list[offset:offset + 50]
                topics = ["0x" + pad_addr(holder) for holder in holders]
                part = await rpc.call("eth_getLogs", [{
                    "fromBlock": hex(cursor), "toBlock": hex(range_end),
                    "topics": [TRANSFER_TOPIC, None, topics[0] if len(topics) == 1 else topics],
                }])
                if not isinstance(part, list):
                    raise RpcError("malformed", "eth_getLogs did not return a list")
                logs.extend(part)
        except RpcError as exc:
            if exc.kind in {"range", "rate_limit"} and step > 1:
                step = max(1, step // 2)
                continue
            raise
        for item in logs:
            topics = item.get("topics") or []
            if len(topics) < 3:
                continue
            holder = topic_to_addr(topics[2])
            token = normalize_evm_address(item.get("address"))
            if holder not in known or token is None or token == ZERO:
                continue
            block_number = hex_int(item.get("blockNumber"))
            token_rows[token] = (chain.key, token, None, None, "transfer_log")
            relation_rows[(holder, token)] = (
                chain.key, holder, token, "transfer_log", block_number,
            )
        cursor = range_end + 1
        if step < chain.logs_max_range:
            step = min(chain.logs_max_range, step * 2)
    return list(token_rows.values()), list(relation_rows.values())


async def discover_tx_to_contracts(
    db: DB,
    chain: ChainCfg,
    rpc: RpcPool,
    cfg: AppCfg,
    candidates: dict[str, tuple[dict[str, Any], int]],
    known_contracts: set[str],
    known_active_calls: set[str],
    checked_block: int,
    seen_at: str,
) -> tuple[int, int, int]:
    """Persist unique top-level tx.to contracts; all RPC results are mandatory."""
    if not candidates:
        return 0, 0, 0

    cached = db.cached_code_statuses(
        chain.key,
        [address for address in candidates if address not in known_contracts],
        cfg.code_cache_ttl_sec,
    )
    contract_addresses = {
        address for address in candidates
        if address in known_contracts or cached.get(address) is True
    }
    unknown = [
        address for address in candidates
        if address not in known_contracts and address not in cached
    ]
    cache_rows: list[tuple[Any, ...]] = []
    offset = 0
    while offset < len(unknown):
        chunk_size = max(1, rpc.call_batch)
        chunk = unknown[offset : offset + chunk_size]
        try:
            results = await rpc.batch(
                [("eth_getCode", [address, "latest"]) for address in chunk]
            )
            if len(results) != len(chunk):
                raise RpcError("partial", "one or more eth_getCode results are missing")
            parsed = [rpc_code_has_bytecode(result) for result in results]
        except Exception:
            rpc.shrink_batch("call")
            raise
        for address, has_code in zip(chunk, parsed):
            cache_rows.append(
                (chain.key, address, int(has_code), seen_at, checked_block)
            )
            if has_code:
                contract_addresses.add(address)
        rpc.grow_batch("call")
        offset += len(chunk)

    # Delay cache writes until every code batch succeeded. A failed code lookup
    # therefore cannot make a partially checked range appear complete.
    db.upsert_code_cache(cache_rows)
    new_contract_rows = []
    discovery_rows = []
    for address in sorted(contract_addresses):
        tx, block_number = candidates[address]
        if address not in known_contracts:
            new_contract_rows.append(
                (chain.key, address, None, None, None, seen_at)
            )
        if address not in known_active_calls:
            discovery_rows.append(
                (
                    chain.key,
                    address,
                    "active_call",
                    block_number,
                    tx.get("hash"),
                    normalize_evm_address(tx.get("from")),
                    seen_at,
                )
            )

    inserted = db.upsert_contracts(new_contract_rows)
    sources_inserted = db.upsert_contract_discoveries(discovery_rows)
    known_contracts.update(contract_addresses)
    known_active_calls.update(row[1] for row in discovery_rows)
    return inserted, sources_inserted, len(unknown)


async def latest_block(rpc: RpcPool) -> int:
    return hex_int(await rpc.call("eth_blockNumber", []))


async def index_chain_cursor(
    db: DB, chain: ChainCfg, rpc: RpcPool, cfg: AppCfg, stop: asyncio.Event,
    role: str, slots: DiscoverySlots, from_block_override: int | None = None,
    to_block_override: int | None = None, monitor: MonitorStore | None = None,
) -> None:
    if from_block_override is not None:
        db.advance_cursor_start(chain.key, role, from_block_override)
    known_contracts = db.contract_addresses(chain.key)
    known_active_calls = db.discovery_addresses(chain.key, "active_call")
    failures = 0
    while not stop.is_set():
        await wait_if_paused(monitor, stop)
        cursor = db.cursor(chain.key, role)
        if cursor is None:
            raise RuntimeError(f"{chain.key}/{role}: cursor not initialized")
        if role == "backfill" and cursor["status"] == "complete":
            log.info("[%s/%s] anchor %s complete", chain.key, role, cursor["anchor_block"])
            return
        try:
            head = max(0, await latest_block(rpc) - chain.confirmations)
            target = head if role == "live" else min(head, int(cursor["anchor_block"]))
            if to_block_override is not None:
                target = min(target, to_block_override)
        except Exception as exc:
            if stop.is_set():
                return
            log.warning("[%s/%s] head unavailable: %s", chain.key, role, type(exc).__name__)
            await asyncio.sleep(cfg.all_down_sleep_sec)
            continue
        start = int(cursor["next_block"])
        if start > target:
            if role == "backfill":
                # Empty migrated ranges are complete without manufacturing a block commit.
                with db._lock, db.conn:
                    db.conn.execute(
                        "UPDATE chain_cursors SET status='complete',updated_at=? WHERE chain=? AND role='backfill'",
                        (datetime.now(timezone.utc).isoformat(), chain.key),
                    )
                return
            try:
                await asyncio.wait_for(stop.wait(), timeout=6)
            except asyncio.TimeoutError:
                pass
            continue
        end = min(target, start + max(1, rpc.block_batch) - 1)
        async with slots.slot(role):
            try:
                numbers = list(range(start, end + 1))
                blocks = await rpc.batch([
                    ("eth_getBlockByNumber", [hex(number), True]) for number in numbers
                ])
                if len(blocks) != len(numbers) or any(not isinstance(block, dict) for block in blocks):
                    raise RpcError("partial", "one or more blocks are missing")
                for expected, block in zip(numbers, blocks):
                    if hex_int(block.get("number")) != expected:
                        raise RpcError("partial", f"wrong block returned for {expected}")
                deployments: list[dict[str, Any]] = []
                candidates: dict[str, tuple[dict[str, Any], int]] = {}
                for number, block in zip(numbers, blocks):
                    for tx in block.get("transactions") or []:
                        if tx.get("to") in (None, "", "0x"):
                            deployments.append(tx)
                        elif cfg.discover_tx_to_contracts:
                            address = normalize_evm_address(tx.get("to"))
                            if address is not None and address not in known_active_calls:
                                candidates.setdefault(address, (tx, number))
                receipts: list[dict[str, Any]] = []
                for offset in range(0, len(deployments), max(1, rpc.receipt_batch)):
                    chunk = deployments[offset:offset + max(1, rpc.receipt_batch)]
                    part = await rpc.batch([
                        ("eth_getTransactionReceipt", [tx["hash"]]) for tx in chunk
                    ])
                    if len(part) != len(chunk) or any(not isinstance(item, dict) for item in part):
                        rpc.shrink_batch("receipt")
                        raise RpcError("partial", "one or more deployment receipts are missing")
                    receipts.extend(part)
                    rpc.grow_batch("receipt")
                now = datetime.now(timezone.utc).isoformat()
                direct_contracts: list[tuple[Any, ...]] = []
                direct_discoveries: list[tuple[Any, ...]] = []
                for tx, receipt in zip(deployments, receipts):
                    address = normalize_evm_address(receipt.get("contractAddress"))
                    if address is None:
                        continue
                    block_number = hex_int(tx.get("blockNumber"))
                    creator = normalize_evm_address(tx.get("from"))
                    direct_contracts.append((
                        chain.key, address, block_number, tx.get("hash"), creator, now,
                    ))
                    direct_discoveries.append((
                        chain.key, address, "direct_deploy", block_number,
                        tx.get("hash"), creator, now,
                    ))
                active_contracts: list[tuple[Any, ...]] = []
                active_discoveries: list[tuple[Any, ...]] = []
                cache_rows: list[tuple[Any, ...]] = []
                found_active: set[str] = set()
                code_checked = 0
                if cfg.discover_tx_to_contracts:
                    (active_contracts, active_discoveries, cache_rows,
                     found_active, code_checked) = await collect_tx_to_contracts(
                        db, chain, rpc, cfg, candidates, known_contracts,
                        known_active_calls, end, now,
                    )
                range_contracts = {
                    row[1] for row in [*direct_contracts, *active_contracts]
                }
                tokens: list[tuple[Any, ...]] = []
                relations: list[tuple[Any, ...]] = []
                if cfg.discover_tokens_from_transfers:
                    tokens, relations = await collect_transfers(
                        chain, rpc, start, end, known_contracts | range_contracts,
                    )
                inserted, source_inserted = db.commit_index_range(
                    chain.key, role, end, [*direct_contracts, *active_contracts],
                    [*direct_discoveries, *active_discoveries], cache_rows,
                    tokens, relations,
                )
                known_contracts.update(range_contracts)
                known_active_calls.update(found_active)
                failures = 0
                rpc.grow_batch("block")
                log.info(
                    "[%s/%s] idx %s..%s head=%s new=%s sources=%s code=%s",
                    chain.key, role, start, end, head, inserted, source_inserted, code_checked,
                )
            except Exception as exc:
                failures += 1
                rpc.shrink_batch("block")
                log.warning(
                    "[%s/%s] range %s-%s not committed: %s",
                    chain.key, role, start, end,
                    exc.kind if isinstance(exc, RpcError) else type(exc).__name__,
                )
                await asyncio.sleep(min(120, cfg.cooldown_min_sec * max(1, failures)))


async def index_chain(
    db: DB,
    chain: ChainCfg,
    rpc: RpcPool,
    cfg: AppCfg,
    stop: asyncio.Event,
    from_block_override: int | None,
    to_block_override: int | None = None,
    once: bool = False,
    monitor: MonitorStore | None = None,
) -> None:
    previous_revision = db.discovery_revision(chain.key)
    activation_rewind = (
        cfg.discover_tx_to_contracts
        and previous_revision is not None
        and previous_revision < 2
    )
    last = db.init_chain(
        chain.key, chain.start_block, cfg.discover_tx_to_contracts
    )
    if (
        from_block_override is not None
        and last < from_block_override - 1
        and not activation_rewind
    ):
        last = from_block_override - 1
        db.set_last_indexed(chain.key, last)
    elif activation_rewind and from_block_override is not None:
        log.info(
            "[%s] --from-block ignored for the one-time tx.to rescan from %s",
            chain.key,
            last + 1,
        )

    log.info("[%s] indexer start from block %s", chain.key, last + 1)
    known_contracts = db.contract_addresses(chain.key)
    known_active_calls = (
        db.discovery_addresses(chain.key, "active_call")
        if cfg.discover_tx_to_contracts else set()
    )
    fixed_target: int | None = None
    range_failures = 0
    while not stop.is_set():
        await wait_if_paused(monitor, stop)
        if stop.is_set():
            break
        if fixed_target is not None and last >= fixed_target:
            log.info("[%s] finite index pass complete at %s", chain.key, last)
            return
        if not once and to_block_override is not None and last >= to_block_override:
            log.info("[%s] configured end block reached at %s", chain.key, last)
            return
        try:
            latest = await latest_block(rpc)
            target = max(0, latest - chain.confirmations)
            if to_block_override is not None:
                target = min(target, to_block_override)
            if once:
                if fixed_target is None:
                    fixed_target = target
                target = fixed_target
        except Exception as exc:
            log.warning("[%s] eth_blockNumber: %s", chain.key, exc)
            if once:
                raise
            await asyncio.sleep(cfg.all_down_sleep_sec)
            continue

        if last >= target:
            if once or (to_block_override is not None and last >= to_block_override):
                log.info("[%s] finite index pass complete at %s", chain.key, last)
                return
            await asyncio.sleep(6)
            continue

        batch_end = min(target, last + max(1, rpc.block_batch))
        numbers = list(range(last + 1, batch_end + 1))
        try:
            blocks = await rpc.batch(
                [("eth_getBlockByNumber", [hex(number), True]) for number in numbers]
            )
            if len(blocks) != len(numbers) or any(not isinstance(block, dict) for block in blocks):
                raise RpcError("partial", f"missing block in {numbers[0]}-{batch_end}")
            for expected, block in zip(numbers, blocks):
                if hex_int(block.get("number")) != expected:
                    raise RpcError("partial", f"wrong block returned for {expected}")

            deployments: list[dict[str, Any]] = []
            tx_to_candidates: dict[str, tuple[dict[str, Any], int]] = {}
            for block_number, block in zip(numbers, blocks):
                for tx in block.get("transactions") or []:
                    if tx.get("to") in (None, "", "0x"):
                        deployments.append(tx)
                    elif cfg.discover_tx_to_contracts:
                        address = normalize_evm_address(tx.get("to"))
                        if address is not None and address not in known_active_calls:
                            tx_to_candidates.setdefault(address, (tx, block_number))

            receipts: list[Any] = []
            for offset in range(0, len(deployments), max(1, rpc.receipt_batch)):
                chunk = deployments[offset : offset + max(1, rpc.receipt_batch)]
                part = await rpc.batch(
                    [("eth_getTransactionReceipt", [tx["hash"]]) for tx in chunk]
                )
                if len(part) != len(chunk) or any(not isinstance(receipt, dict) for receipt in part):
                    rpc.shrink_batch("receipt")
                    raise RpcError("partial", "one or more deployment receipts are missing")
                receipts.extend(part)
                rpc.grow_batch("receipt")

            now = datetime.now(timezone.utc).isoformat()
            new_rows: list[tuple[Any, ...]] = []
            direct_source_rows: list[tuple[Any, ...]] = []
            for tx, receipt in zip(deployments, receipts):
                address = normalize_evm_address(receipt.get("contractAddress"))
                if address is None:
                    continue
                block_number = hex_int(tx.get("blockNumber"))
                creator = normalize_evm_address(tx.get("from"))
                new_rows.append(
                    (
                        chain.key,
                        address,
                        block_number,
                        tx.get("hash"),
                        creator,
                        now,
                    )
                )
                direct_source_rows.append(
                    (
                        chain.key,
                        address,
                        "direct_deploy",
                        block_number,
                        tx.get("hash"),
                        creator,
                        now,
                    )
                )
            inserted = db.upsert_contracts(new_rows)
            db.upsert_contract_discoveries(direct_source_rows)
            known_contracts.update(row[1] for row in new_rows)

            active_inserted = 0
            active_sources = 0
            code_checked = 0
            if cfg.discover_tx_to_contracts:
                active_inserted, active_sources, code_checked = await discover_tx_to_contracts(
                    db,
                    chain,
                    rpc,
                    cfg,
                    tx_to_candidates,
                    known_contracts,
                    known_active_calls,
                    batch_end,
                    now,
                )

            if cfg.discover_tokens_from_transfers:
                await harvest_transfers(db, chain, rpc, numbers[0], batch_end)

            # The cursor moves only after every mandatory block, receipt and log range succeeded.
            db.set_last_indexed(chain.key, batch_end)
            last = batch_end
            range_failures = 0
            rpc.grow_batch("block")
            log.info(
                "[%s] idx %s..%s safe_head=%s lag=%s new=%s active_sources=%s "
                "code_checked=%s batch=%s",
                chain.key,
                numbers[0],
                batch_end,
                target,
                max(0, target - last),
                inserted + active_inserted,
                active_sources,
                code_checked,
                len(numbers),
            )
        except Exception as exc:
            range_failures += 1
            rpc.shrink_batch("block")
            log.warning(
                "[%s] range %s-%s not committed: %s",
                chain.key,
                numbers[0],
                batch_end,
                exc,
            )
            if once and range_failures >= max(1, cfg.max_retries):
                raise
            await asyncio.sleep(cfg.cooldown_min_sec)


async def harvest_transfers(
    db: DB, chain: ChainCfg, rpc: RpcPool, start: int, end: int
) -> None:
    known = db.contract_addresses(chain.key)
    if not known:
        return

    step = max(1, min(chain.logs_max_range, end - start + 1))
    token_rows: dict[str, tuple[str, str, None, None, str]] = {}
    relation_rows: dict[
        tuple[str, str], tuple[str, str, str, str, int]
    ] = {}
    known_list = sorted(known)
    holder_chunk_size = 50
    cursor = start
    while cursor <= end:
        range_end = min(end, cursor + step - 1)
        try:
            logs: list[dict[str, Any]] = []
            for offset in range(0, len(known_list), holder_chunk_size):
                holders = known_list[offset : offset + holder_chunk_size]
                holder_topics = ["0x" + pad_addr(holder) for holder in holders]
                recipient_filter: str | list[str] = (
                    holder_topics[0] if len(holder_topics) == 1 else holder_topics
                )
                part = await rpc.call(
                    "eth_getLogs",
                    [
                        {
                            "fromBlock": hex(cursor),
                            "toBlock": hex(range_end),
                            "topics": [TRANSFER_TOPIC, None, recipient_filter],
                        }
                    ],
                )
                if not isinstance(part, list):
                    raise RpcError("malformed", "eth_getLogs did not return a list")
                logs.extend(part)
        except RpcError as exc:
            if exc.kind in ("range", "rate_limit") and step > 1:
                step = max(1, step // 2)
                log.info("[%s] eth_getLogs %s, range -> %s", chain.key, exc.kind, step)
                continue
            raise

        for item in logs:
            topics = item.get("topics") or []
            if len(topics) < 3:
                continue
            contract = topic_to_addr(topics[2])
            if contract not in known:
                continue
            token = (item.get("address") or "").lower()
            if not token or token == ZERO:
                continue
            block_number = hex_int(item.get("blockNumber"))
            token_rows[token] = (chain.key, token, None, None, "transfer_log")
            relation_rows[(contract, token)] = (
                chain.key,
                contract,
                token,
                "transfer_log",
                block_number,
            )
        cursor = range_end + 1
        if step < chain.logs_max_range:
            step = min(chain.logs_max_range, step * 2)

    db.upsert_tokens(list(token_rows.values()))
    db.upsert_contract_tokens(list(relation_rows.values()))


# ---------------------------------------------------------------------------
# Balances
# ---------------------------------------------------------------------------


def classify_address_scan(
    total_usd: float,
    chain_rows: list[dict[str, Any]],
    min_usd: float,
) -> tuple[str, int]:
    coverage = sum(1 for row in chain_rows if row.get("has_code") is not None)
    if total_usd >= min_usd:
        return "qualifying", coverage
    if chain_rows and all(row["status"] in ("complete", "absent") for row in chain_rows):
        return "below", coverage
    return "incomplete", coverage


async def scan_address_chain(
    db: DB,
    chain: ChainCfg,
    rpc: RpcPool | None,
    prices: PriceBook,
    cfg: AppCfg,
    address: str,
    balance_sem: asyncio.Semaphore,
    timeout_sec: float = 20,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    # Waiting for our local scheduler is not counted as a slow RPC.
    async with balance_sem:
        policy = BALANCE_RPC.set(True)
        try:
            return await _scan_address_chain(db, chain, rpc, prices, cfg, address, timeout_sec)
        finally:
            BALANCE_RPC.reset(policy)


async def _scan_address_chain(
    db: DB, chain: ChainCfg, rpc: RpcPool | None, prices: PriceBook, cfg: AppCfg,
    address: str, timeout_sec: float,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    base_row: dict[str, Any] = {
        "chain": chain.key,
        "has_code": None,
        "status": "rpc_error",
        "native_amount": None,
        "native_raw": None,
        "native_usd": None,
        "observed_native_usd": None,
        "included_native_usd": None,
        "excluded_usd": 0.0,
        "valuation_status": None,
        "checked_tokens": [],
        "failed_tokens": [],
        "tokens_usd": None,
        "total_usd": None,
        "note": None,
    }
    if rpc is None:
        base_row["note"] = "no usable RPC"
        return base_row, []

    token_rows: list[dict[str, Any]] = []
    try:
        async with asyncio.timeout(timeout_sec):
            code = await rpc.call("eth_getCode", [address, "latest"])
            if not isinstance(code, str) or not code.startswith("0x") or len(code) % 2:
                raise RpcError("malformed", "eth_getCode missing result")
            bytes.fromhex(code[2:])
            if not code[2:].lstrip("0"):
                return {
                    **base_row,
                    "has_code": 0,
                    "status": "absent",
                    "native_amount": 0.0,
                    "native_usd": 0.0,
                    "tokens_usd": 0.0,
                    "total_usd": 0.0,
                }, []

            base_row["has_code"] = 1
            tokens = db.tokens_for_contract(chain.key, address)
            for token in tokens:
                if token["decimals"] is not None and token["symbol"] is not None:
                    continue
                decimals = token["decimals"]
                symbol = token["symbol"]
                try:
                    if decimals is None:
                        raw_decimals = await rpc.call(
                            "eth_call",
                            [{"to": token["address"], "data": DECIMALS_SEL}, "latest"],
                        )
                        value = decode_uint(raw_decimals)
                        decimals = value if 0 <= value <= 36 else None
                    if symbol is None:
                        raw_symbol = await rpc.call(
                            "eth_call",
                            [{"to": token["address"], "data": SYMBOL_SEL}, "latest"],
                        )
                        symbol = decode_string(raw_symbol)
                except Exception as exc:
                    log.debug("[%s] token metadata %s: %s", chain.key, token["address"], exc)
                db.upsert_tokens(
                    [
                        (
                            chain.key,
                            token["address"],
                            symbol,
                            decimals,
                            token["source"] or "meta",
                        )
                    ]
                )
            tokens = db.tokens_for_contract(chain.key, address)

            price_keys = [llama_key(chain)] + [
                llama_key(chain, token["address"]) for token in tokens
            ]
            book = await prices.fetch(price_keys)
            raw_native = await rpc.call("eth_getBalance", [address, "latest"])
            if raw_native is None:
                raise RpcError("partial", "native balance missing")
            native_raw = decode_uint(raw_native)
            native_amount = native_raw / (10 ** chain.native_decimals)
            native_price = book.get(llama_key(chain))
            observed_native_usd = native_amount * native_price if native_price is not None else None
            policy = db.valuation_policy(chain.key, address, "native")
            valuation_status = "included"
            included_native_usd = observed_native_usd
            excluded_usd = 0.0
            anomaly = False
            anomaly_note: str | None = None
            if policy is not None and policy["policy"] in {"exclude_from_total", "quarantine"}:
                valuation_status = str(policy["policy"])
                included_native_usd = 0.0
                excluded_usd = float(observed_native_usd or 0.0)
                anomaly_note = str(policy["reason"])
            elif (
                observed_native_usd is not None
                and observed_native_usd >= cfg.native_anomaly_usd
                and not (policy is not None and policy["policy"] == "include")
            ):
                match, second_value = await rpc.confirmed_call(
                    "eth_getBalance", [address, "latest"], raw_native
                )
                early_raw: int | None = None
                try:
                    early_raw = decode_uint(
                        await rpc.call("eth_getBalance", [address, "0x0"])
                    )
                except Exception:
                    pass
                evidence = {
                    "threshold_usd": cfg.native_anomaly_usd,
                    "second_rpc_value": second_value,
                    "lifecycle": chain.lifecycle,
                    "early_balance_ratio": (
                        early_raw / native_raw if early_raw is not None and native_raw else None
                    ),
                }
                db.save_anomaly(
                    chain.key, address, "native", native_raw, observed_native_usd,
                    match, early_raw, "quarantined", evidence,
                )
                valuation_status = "anomalous_balance"
                included_native_usd = 0.0
                excluded_usd = float(observed_native_usd)
                anomaly = True
                anomaly_note = "large native balance quarantined pending valuation policy"
            native_usd = observed_native_usd
            base_row.update(
                native_raw=str(native_raw), native_amount=native_amount,
                native_usd=native_usd, observed_native_usd=observed_native_usd,
                included_native_usd=included_native_usd, excluded_usd=excluded_usd,
                valuation_status=valuation_status, tokens_usd=0.0,
                total_usd=included_native_usd or 0.0,
            )
            token_rows.append({
                "chain": chain.key, "token": "native", "symbol": chain.native_symbol,
                "raw_amount": native_raw, "amount": native_amount,
                "price_usd": native_price, "usd_value": native_usd,
                "priced": native_price is not None,
                "valuation_status": valuation_status,
            })
            checked_tokens = ["native"]
            failed_tokens: list[str] = []
            base_row["checked_tokens"] = checked_tokens
            base_row["failed_tokens"] = failed_tokens

            holder_data = BALANCE_OF_SEL + pad_addr(address)
            token_rpc_error = False
            unpriced_positive = native_raw > 0 and native_price is None
            tokens_usd = 0.0
            for offset in range(0, len(tokens), max(1, rpc.call_batch)):
                chunk = tokens[offset : offset + max(1, rpc.call_batch)]
                part = await rpc.batch_partial(
                    [
                        (
                            "eth_call",
                            [{"to": token["address"], "data": holder_data}, "latest"],
                        )
                        for token in chunk
                    ]
                )
                token_rpc_error |= len(part) != len(chunk)


                for token, result in zip(chunk, part):
                    if isinstance(result, Exception) or result is None:
                        token_rpc_error = True
                        failed_tokens.append(token["address"])
                        continue
                    try:
                        raw_amount = decode_uint(result)
                    except Exception:
                        token_rpc_error = True
                        failed_tokens.append(token["address"])
                        continue
                    checked_tokens.append(token["address"])
                    if raw_amount <= 0:
                        continue
                    decimals = token["decimals"]
                    amount = raw_amount / (10 ** int(decimals)) if decimals is not None else None
                    price = book.get(llama_key(chain, token["address"]))
                    priced = amount is not None and price is not None
                    usd_value = amount * price if priced else None
                    if usd_value is not None:
                        tokens_usd += usd_value
                    else:
                        unpriced_positive = True
                    token_rows.append(
                        {
                            "chain": chain.key,
                            "token": token["address"],
                            "symbol": token["symbol"] or token["address"][:10],
                            "raw_amount": raw_amount,
                            "amount": amount,
                            "price_usd": price,
                            "usd_value": usd_value,
                            "priced": priced,
                        }
                    )
                base_row.update(
                    tokens_usd=tokens_usd,
                    total_usd=(included_native_usd or 0.0) + tokens_usd,
                )

            known_total = (included_native_usd or 0.0) + tokens_usd
            if token_rpc_error:
                chain_status = "partial"
                note = "one or more token balance calls failed"
            elif anomaly:
                chain_status = "anomalous_balance"
                note = anomaly_note
            elif unpriced_positive:
                chain_status = "price_missing"
                note = "positive balance without metadata or price"
            else:
                chain_status = "complete"
                note = anomaly_note
            if chain.lifecycle != "active":
                lifecycle_note = f"network lifecycle={chain.lifecycle}"
                note = f"{note}; {lifecycle_note}" if note else lifecycle_note
            return {
                **base_row,
                "has_code": 1,
                "status": chain_status,
                "native_amount": native_amount,
                "native_usd": native_usd,
                "native_raw": str(native_raw),
                "observed_native_usd": observed_native_usd,
                "included_native_usd": included_native_usd,
                "excluded_usd": excluded_usd,
                "valuation_status": valuation_status,
                "checked_tokens": checked_tokens,
                "failed_tokens": failed_tokens,
                "tokens_usd": tokens_usd,
                "total_usd": known_total,
                "note": note,
            }, token_rows
    except Exception as exc:
        base_row["note"] = "chain_timeout" if isinstance(exc, TimeoutError) else type(exc).__name__
        if isinstance(exc, RpcError):
            base_row["note"] = exc.kind
        if base_row["total_usd"] is not None:
            base_row["status"] = "partial"
        return base_row, token_rows


class RabbyClient:
    """Optional public GETs only; auth failures disable this client for the run."""

    def __init__(self, cfg: AppCfg, transport: httpx.AsyncBaseTransport | None = None):
        self.interval = cfg.rabby_request_interval_sec
        self.client = httpx.AsyncClient(
            base_url="https://api.rabby.io", timeout=cfg.rabby_timeout_sec,
            headers={"Accept": "application/json"}, transport=transport,
        )
        self.lock = asyncio.Lock()
        self.next_request = 0.0
        self.cooldown_until = 0.0
        self.disabled = False

    async def close(self) -> None:
        await self.client.aclose()

    async def get(self, path: str, address: str) -> Any:
        async with self.lock:
            if self.disabled:
                raise RpcError("auth", "Rabby disabled for this run", permanent=True)
            if time.monotonic() < self.cooldown_until:
                raise RpcError("cooldown", "Rabby cooling down")
            await asyncio.sleep(max(0, self.next_request - time.monotonic()))
            self.next_request = time.monotonic() + self.interval
            try:
                response = await self.client.get(path, params={"id": address})
                if response.status_code in (401, 403):
                    self.disabled = True
                    raise RpcError("auth", "Rabby authentication required", permanent=True)
                if response.status_code == 429:
                    self.cooldown_until = time.monotonic() + (_parse_retry_after(response.headers) or 60)
                    raise RpcError("rate_limit", "Rabby rate limit")
                response.raise_for_status()
                data = response.json()
                if isinstance(data, dict) and (
                    data.get("error_code") not in (None, 0) or data.get("error")
                ):
                    raise RpcError("api_error", "Rabby error envelope")
                return data
            except (httpx.HTTPError, ValueError) as exc:
                self.cooldown_until = time.monotonic() + 30
                raise RpcError("api_error", type(exc).__name__) from exc

    async def snapshot(self, address: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        total = await self.get("/v1/user/total_balance", address)
        tokens = await self.get("/v1/user/cache_token_list", address)
        if not isinstance(total, dict) or not isinstance(total.get("chain_list"), list):
            raise RpcError("malformed", "Rabby chain_list missing")
        if not isinstance(tokens, list) or any(not isinstance(t, dict) for t in tokens):
            raise RpcError("malformed", "Rabby token list missing")
        return total, tokens


def _nonnegative_number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        result = float(value)
        return result if math.isfinite(result) and result >= 0 else None
    except (TypeError, ValueError):
        return None


def build_rabby_estimate(
    total: dict[str, Any], tokens: list[dict[str, Any]],
    chains: dict[str, ChainCfg], code_rows: list[dict[str, Any]],
    threshold: float, uncertainty: float,
) -> dict[str, Any]:
    # Map by numeric chain ID. Rabby aliases (eth, arb, xdai...) are not our keys.
    by_id: dict[int, dict[str, Any]] = {}
    aliases: dict[str, str] = {}
    supported = {chain.chain_id: key for key, chain in chains.items()}
    for item in total.get("chain_list", []):
        if not isinstance(item, dict):
            continue
        try:
            chain_id = int(item["community_id"])
        except (KeyError, ValueError, TypeError):
            continue
        if chain_id in by_id:
            raise RpcError("malformed", "duplicate Rabby chain")
        by_id[chain_id] = item
        if chain_id in supported:
            aliases[str(item.get("id"))] = supported[chain_id]
    codes = {row["chain"]: row.get("has_code") for row in code_rows}
    parts, token_parts = [], []
    known = 0.0
    covered = 0
    missing = False
    for key, chain in chains.items():
        code = codes.get(key)
        item = by_id.get(chain.chain_id)
        amount = _nonnegative_number(item.get("usd_value")) if item else None
        if code == 0:
            amount = 0.0  # Never count EOA holdings.
        elif code is None:
            amount = None
        if amount is None:
            missing = True
        else:
            known += amount
            covered += 1
        parts.append({"chain": key, "has_code": code, "estimated_usd": amount})
    for token in tokens:
        key = aliases.get(str(token.get("chain")))
        if key is None or codes.get(key) != 1 or token.get("is_scam") is True:
            continue
        amount = _nonnegative_number(token.get("amount"))
        price = _nonnegative_number(token.get("price"))
        if amount is None:
            missing = True
            continue
        if amount <= 0:
            continue
        priced = price is not None and price > 0
        missing |= not priced
        token_parts.append({
            "chain": key, "token": str(token.get("id", "")),
            "symbol": str(token.get("symbol") or token.get("id", "")),
            "decimals": token.get("decimals"),
            "raw_amount": str(token.get("raw_amount_str", token.get("raw_amount", ""))),
            "amount": amount, "price_usd": price if priced else None,
            "usd_value": amount * price if priced else None, "priced": priced,
        })
    # If estimate = real * (1 +/- uncertainty), these are the real-value bounds.
    value = known if covered else None
    delta = Decimal(str(uncertainty))
    lower = float(Decimal(str(value)) / (1 + delta)) if value is not None else None
    upper = float(Decimal(str(value)) / (1 - delta)) if value is not None and not missing else None
    if missing:
        status = "incomplete"
    elif lower >= threshold:
        status = "estimated_above"
    elif upper >= threshold:
        status = "near_threshold"
    else:
        status = "estimated_below"
    return {
        "status": status, "estimated_usd": value, "lower_usd": lower, "upper_usd": upper,
        "coverage": covered, "total_networks": len(chains), "uncertainty": uncertainty,
        "chains": parts, "tokens": token_parts,
        "note": "Rabby estimate only; uncertainty is an assumption, not a guarantee"
                + ("; missing code, network balance or token price; upper bound unknown" if missing else ""),
    }


async def rabby_fallback_loop(
    db: DB, chains: dict[str, ChainCfg], cfg: AppCfg, stop: asyncio.Event,
    rpc_finished: asyncio.Event, once: bool,
) -> None:
    client = RabbyClient(cfg)
    snapshot_at = time.time() if once else None
    try:
        while not stop.is_set():
            if client.disabled:
                return
            row = db.pending_rabby_scan(
                cfg.balance_retry_sec, snapshot_at,
                cfg.rabby_token_discovery_after_sec,
            )
            if row is None or time.monotonic() < client.cooldown_until:
                if once and rpc_finished.is_set() and row is None:
                    return
                if once and time.monotonic() < client.cooldown_until:
                    # A finite pass must not wait through service cooldowns.
                    return
                try:
                    await asyncio.wait_for(stop.wait(), timeout=1)
                except TimeoutError:
                    continue
                return
            try:
                total, tokens = await client.snapshot(row["address"])
                estimate = build_rabby_estimate(
                    total, tokens, chains,
                    [dict(part) for part in db.address_scan_chains(row["id"])],
                    cfg.min_usd, cfg.rabby_uncertainty,
                )
            except Exception as exc:
                estimate = {
                    "status": "incomplete", "estimated_usd": None, "lower_usd": None,
                    "upper_usd": None, "coverage": 0, "total_networks": len(chains),
                    "uncertainty": cfg.rabby_uncertainty, "chains": [], "tokens": [],
                    "note": "Rabby unavailable: " + (exc.kind if isinstance(exc, RpcError) else type(exc).__name__),
                }
            db.save_rabby_estimate(row["address"], estimate, row["id"])
            discovered = db.ingest_indexed_tokens(
                row["address"], estimate.get("tokens") or [], "rabby_discovery"
            )
            log.info("[rabby] %s estimate=%s status=%s coverage=%s/%s",
                     row["address"], estimate["estimated_usd"], estimate["status"],
                     estimate["coverage"], len(chains))
            if discovered:
                log.info(
                    "[rabby] %s discovered %s token contracts for RPC verification",
                    row["address"], discovered,
                )
    finally:
        await client.close()


async def check_multichain_balances(
    db: DB,
    chains: dict[str, ChainCfg],
    pools: dict[str, RpcPool],
    prices: PriceBook,
    cfg: AppCfg,
    min_usd: float,
    stop: asyncio.Event,
    once: bool = False,
    monitor: MonitorStore | None = None,
) -> None:
    log.info("multichain balance checker online (%s networks, %s addresses, %s chain slots)",
             len(chains), cfg.balance_concurrency, cfg.balance_chain_concurrency)
    chain_sem = asyncio.Semaphore(cfg.balance_chain_concurrency)
    rpc_finished = asyncio.Event()
    fallback = asyncio.create_task(
        rabby_fallback_loop(db, chains, cfg, stop, rpc_finished, once), name="rabby-fallback"
    ) if cfg.rabby_fallback else None

    async def scan_one(address: str) -> None:
        started = time.monotonic()
        due_keys = db.due_address_chains(address, list(chains), snapshot_at)
        if not due_keys:
            return
        jobs = {
            asyncio.create_task(scan_address_chain(
                db, chain, pools.get(chain.key), prices, cfg, address, chain_sem,
                cfg.balance_chain_timeout_sec,
            )): chain.key for chain in chains.values() if chain.key in due_keys
        }
        try:
            done, unfinished = await asyncio.wait(jobs, timeout=cfg.balance_address_timeout_sec)
            results = []
            for task, key in jobs.items():
                if task in done:
                    results.append(task.result())
                else:
                    results.append(({
                        "chain": key, "has_code": None, "status": "rpc_error",
                        "total_usd": None, "note": "address_timeout",
                    }, []))
            for task in unfinished:
                task.cancel()
            await asyncio.gather(*unfinished, return_exceptions=True)
        finally:
            for task in jobs:
                if not task.done():
                    task.cancel()
            await asyncio.gather(*jobs, return_exceptions=True)
        for chain_row, tokens in results:
            db.save_address_chain_state(address, chain_row, tokens, min_usd)
        chain_rows, token_rows = db.current_address_parts(address, list(chains))
        total_usd = sum(float(row.get("total_usd") or 0.0) for row in chain_rows)
        status, coverage = classify_address_scan(total_usd, chain_rows, min_usd)
        notes = sorted({row["status"] for row in chain_rows
                        if row["status"] not in ("complete", "absent")})
        db.save_address_scan(
            address, status, total_usd, coverage, len(chains),
            chain_rows, token_rows, ", ".join(notes) or None,
        )
        db.apply_address_schedule(address, status)
        log.info(
            "[balances] %s $%.2f %s coverage=%s/%s refreshed=%s elapsed=%.2fs",
            address, total_usd, status, coverage, len(chains), len(due_keys),
            time.monotonic() - started,
        )

    active: dict[asyncio.Task[Any], str] = {}
    stop_job = asyncio.create_task(stop.wait())
    snapshot_at = time.time() if once else None
    try:
        while not stop.is_set():
            await wait_if_paused(monitor, stop)
            if stop.is_set():
                break
            pending = db.pending_addresses(
                cfg.recheck_interval_sec, limit=max(80, cfg.balance_concurrency * 2),
                retry_sec=cfg.balance_retry_sec, as_of=snapshot_at,
            )
            in_flight = set(active.values())
            for address in pending:
                if len(active) >= cfg.balance_concurrency:
                    break
                if address not in in_flight:
                    active[asyncio.create_task(scan_one(address), name=f"balance-{address}")] = address
                    in_flight.add(address)
            if not active:
                if once:
                    break
                try:
                    await asyncio.wait_for(stop.wait(), timeout=3)
                except TimeoutError:
                    pass
                continue
            done, _ = await asyncio.wait([*active, stop_job], return_when=asyncio.FIRST_COMPLETED)
            for task in done:
                if task is not stop_job:
                    del active[task]
                    task.result()
            if fallback is not None and fallback.done():
                fallback.result()
        rpc_finished.set()
        if once and fallback is not None and not stop.is_set():
            await fallback
    finally:
        rpc_finished.set()
        jobs = [*active, stop_job] + ([fallback] if fallback is not None else [])
        for task in jobs:
            if not task.done():
                task.cancel()
        await asyncio.gather(*jobs, return_exceptions=True)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
OK_FILL = PatternFill("solid", fgColor="C6EFCE")
BELOW_FILL = PatternFill("solid", fgColor="FFF2CC")
INCOMPLETE_FILL = PatternFill("solid", fgColor="F4CCCC")
THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
MONEY = '#,##0.00"$"'


def _style_header(ws) -> None:
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws, widths: dict[int, int]) -> None:
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w


def export_xlsx(
    db: DB,
    cfg: AppCfg,
    chains: dict[str, ChainCfg],
    min_usd: float,
    snapshot: bool = False,
) -> tuple[Path, Path, Path]:
    cfg.export_dir.mkdir(parents=True, exist_ok=True)
    if snapshot:
        day = datetime.now(timezone.utc).strftime("%Y%m%d")
        folder = cfg.export_dir / "archive"
        folder.mkdir(parents=True, exist_ok=True)
        paths = (
            folder / f"qualifying_{day}.xlsx",
            folder / f"below_threshold_{day}.xlsx",
            folder / f"incomplete_{day}.xlsx",
        )
    else:
        paths = (
            cfg.export_dir / "qualifying.xlsx",
            cfg.export_dir / "below_threshold.xlsx",
            cfg.export_dir / "incomplete.xlsx",
        )

    headers = [
        "Address",
        "Total USD",
        "Coverage",
        "Status",
        "Networks with bytecode",
        "Network balances",
        "Tokens",
        "Discovery sources",
        "First seen (UTC)",
        "Last seen (UTC)",
        "Checked (UTC)",
        "Notes",
        "Rabby estimate USD",
        "Rabby lower USD (assumed tolerance)",
        "Rabby upper USD (assumed tolerance)",
        "Rabby status",
        "Rabby coverage",
        "Rabby checked (UTC)",
        "Rabby networks / tokens",
        "Rabby notes",
    ]

    def prepared(row: sqlite3.Row) -> list[Any]:
        chain_parts = db.address_scan_chains(row["id"])
        token_parts = db.address_scan_tokens(row["id"])
        sources = db.address_sources(row["address"])
        with_code = [
            chains[item["chain"]].display_name if item["chain"] in chains else item["chain"]
            for item in chain_parts
            if item["has_code"] == 1
        ]
        networks = []
        for item in chain_parts:
            if item["has_code"] != 1:
                continue
            amount = float(item["total_usd"] or 0.0)
            observed = item["observed_native_usd"] if "observed_native_usd" in item.keys() else None
            excluded = float(item["excluded_usd"] or 0.0) if "excluded_usd" in item.keys() else 0.0
            detail = f", observed native=${float(observed):,.2f}, excluded=${excluded:,.2f}" if excluded else ""
            networks.append(
                f"{item['chain']}=${amount:,.2f} ({item['status']}{detail}; "
                f"lifecycle={chains[item['chain']].lifecycle if item['chain'] in chains else 'unknown'})"
            )
        token_values = []
        for item in token_parts:
            label = item["symbol"] or item["token"][:10]
            if item["priced"]:
                valuation = item["valuation_status"] if "valuation_status" in item.keys() else "included"
                suffix = f" ({valuation})" if valuation != "included" else ""
                token_values.append(
                    f"{item['chain']}:{label}=${float(item['usd_value'] or 0):,.2f}{suffix}"
                )
            else:
                amount = item["amount"] if item["amount"] is not None else item["raw_amount"]
                token_values.append(f"{item['chain']}:{label}={amount} (unpriced)")
        source_values = [
            f"{item['chain']}:{item['source']}:{item['observed_block']}:"
            f"{item['observed_tx'] or ''}"
            for item in sources
        ]
        seen = [item["first_seen_at"] for item in sources if item["first_seen_at"]]
        estimate = db.rabby_estimate(row["address"])
        estimate_values: list[Any] = [None] * 8
        if estimate is not None:
            summary = "; ".join(
                f"{part['chain']}=${part['estimated_usd']}"
                for part in json.loads(estimate["chain_parts"]) if part.get("has_code") == 1
            )
            tokens_text = "; ".join(
                f"{part['chain']}:{part['symbol']}={part['amount']}"
                + (f" (${part['usd_value']:.2f})" if part["priced"] else " (unpriced)")
                for part in json.loads(estimate["token_parts"])
            )
            estimate_values = [
                estimate["estimated_usd"], estimate["lower_usd"], estimate["upper_usd"],
                estimate["status"], f"{estimate['coverage']}/{estimate['total_networks']}",
                estimate["checked_at"], (summary + " | " + tokens_text)[:32000],
                (estimate["note"] or "") + f"; assumed deviation={estimate['uncertainty']:.0%}"
                + ("; relates to an older RPC scan" if estimate["rpc_scan_id"] != row["id"] else ""),
            ]
        return [
            row["address"],
            float(row["total_usd"] or 0.0),
            f"{row['coverage']}/{row['total_networks']}",
            row["status"],
            "; ".join(with_code),
            "; ".join(networks),
            "; ".join(token_values),
            "; ".join(source_values),
            min(seen) if seen else None,
            max(seen) if seen else None,
            row["scanned_at"],
            row["note"],
        ] + estimate_values

    def build(path: Path, rows: list[sqlite3.Row], title: str, fill: PatternFill) -> None:
        tmp = path.with_suffix(".tmp.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]
        ws.append(headers)
        for row in rows:
            ws.append(prepared(row))
            for cell in ws[ws.max_row]:
                cell.border = THIN
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.fill = fill
            ws.cell(ws.max_row, 2).number_format = MONEY
        info = wb.create_sheet("Parameters")
        info.append(["Parameter", "Value"])
        info.append(["USD threshold", min_usd])
        info.append(["Exported UTC", datetime.now(timezone.utc).isoformat()])
        info.append(["Database", str(db.path)])
        info.append(["Rows", len(rows)])
        info.append(["Rabby", "Supplementary estimate; never changes RPC status or total"])
        _style_header(ws)
        _style_header(info)
        _autosize(
            ws,
            {1: 46, 2: 16, 3: 12, 4: 14, 5: 34, 6: 60, 7: 70, 8: 70,
             9: 24, 10: 24, 11: 24, 12: 35, 13: 20, 14: 24, 15: 24,
             16: 22, 17: 16, 18: 24, 19: 70, 20: 60},
        )
        _autosize(info, {1: 24, 2: 70})
        wb.save(tmp)
        tmp.replace(path)

    latest = db.latest_address_scans()
    # A database can contain scans made with a different --min-usd value.
    # Classify exports from the actual latest total, not the historic status label.
    qualifying = [
        row for row in latest
        if float(row["total_usd"] or 0.0) >= min_usd
    ]
    below = [
        row for row in latest
        if row["status"] == "below"
        and 0.0 < float(row["total_usd"] or 0.0) < min_usd
    ]
    incomplete = [
        row for row in latest
        if row["status"] == "incomplete" and float(row["total_usd"] or 0.0) < min_usd
    ]
    build(paths[0], qualifying, "Qualifying", OK_FILL)
    build(paths[1], below, "Below threshold", BELOW_FILL)
    build(paths[2], incomplete, "Incomplete", INCOMPLETE_FILL)
    log.info(
        "xlsx qualifying=%s below=%s incomplete=%s",
        len(qualifying),
        len(below),
        len(incomplete),
    )
    return paths


def export_bundle_process(
    db_path: Path, cfg: AppCfg, chains: dict[str, ChainCfg], min_usd: float,
) -> None:
    """Spawn target: export from query-only connections, never the scanner connection."""
    db = DB(db_path, read_only=True)
    sui_store = SuiStore(db_path, read_only=True)
    try:
        export_xlsx(db, cfg, chains, min_usd, snapshot=False)
        export_sui_xlsx(sui_store, cfg.export_dir, min_usd, snapshot=False)
    finally:
        sui_store.close()
        db.close()


async def run_export_process(
    db: DB, cfg: AppCfg, chains: dict[str, ChainCfg], lock: asyncio.Lock,
) -> tuple[Path, ...]:
    async with lock:
        context = multiprocessing.get_context("spawn")
        process = context.Process(
            target=export_bundle_process,
            args=(db.path, cfg, chains, cfg.min_usd),
            name="xlsx-export",
        )
        process.start()
        await asyncio.to_thread(process.join)
        if process.exitcode != 0:
            raise RuntimeError(f"XLSX export process exited with {process.exitcode}")
        return tuple(
            cfg.export_dir / name for name in (
                "qualifying.xlsx", "below_threshold.xlsx", "incomplete.xlsx",
                "sui_qualifying.xlsx", "sui_below_threshold.xlsx",
                "sui_incomplete.xlsx", "sui_packages.xlsx",
            )
        )


def copy_daily_report_snapshot(export_dir: Path, now: datetime | None = None) -> list[Path]:
    local = (now or datetime.now(timezone.utc)).astimezone(ZoneInfo("Europe/Moscow"))
    folder = export_dir / "archive" / local.strftime("%Y%m%d")
    folder.mkdir(parents=True, exist_ok=True)
    copied: list[Path] = []
    for source in export_dir.glob("*.xlsx"):
        target = folder / source.name
        temporary = target.with_suffix(".tmp.xlsx")
        shutil.copy2(source, temporary)
        temporary.replace(target)
        copied.append(target)
    return copied


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------


def seed_tokens(db: DB, seed: dict[str, list[dict[str, Any]]], enabled: list[str]) -> None:
    rows = []
    for chain in enabled:
        for t in seed.get(chain, []):
            rows.append((chain, t["address"].lower(), t.get("symbol"), t.get("decimals"), "seed"))
    db.upsert_tokens(rows)


async def safe_head(rpc: RpcPool) -> int | None:
    try:
        return await latest_block(rpc)
    except Exception as e:
        log.warning("[%s] не удалось снять head: %s", rpc.chain, e)
        return None


def write_status(path: Path, lines: list[str]) -> None:
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


async def wait_if_paused(monitor: MonitorStore | None, stop: asyncio.Event) -> None:
    while monitor is not None and monitor.setting("scanner_paused", "0") == "1":
        if stop.is_set():
            return
        try:
            await asyncio.wait_for(stop.wait(), timeout=1)
        except asyncio.TimeoutError:
            pass


def backup_database(db: DB, backup_dir: Path, keep: int = 7) -> Path:
    backup_dir.mkdir(parents=True, exist_ok=True)
    target = backup_dir / f"contracts_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S_%f')}.db"
    destination = sqlite3.connect(target)
    try:
        with db._lock:
            db.conn.backup(destination)
    finally:
        destination.close()
    backups = sorted(backup_dir.glob("contracts_*.db"), key=lambda item: item.stat().st_mtime)
    for old in backups[:-max(1, keep)]:
        old.unlink(missing_ok=True)
    return target


async def print_startup_status(
    selected: list[ChainCfg], pools: dict[str, RpcPool], db: DB
) -> None:
    header = [
        "",
        "=== DeFi scanner start ===",
        f"UTC        {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')}",
        f"DB         {db.path}",
        "",
        f"{'сеть':<14} {'start':>12} {'last_indexed':>14} {'head':>12} {'осталось':>12} {'контрактов':>12}",
        "-" * 80,
    ]
    for line in header:
        console.info(line)
        log.info(line if line else " ")

    rows = []
    for c in selected:
        last = db.last_indexed(c.key)
        if last == 0:
            last = max(c.start_block - 1, 0)
        head = await safe_head(pools[c.key])
        if head is not None:
            head = max(0, head - c.confirmations)
        remain = (head - last) if head is not None else None
        counts = db.chain_counts(c.key)
        line = (
            f"{c.key:<14} {c.start_block:>12,} {last:>14,} "
            f"{(f'{head:,}' if head is not None else 'n/a'):>12} "
            f"{(f'{remain:,}' if remain is not None else 'n/a'):>12} "
            f"{counts['contracts']:>12,}"
        )
        console.info(line)
        log.info(line)
        rows.append((c.key, last, head, remain, counts))

    console.info("-" * 80)
    console.info("Дальше работа в фоне. Логи: logs/scanner.log   статус: logs/status.txt")
    console.info("Excel: reports/qualifying.xlsx, below_threshold.xlsx, incomplete.xlsx")
    console.info("Остановка: Ctrl+C  (или systemctl stop / kill)")
    console.info("")


async def heartbeat_loop(
    db: DB,
    cfg: AppCfg,
    selected: list[ChainCfg],
    pools: dict[str, RpcPool],
    stop: asyncio.Event,
    monitor: MonitorStore,
    run_id: str,
) -> None:
    status_path = cfg.log_dir / "status.txt"
    previous: dict[str, tuple[float, int]] = {}
    discovery_keys = {chain.key for chain in selected}
    last_prune_day = ""
    while not stop.is_set():
        try:
            await asyncio.wait_for(stop.wait(), timeout=cfg.heartbeat_sec)
            break
        except asyncio.TimeoutError:
            pass
        sampled_at = datetime.now(timezone.utc)
        paused = monitor.setting("scanner_paused", "0") == "1"
        monitor.heartbeat(run_id, "paused" if paused else "running")
        lines = [
            f"updated_utc {sampled_at.isoformat()}",
            f"{'chain':<14} {'last':>12} {'safe_head':>12} {'lag':>10} {'contracts':>10} {'rpc':<30} {'cooldown':>9} {'error'}",
        ]
        for key, pool in pools.items():
            c = cfg.chains[key]
            role = "live" if key in discovery_keys else "balance"
            live_cursor = db.cursor(c.key, "live") if key in discovery_keys else None
            backfill_cursor = db.cursor(c.key, "backfill") if key in discovery_keys else None
            last = int(live_cursor["last_committed"]) if live_cursor is not None else None
            head = max(0, pool.last_head - c.confirmations) if pool.last_head is not None else None
            lag = (head - last) if head is not None and last is not None else None
            cnt = db.chain_counts(c.key)
            method_cooldowns = pool.method_cooldowns()
            cooldown = max(method_cooldowns.values(), default=0.0)
            errors = sorted({
                state.last_error
                for endpoint in pool.endpoints
                for state in endpoint.method_health.values()
                if state.last_error
            })
            old = previous.get(f"{c.key}:live")
            blocks_per_hour = None
            if old and last is not None and sampled_at.timestamp() > old[0]:
                blocks_per_hour = max(
                    0.0, (last - old[1]) * 3600 / (sampled_at.timestamp() - old[0])
                )
            if last is not None:
                previous[f"{c.key}:live"] = (sampled_at.timestamp(), last)
            metrics = pool.take_metrics()
            monitor.add_chain_sample(
                run_id=run_id, chain=c.key, role=role, cursor=last,
                safe_head=head, lag=lag, blocks_per_hour=blocks_per_hour,
                contracts=cnt["contracts"],
                direct_deploy=db.discovery_count(c.key, "direct_deploy"),
                active_call=db.discovery_count(c.key, "active_call"),
                active_rpc=pool.active_endpoint(), cooldown_sec=cooldown,
                rpc_requests=metrics["requests"], rpc_successes=metrics["successes"],
                rpc_errors=metrics["errors"], latency_p50_ms=metrics["latency_p50_ms"],
                latency_p95_ms=metrics["latency_p95_ms"],
                errors_json=metrics["errors_by_type"],
            )
            if backfill_cursor is not None:
                backfill_last = int(backfill_cursor["last_committed"])
                backfill_anchor = int(backfill_cursor["anchor_block"])
                old_backfill = previous.get(f"{c.key}:backfill")
                backfill_bph = None
                if old_backfill and sampled_at.timestamp() > old_backfill[0]:
                    backfill_bph = max(
                        0.0, (backfill_last - old_backfill[1]) * 3600
                        / (sampled_at.timestamp() - old_backfill[0])
                    )
                previous[f"{c.key}:backfill"] = (sampled_at.timestamp(), backfill_last)
                monitor.add_chain_sample(
                    run_id=run_id, chain=c.key, role="backfill", cursor=backfill_last,
                    safe_head=backfill_anchor, lag=max(0, backfill_anchor - backfill_last),
                    blocks_per_hour=backfill_bph, contracts=cnt["contracts"],
                    direct_deploy=db.discovery_count(c.key, "direct_deploy"),
                    active_call=db.discovery_count(c.key, "active_call"),
                    active_rpc=pool.active_endpoint(), cooldown_sec=cooldown,
                    rpc_requests=0, rpc_successes=0, rpc_errors=0,
                    latency_p50_ms=None, latency_p95_ms=None, errors_json={},
                )
            line = (
                f"{c.key:<14} {(f'{last:,}' if last is not None else 'balance'):>12} "
                f"{(f'{head:,}' if head is not None else 'n/a'):>12} "
                f"{(f'{lag:,}' if lag is not None else 'n/a'):>10} "
                f"{cnt['contracts']:>10,} {pool.active_endpoint():<30} "
                f"{cooldown:>8.0f}s {','.join(errors) or '-'}"
            )
            lines.append(line)
            active_method_cooldowns = [
                f"{name}={seconds:.0f}s" for name, seconds in method_cooldowns.items()
                if seconds > 0
            ]
            if active_method_cooldowns:
                lines.append("  method_cooldowns " + ", ".join(active_method_cooldowns))
            lines.append(
                f"  adaptive_batch blocks={pool.block_batch} receipts={pool.receipt_batch} calls={pool.call_batch}"
            )
            if backfill_cursor is not None:
                lines.append(
                    f"{c.key + '/backfill':<14} {int(backfill_cursor['last_committed']):>12,} "
                    f"{int(backfill_cursor['anchor_block']):>12,} "
                    f"{max(0, int(backfill_cursor['anchor_block']) - int(backfill_cursor['last_committed'])):>10,} "
                    f"status={backfill_cursor['status']}"
                )
            log.info("heartbeat %s", line)
            db.save_rpc_health(pool.health_rows())
        aggregate = db.aggregate_counts()
        lines.append(
            "address_scans "
            + " ".join(f"{status}={count}" for status, count in sorted(aggregate.items()))
        )
        write_status(status_path, lines)
        snapshot = db.monitoring_snapshot(cfg.min_usd)
        reports: dict[str, dict[str, Any]] = {}
        export_times: list[str] = []
        for name in ("qualifying.xlsx", "below_threshold.xlsx", "incomplete.xlsx"):
            path = cfg.export_dir / name
            if path.exists():
                stat = path.stat()
                mtime = datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat()
                reports[name] = {"bytes": stat.st_size, "updated_at": mtime}
                export_times.append(mtime)
        wal = Path(str(db.path) + "-wal")
        monitor.add_aggregate_sample(
            run_id=run_id, **snapshot,
            db_bytes=db.path.stat().st_size if db.path.exists() else 0,
            wal_bytes=wal.stat().st_size if wal.exists() else 0,
            reports_json=reports,
            last_export_at=max(export_times) if export_times else None,
        )
        prune_day = sampled_at.strftime("%Y%m%d")
        if prune_day != last_prune_day:
            monitor.prune(30)
            last_prune_day = prune_day
        # короткий пинг в консоль, чтобы по ssh было видно что жив
        console.info("heartbeat  " + " | ".join(
            f"{c.key}:live={int(db.cursor(c.key, 'live')['last_committed']) if db.cursor(c.key, 'live') else 'n/a'}"
            for c in selected
        ))


async def export_loop(
    db: DB,
    cfg: AppCfg,
    selected: list[ChainCfg],
    stop: asyncio.Event,
    monitor: MonitorStore | None = None,
    sui_store: SuiStore | None = None,
    export_lock: asyncio.Lock | None = None,
) -> None:
    mapping = {c.key: c for c in selected}
    export_lock = export_lock or asyncio.Lock()
    last_export_revision = int(
        monitor.setting("last_export_revision", "-1") if monitor else -1
    )
    while not stop.is_set():
        try:
            await asyncio.wait_for(stop.wait(), timeout=cfg.export_every_sec)
            break
        except asyncio.TimeoutError:
            pass
        try:
            revision = db.data_revision()
            if revision != last_export_revision:
                if monitor:
                    monitor.set_setting("exporter_state", "running")
                await run_export_process(db, cfg, mapping, export_lock)
                last_export_revision = revision
                if monitor:
                    monitor.set_setting("last_export_revision", revision)
            local = datetime.now(timezone.utc).astimezone(ZoneInfo("Europe/Moscow"))
            snapshot_day = monitor.setting("last_snapshot_day", "") if monitor else ""
            if cfg.daily_snapshot and local.hour >= 3 and snapshot_day != local.strftime("%Y%m%d"):
                await asyncio.to_thread(copy_daily_report_snapshot, cfg.export_dir)
                if monitor:
                    monitor.set_setting("last_snapshot_day", local.strftime("%Y%m%d"))
            if monitor:
                monitor.set_setting("exporter_state", "idle")
                monitor.set_setting("consecutive_export_failures", "0")
                monitor.resolve_incident("export:failed")
        except Exception as e:
            log.warning("periodic export: %s", e)
            if monitor:
                monitor.set_setting("exporter_state", "failed")
                failures = int(monitor.setting("consecutive_export_failures", "0") or 0) + 1
                monitor.set_setting("consecutive_export_failures", failures)
                if failures >= 2:
                    monitor.open_incident(
                        "export:failed", "critical", "export_failure",
                        f"XLSX export failed {failures} times: {type(e).__name__}",
                    )


async def control_loop(
    db: DB,
    cfg: AppCfg,
    chains: dict[str, ChainCfg],
    monitor: MonitorStore,
    stop: asyncio.Event,
    export_lock: asyncio.Lock,
) -> None:
    while not stop.is_set():
        for request in monitor.pending_controls():
            try:
                if request.action == "pause":
                    monitor.set_runtime_state("pausing", "finishing current block/address work")
                    monitor.set_setting("scanner_paused", "1")
                    result = "pause requested"
                elif request.action == "resume":
                    monitor.set_setting("scanner_paused", "0")
                    result = "scanner resumed"
                elif request.action == "export":
                    paths = await run_export_process(db, cfg, chains, export_lock)
                    result = "exported: " + ", ".join(path.name for path in paths)
                elif request.action == "backup":
                    path = await asyncio.to_thread(backup_database, db, ROOT / "backups", 7)
                    result = f"backup: {path.name}"
                else:
                    raise ValueError("unsupported scanner control action")
                monitor.resolve_incident(f"control:{request.action}")
                monitor.finish_control(request.id, "complete", result)
            except Exception as exc:
                monitor.finish_control(request.id, "failed", f"{type(exc).__name__}: {exc}")
                monitor.open_incident(
                    f"control:{request.action}", "critical", "control_failure",
                    f"{request.action} failed: {type(exc).__name__}",
                )
        try:
            await asyncio.wait_for(stop.wait(), timeout=1)
        except asyncio.TimeoutError:
            pass


async def daily_backup_loop(
    db: DB, monitor: MonitorStore, stop: asyncio.Event,
) -> None:
    while not stop.is_set():
        last = float(monitor.setting("last_backup_at", "0") or 0)
        if time.time() - last >= 86400:
            try:
                path = await asyncio.to_thread(backup_database, db, ROOT / "backups", 7)
                monitor.set_setting("last_backup_at", time.time())
                monitor.set_setting("last_backup_file", path.name)
                monitor.resolve_incident("backup:failed")
            except Exception as exc:
                monitor.open_incident(
                    "backup:failed", "critical", "backup_failure",
                    f"Daily SQLite backup failed: {type(exc).__name__}",
                )
        try:
            await asyncio.wait_for(stop.wait(), timeout=3600)
        except asyncio.TimeoutError:
            pass


async def supervised(
    name: str, factory, stop: asyncio.Event, restart_sec: int,
    monitor: MonitorStore | None = None,
) -> None:
    while not stop.is_set():
        worker: asyncio.Task[Any] | None = None
        try:
            worker = asyncio.create_task(factory(), name=f"worker-{name}")
            done, _ = await asyncio.wait({worker}, timeout=10)
            if worker not in done and monitor:
                monitor.resolve_incident(f"worker:{name}")
            await worker
            if monitor:
                monitor.resolve_incident(f"worker:{name}")
            return
        except asyncio.CancelledError:
            if worker and not worker.done():
                worker.cancel()
                await asyncio.gather(worker, return_exceptions=True)
            raise
        except Exception:
            log.error("[%s] воркер упал, рестарт через %sс\n%s", name, restart_sec, traceback.format_exc())
            if monitor:
                monitor.open_incident(
                    f"worker:{name}", "critical", "worker_crash",
                    f"Worker {name} crashed and will restart",
                )
            try:
                await asyncio.wait_for(stop.wait(), timeout=restart_sec)
            except asyncio.TimeoutError:
                continue


async def run(args: argparse.Namespace) -> None:
    cfg = load_config(ROOT / "config.yaml")
    setup_logging(cfg.log_dir)
    monitor = MonitorStore(cfg.monitoring_db_path)
    profile = os.getenv("SCANNER_LOAD_PROFILE", "").strip().lower()
    if not profile:
        profile = (monitor.setting("load_profile", "conservative") or "conservative").lower()
    if profile not in LOAD_PROFILES:
        log.warning("unknown load profile %s; using conservative", profile)
        profile = "conservative"
    apply_load_profile(cfg, profile)
    if args.min_usd is not None:
        cfg.min_usd = args.min_usd
    if getattr(args, "rabby_fallback", None) is not None:
        cfg.rabby_fallback = args.rabby_fallback
    if getattr(args, "tx_to_contracts", None) is not None:
        cfg.discover_tx_to_contracts = args.tx_to_contracts

    wanted = (
        [item.strip() for item in args.chains.split(",") if item.strip()]
        if args.chains
        else cfg.default_chains
    )
    sui_cfg = SuiConfig.from_mapping(cfg.sui)
    unknown = [key for key in wanted if key not in cfg.chains and key != "sui"]
    if unknown:
        raise SystemExit(
            f"unknown network(s): {', '.join(unknown)}. Available: {', '.join((*cfg.chains, 'sui'))}"
        )
    selected = [cfg.chains[key] for key in wanted if key in cfg.chains]
    sui_requested = "sui" in wanted and sui_cfg.enabled
    balance_chains = {key: chain for key, chain in cfg.chains.items() if chain.enabled}

    db = DB(cfg.db_path)
    sui_store = SuiStore(cfg.db_path)
    db.seed_valuation_policies(cfg.valuation_policies)
    revalued = db.revalue_system_balances(cfg.min_usd)
    if revalued:
        log.info("revalued %s historical address scans using asset policies", revalued)
    reclassified = db.reclassify_latest_scans(cfg.min_usd)
    if reclassified:
        log.info("reclassified %s latest scans at threshold $%s", reclassified, cfg.min_usd)
    seed_tokens(db, load_token_seed(ROOT / "tokens.yaml"), list(balance_chains))
    cfg.export_dir.mkdir(parents=True, exist_ok=True)
    export_lock = asyncio.Lock()

    if args.export_only:
        await run_export_process(db, cfg, balance_chains, export_lock)
        sui_store.close()
        db.close()
        monitor.close()
        return

    stop = asyncio.Event()
    loop = asyncio.get_running_loop()
    for sig in (signal.SIGINT, signal.SIGTERM):
        try:
            loop.add_signal_handler(sig, stop.set)
        except (NotImplementedError, RuntimeError):
            try:
                signal.signal(sig, lambda *_args: loop.call_soon_threadsafe(stop.set))
            except (ValueError, OSError):
                pass
    mode = "index-only" if args.index_only else "balances-only" if args.balances_only else "full"
    run_id = None if args.rpc_check else monitor.begin_run(
        mode, cfg.min_usd, profile, vars(args)
    )
    global_rpc_sem = asyncio.Semaphore(cfg.global_rpc_concurrency)
    prices = PriceBook(cfg.http_timeout_sec, cfg.price_batch_size)
    pools: dict[str, RpcPool] = {}
    sui_client: BlockberryClient | None = None
    should_export = not args.rpc_check
    try:
        if sui_requested:
            blockberry_key = os.getenv("BLOCKBERRY_API_KEY", "").strip()
            if blockberry_key:
                sui_client = BlockberryClient(
                    blockberry_key, sui_cfg, global_rpc_sem,
                    [value.strip() for value in os.getenv("SUI_RPC", "").split(",") if value.strip()],
                )
                await sui_client.__aenter__()
            else:
                log.warning("[sui] BLOCKBERRY_API_KEY is not configured; Sui is disabled")
        needed = (
            balance_chains
            if not args.index_only or args.rpc_check
            else {chain.key: chain for chain in selected}
        )
        for chain in needed.values():
            if not chain.rpc:
                log.warning("[%s] no RPC configured", chain.key)
                continue
            pool = RpcPool(
                chain.key,
                chain.rpc,
                cfg,
                expected_chain_id=chain.chain_id,
                global_sem=global_rpc_sem,
                health_store=db,
            )
            await pool.__aenter__()
            pools[chain.key] = pool

        reports = await asyncio.gather(
            *(pool.preflight() for pool in pools.values()),
            return_exceptions=True,
        )
        for (key, _pool), report in zip(pools.items(), reports):
            if isinstance(report, Exception):
                log.warning("[%s] RPC preflight failed: %s", key, report)
                if args.rpc_check:
                    console.info(f"{key}: ERROR {type(report).__name__}")
                continue
            ok_count = sum(1 for row in report if row.get("ok"))
            log.info("[%s] RPC preflight %s/%s usable", key, ok_count, len(report))
            if args.rpc_check:
                heads = [row["head"] for row in report if row.get("ok") and row.get("head")]
                details = ", ".join(
                    f"{row['endpoint']}={'ok' if row.get('ok') else row.get('error', 'error')}"
                    for row in report
                )
                head_text = f" head={max(heads):,}" if heads else ""
                console.info(f"{key}: {ok_count}/{len(report)}{head_text}  {details}")

        if sui_requested:
            if sui_client is None:
                if args.rpc_check:
                    console.info("sui: ERROR missing BLOCKBERRY_API_KEY")
            else:
                try:
                    report = await sui_client.preflight()
                    capabilities = (
                        f"discovery={'ok' if report.get('discovery_ok') else 'error'} "
                        f"defi={'ok' if report.get('defi_ok') else report.get('defi_error') or 'error'}"
                    )
                    if args.rpc_check:
                        console.info(
                            f"sui: ok head={report['head']:,} provider=Blockberry-indexed "
                            f"{capabilities}"
                        )
                    else:
                        log.info(
                            "[sui] Blockberry preflight head=%s %s",
                            report["head"], capabilities,
                        )
                except Exception as exc:
                    log.warning("[sui] Blockberry preflight failed: %s", type(exc).__name__)
                    if args.rpc_check:
                        console.info(f"sui: ERROR {type(exc).__name__}")
                    await sui_client.__aexit__(None, None, None)
                    sui_client = None

        if args.rpc_check:
            return

        discovery_ready = [
            chain for chain in selected
            if chain.key in pools and pools[chain.key].usable()
        ]
        if not discovery_ready and sui_client is None and not args.balances_only:
            raise RuntimeError("none of the selected discovery networks has an RPC")
        await print_startup_status(discovery_ready, pools, db)
        for chain in discovery_ready:
            db.init_chain(chain.key, chain.start_block, cfg.discover_tx_to_contracts)
            head = pools[chain.key].last_head
            if head is None:
                head = await latest_block(pools[chain.key])
            db.init_chain_cursors(
                chain.key, chain.start_block,
                max(0, int(head) - chain.confirmations), lookback=2,
            )

        if args.once:
            if cfg.run_indexer and not args.balances_only:
                index_jobs = [
                        index_chain(
                            db,
                            chain,
                            pools[chain.key],
                            cfg,
                            stop,
                            args.from_block,
                            args.to_block,
                            once=True,
                            monitor=monitor,
                        )
                        for chain in discovery_ready
                ]
                if sui_client is not None:
                    index_jobs.append(sui_discovery_loop(
                        sui_store, sui_client, sui_cfg, stop, once=True,
                        from_checkpoint=args.from_block, to_checkpoint=args.to_block,
                        monitor=monitor, run_id=run_id,
                    ))
                await asyncio.gather(*index_jobs)
            if cfg.run_balance_checker and not args.index_only:
                balance_jobs = [check_multichain_balances(
                    db, balance_chains, pools, prices, cfg, cfg.min_usd, stop,
                    once=True, monitor=monitor,
                )]
                if sui_client is not None:
                    balance_jobs.append(sui_balance_loop(
                        sui_store, sui_client, sui_cfg, cfg.min_usd, stop, once=True,
                    ))
                await asyncio.gather(*balance_jobs)
            await run_export_process(db, cfg, balance_chains, export_lock)
            should_export = False
            return

        tasks: list[asyncio.Task[Any]] = []
        discovery_slots = DiscoverySlots(live_slots=2, backfill_slots=1)
        if cfg.run_indexer and not args.balances_only:
            for chain in discovery_ready:
                pool = pools[chain.key]
                for role in ("live", "backfill"):
                    tasks.append(
                        asyncio.create_task(
                            supervised(
                                f"idx-{chain.key}-{role}",
                                lambda c=chain, p=pool, r=role: index_chain_cursor(
                                    db, c, p, cfg, stop, r, discovery_slots,
                                    args.from_block, args.to_block, monitor,
                                ),
                                stop, cfg.task_restart_sec, monitor,
                            ),
                            name=f"idx-{chain.key}-{role}",
                        )
                    )
            if sui_client is not None:
                tasks.append(asyncio.create_task(
                    supervised(
                        "idx-sui",
                        lambda: sui_discovery_loop(
                            sui_store, sui_client, sui_cfg, stop,
                            from_checkpoint=args.from_block, to_checkpoint=args.to_block,
                            monitor=monitor, run_id=run_id,
                        ),
                        stop, cfg.task_restart_sec, monitor,
                    ),
                    name="idx-sui",
                ))
        if cfg.run_balance_checker and not args.index_only:
            tasks.append(
                asyncio.create_task(
                    supervised(
                        "balances",
                        lambda: check_multichain_balances(
                            db,
                            balance_chains,
                            pools,
                            prices,
                            cfg,
                            cfg.min_usd,
                            stop,
                            monitor=monitor,
                        ),
                        stop,
                        cfg.task_restart_sec,
                        monitor,
                    ),
                    name="balances",
                )
            )
            if sui_client is not None:
                tasks.append(asyncio.create_task(
                    supervised(
                        "balances-sui",
                        lambda: sui_balance_loop(
                            sui_store, sui_client, sui_cfg, cfg.min_usd, stop,
                        ),
                        stop, cfg.task_restart_sec, monitor,
                    ),
                    name="balances-sui",
                ))
        tasks.append(
            asyncio.create_task(
                export_loop(
                    db, cfg, list(balance_chains.values()), stop, monitor,
                    sui_store, export_lock,
                ),
                name="xlsx",
            )
        )
        tasks.append(
            asyncio.create_task(
                heartbeat_loop(db, cfg, discovery_ready, pools, stop, monitor, run_id),
                name="heartbeat",
            )
        )
        tasks.append(
            asyncio.create_task(
                control_loop(db, cfg, balance_chains, monitor, stop, export_lock),
                name="control",
            )
        )
        tasks.append(
            asyncio.create_task(daily_backup_loop(db, monitor, stop), name="daily-backup")
        )
        await asyncio.gather(*tasks)
    except asyncio.CancelledError:
        pass
    finally:
        if run_id is not None:
            monitor.heartbeat(run_id, "stopping")
        stop.set()
        if should_export:
            try:
                await run_export_process(db, cfg, balance_chains, export_lock)
            except Exception as exc:
                log.warning("final export: %s", exc)
        for pool in pools.values():
            await pool.__aexit__(None, None, None)
        if sui_client is not None:
            await sui_client.__aexit__(None, None, None)
        sui_store.close()
        db.close()
        if run_id is not None:
            monitor.finish_run(run_id)
        monitor.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="DeFi contract liquidity scanner")
    parser.add_argument("--chains", help="discovery networks, comma separated")
    parser.add_argument("--from-block", type=int, dest="from_block")
    parser.add_argument("--to-block", type=int, dest="to_block")
    parser.add_argument("--min-usd", type=float, dest="min_usd")
    parser.add_argument("--index-only", action="store_true")
    parser.add_argument("--balances-only", action="store_true")
    parser.add_argument("--export-only", action="store_true")
    parser.add_argument("--rpc-check", action="store_true", help="validate all configured RPCs and exit")
    parser.add_argument("--once", action="store_true", help="finish the current finite pass and exit")
    parser.add_argument(
        "--rabby-fallback", action=argparse.BooleanOptionalAction, default=None,
        help="supplement failed/slow RPC scans with a separately labelled Rabby estimate",
    )
    parser.add_argument(
        "--tx-to-contracts", action=argparse.BooleanOptionalAction, default=None,
        help="discover unique top-level tx.to recipients that currently have bytecode",
    )
    args = parser.parse_args()
    try:
        asyncio.run(run(args))
    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()
